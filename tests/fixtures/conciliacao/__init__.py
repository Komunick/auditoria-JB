"""Geradores de fixtures SINTETICAS da Conciliacao Fiscal.

Nenhuma planilha real entra no Git (constituicao, principio VI). Tudo aqui e
gerado em tempo de teste, numa pasta temporaria, a partir de CNPJ sintetico
valido, razao social ficticia e valores inventados.

Os layouts seguem `specs/003-conciliacao-fiscal/contracts/xlsx-inputs.md`:

- Quadro 50-5: bloco-resumo com "Receita Total Informada" e "Receita Total
  calculada"; sem DIMP (PIX/nao-PIX ficam NULL, nunca zero).
- Quadro 50-3/DIMP: o mesmo bloco-resumo mais o bloco de movimentos
  reconhecido por "Instituicao Financeira" + "Total DIMP".
- Planilha-mestre: secoes anuais "RELATORIO AAAA", meses reconhecidos pelo
  texto e uma linha TOTAL por secao.

As variacoes invalidas (pacote falso, macro, vinculo externo, traversal, razao
de compressao) sao produzidas por manipulacao direta do ZIP, porque o objetivo
e exercitar a validacao ANTES do openpyxl.
"""

from __future__ import annotations

import os
import shutil
import zipfile
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

VERSAO_FIXTURES = "1"

# ----------------------------------------------------------------------
# CNPJ sintetico

def dv_cnpj(base12: str) -> str:
    """Devolve os dois digitos verificadores de uma base de 12 digitos."""
    def _dv(numero: str) -> str:
        pesos = list(range(len(numero) + 1, 1, -1))
        pesos = [p if p <= 9 else p - 8 for p in pesos]
        soma = sum(int(d) * p for d, p in zip(numero, pesos))
        resto = soma % 11
        return "0" if resto < 2 else str(11 - resto)

    d1 = _dv(base12)
    return d1 + _dv(base12 + d1)


def cnpj_sintetico(base12: str = "112223330001") -> str:
    """CNPJ so-digitos, valido no calculo dos DV, sem dono real."""
    return base12 + dv_cnpj(base12)


def formatar_cnpj(cnpj: str) -> str:
    """14 digitos -> 00.000.000/0000-00 (como aparece nas fontes)."""
    d = "".join(c for c in cnpj if c.isdigit())
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"


CNPJ_A = cnpj_sintetico("112223330001")          # 11.222.333/0001-81
CNPJ_B = cnpj_sintetico("445556660001")          # outra empresa sintetica
RAZAO_A = "EMPRESA SINTETICA ALFA LTDA"
RAZAO_B = "EMPRESA SINTETICA BETA LTDA"

_IDENTIFICACAO = "SECRETARIA DA FAZENDA DO ESTADO DA BAHIA"

# Rotulos exatos do contrato (acento/caixa/espaco nao podem ser exigidos pelo
# parser, mas a fixture representa a fonte como ela chega).
ROTULOS_RESUMO = {
    "A": "Período",
    "B": "CNPJ",
    "C": "Receita Total Informada",
    "D": "Nº de documentos",
    "E": "Receita com substituição tributária de ICMS",
    "F": "Receita sem substituição tributária",
    "G": "Receita Total calculada",
    "H": "Nº de documentos calculados",
    "I": "Receita calculada com ST",
    "J": "Receita calculada sem ST",
}

ROTULOS_DIMP = {
    "A": "Período",
    "B": "CNPJ",
    "C": "Instituição Financeira",
    "D": "Débito",
    "E": "Crédito",
    "F": "Transferência",
    "G": "PIX",
    "H": "Voucher",
    "I": "Outras",
    "J": "Total DIMP",
}

MESES = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]


def dec(valor) -> Decimal:
    """Decimal de centavos exatos (o dominio nunca usa float)."""
    return Decimal(str(valor)).quantize(Decimal("0.01"))


# ----------------------------------------------------------------------
# Resumo comum aos dois layouts

def _escrever_resumo(aba, linha_cabecalho: int, *, competencia: str,
                     cnpj: str, valores: dict, rotulos: dict) -> int:
    """Escreve cabecalho + a linha mensal. Devolve a linha de dados."""
    for col, texto in rotulos.items():
        aba[f"{col}{linha_cabecalho}"] = texto
        aba[f"{col}{linha_cabecalho}"].font = Font(bold=True)

    linha = linha_cabecalho + 1
    aba[f"A{linha}"] = competencia
    aba[f"B{linha}"] = formatar_cnpj(cnpj)
    aba[f"C{linha}"] = valores["receita_declarada"]
    aba[f"D{linha}"] = valores.get("docs_declarados", 10)
    aba[f"E{linha}"] = valores["declarada_com_st"]
    aba[f"F{linha}"] = valores["declarada_sem_st"]
    aba[f"G{linha}"] = valores["receita_calculada"]
    aba[f"H{linha}"] = valores.get("docs_calculados", 10)
    aba[f"I{linha}"] = valores["calculada_com_st"]
    aba[f"J{linha}"] = valores["calculada_sem_st"]
    return linha


def _cabecalho_orgao(aba, razao_social: str, cnpj: str, titulo: str) -> None:
    aba["A1"] = _IDENTIFICACAO
    aba["A1"].font = Font(bold=True)
    aba["A2"] = titulo
    aba["A3"] = f"Contribuinte: {razao_social}"
    aba["A4"] = f"CNPJ: {formatar_cnpj(cnpj)}"


def valores_resumo(*, declarada_com_st="10000.00", declarada_sem_st="30000.00",
                   calculada_com_st="10500.00",
                   calculada_sem_st="31000.00") -> dict:
    """Resumo que FECHA: com ST + sem ST = total, nos dois lados."""
    d_com, d_sem = dec(declarada_com_st), dec(declarada_sem_st)
    c_com, c_sem = dec(calculada_com_st), dec(calculada_sem_st)
    return {
        "receita_declarada": d_com + d_sem,
        "declarada_com_st": d_com,
        "declarada_sem_st": d_sem,
        "receita_calculada": c_com + c_sem,
        "calculada_com_st": c_com,
        "calculada_sem_st": c_sem,
    }


# ----------------------------------------------------------------------
# Quadro 50-5

def gerar_quadro_50_5(caminho: str, *, cnpj: str = CNPJ_A,
                      competencia: str = "202401",
                      razao_social: str = RAZAO_A,
                      valores: dict | None = None,
                      linha_cabecalho: int = 6,
                      rotulos: dict | None = None,
                      abas_extras: tuple[str, ...] = (),
                      aba_oculta: str = "") -> str:
    """Quadro 50-5 valido: um relatorio, um CNPJ, uma competencia, uma linha.

    `linha_cabecalho` permite gerar a variacao com cabecalho deslocado.
    `abas_extras`/`aba_oculta` exercitam a varredura de multiplas abas sem
    tornar o relatorio ambiguo (as extras nao levam bloco reconhecivel).
    """
    wb = Workbook()
    aba = wb.active
    aba.title = "Quadro 50-5"
    _cabecalho_orgao(aba, razao_social, cnpj,
                     "Quadro 50-5 - Receita Declarada x Receita Calculada")
    _escrever_resumo(aba, linha_cabecalho, competencia=competencia, cnpj=cnpj,
                     valores=valores or valores_resumo(),
                     rotulos=rotulos or ROTULOS_RESUMO)
    for nome in abas_extras:
        extra = wb.create_sheet(nome)
        extra["A1"] = "Observações gerais do relatório"
    if aba_oculta:
        oculta = wb.create_sheet(aba_oculta)
        oculta["A1"] = "Notas internas"
        oculta.sheet_state = "hidden"
    wb.save(caminho)
    return caminho


# ----------------------------------------------------------------------
# Quadro 50-3 / DIMP

def movimento(instituicao: str, *, debito="0.00", credito="0.00",
              transferencia="0.00", pix="0.00", voucher="0.00",
              outras="0.00") -> dict:
    """Movimento DIMP que FECHA: total = soma dos seis componentes."""
    partes = {
        "debito": dec(debito), "credito": dec(credito),
        "transferencia": dec(transferencia), "pix": dec(pix),
        "voucher": dec(voucher), "outras": dec(outras),
    }
    partes["instituicao"] = instituicao
    partes["total_dimp"] = sum(
        (partes[k] for k in ("debito", "credito", "transferencia", "pix",
                             "voucher", "outras")), Decimal("0.00"))
    return partes


MOVIMENTOS_PADRAO = (
    movimento("BANCO SINTETICO S.A.", debito="4000.00", credito="6000.00",
              pix="7000.00"),
    movimento("ADQUIRENTE SINTETICA LTDA", credito="9000.00",
              transferencia="1500.00", outras="500.00"),
    # Mesma instituicao aparecendo de novo: a linha original nao pode sumir na
    # agregacao da interface.
    movimento("BANCO SINTETICO S.A.", pix="2000.00", voucher="1000.00"),
)


def gerar_quadro_50_3_dimp(caminho: str, *, cnpj: str = CNPJ_A,
                           competencia: str = "202401",
                           razao_social: str = RAZAO_A,
                           valores: dict | None = None,
                           movimentos=MOVIMENTOS_PADRAO,
                           linha_cabecalho: int = 6,
                           linhas_estranhas=(),
                           incluir_bloco_dimp: bool = True) -> str:
    """Quadro 50-3 com bloco DIMP.

    `linhas_estranhas` recebe tuplas (competencia, cnpj, movimento) para as
    linhas de outro periodo/CNPJ que devem ser EXCLUIDAS com aviso.
    `incluir_bloco_dimp=False` produz o caso DIMP_OBRIGATORIA_AUSENTE.
    """
    wb = Workbook()
    aba = wb.active
    aba.title = "Quadro 50-3"
    _cabecalho_orgao(aba, razao_social, cnpj,
                     "Quadro 50-3 - Receita Calculada x DIMP")
    linha_dados = _escrever_resumo(aba, linha_cabecalho,
                                   competencia=competencia, cnpj=cnpj,
                                   valores=valores or valores_resumo(),
                                   rotulos=ROTULOS_RESUMO)

    if not incluir_bloco_dimp:
        wb.save(caminho)
        return caminho

    linha = linha_dados + 3
    aba[f"A{linha}"] = "Demonstrativo DIMP"
    linha += 1
    for col, texto in ROTULOS_DIMP.items():
        aba[f"{col}{linha}"] = texto
        aba[f"{col}{linha}"].font = Font(bold=True)

    todas = [(competencia, cnpj, m) for m in movimentos]
    todas.extend(linhas_estranhas)
    for comp, doc, mov in todas:
        linha += 1
        aba[f"A{linha}"] = comp
        aba[f"B{linha}"] = formatar_cnpj(doc) if doc else ""
        aba[f"C{linha}"] = mov["instituicao"]
        aba[f"D{linha}"] = mov["debito"]
        aba[f"E{linha}"] = mov["credito"]
        aba[f"F{linha}"] = mov["transferencia"]
        aba[f"G{linha}"] = mov["pix"]
        aba[f"H{linha}"] = mov["voucher"]
        aba[f"I{linha}"] = mov["outras"]
        aba[f"J{linha}"] = mov["total_dimp"]

    wb.save(caminho)
    return caminho


def total_pix(movimentos=MOVIMENTOS_PADRAO) -> Decimal:
    return sum((m["pix"] for m in movimentos), Decimal("0.00"))


def total_dimp(movimentos=MOVIMENTOS_PADRAO) -> Decimal:
    return sum((m["total_dimp"] for m in movimentos), Decimal("0.00"))


def total_nao_pix(movimentos=MOVIMENTOS_PADRAO) -> Decimal:
    return total_dimp(movimentos) - total_pix(movimentos)


# ----------------------------------------------------------------------
# Planilha-mestre (modelo recebido para preenchimento)

def gerar_modelo_mestre(caminho: str, *, anos=(2023, 2024),
                        cnpj: str = CNPJ_A, razao_social: str = RAZAO_A,
                        meses=None, com_total: bool = True) -> str:
    """Modelo com secoes anuais 'RELATORIO AAAA' e uma linha por mes.

    Traz celulas, formulas, estilos, larguras e mesclagens FORA do mapa
    (C, D, E, G, H, I, J, L, M) para o teste de preservacao.
    """
    meses = list(meses or MESES)
    wb = Workbook()
    aba = wb.active
    aba.title = "Conciliação"

    aba["A1"] = razao_social
    aba["A1"].font = Font(bold=True, size=14)
    aba["A2"] = f"CNPJ: {formatar_cnpj(cnpj)}"
    aba.merge_cells("A1:M1")
    for col, largura in (("A", 18), ("B", 12), ("C", 16), ("G", 16),
                         ("L", 40), ("M", 40)):
        aba.column_dimensions[col].width = largura

    linha = 4
    for ano in anos:
        aba[f"A{linha}"] = f"RELATÓRIO {ano}"
        aba[f"A{linha}"].font = Font(bold=True)
        linha += 1
        # Rotulos conforme o mapeamento de contracts/xlsx-outputs.md: H e a
        # diferenca (G-C) e J e a calculada SEM ST derivada por formula
        # (G-I) — o legado calcula, nao copia.
        for col, texto in (("A", "MÊS"), ("C", "RECEITA DECLARADA"),
                           ("D", "DECLARADA COM ST"),
                           ("E", "DECLARADA SEM ST"),
                           ("G", "RECEITA CALCULADA"),
                           ("H", "DIFERENÇA (CALC - DECL)"),
                           ("I", "CALCULADA COM ST"),
                           ("J", "CALCULADA SEM ST"), ("L", "PIX"),
                           ("M", "CARTÃO/OUTROS")):
            aba[f"{col}{linha}"] = texto
            aba[f"{col}{linha}"].font = Font(bold=True)
        linha += 1
        primeira = linha
        for mes in meses:
            aba[f"A{linha}"] = mes
            # Coluna B fica FORA do mapa: guarda uma marca que a copia deve
            # preservar intacta.
            aba[f"B{linha}"] = f"ref-{ano}-{linha}"
            aba[f"L{linha}"].alignment = Alignment(wrap_text=True)
            aba[f"M{linha}"].alignment = Alignment(wrap_text=True)
            linha += 1
        if com_total:
            aba[f"A{linha}"] = "TOTAL"
            aba[f"A{linha}"].font = Font(bold=True)
            linha += 1
        # Rodape fora do mapa, com formula propria que nao pode ser tocada.
        aba[f"A{linha}"] = "Conferido por"
        aba[f"B{linha}"] = f"=COUNTA(A{primeira}:A{primeira + len(meses) - 1})"
        linha += 3

    resumo = wb.create_sheet("Instruções")
    resumo["A1"] = "Preencher somente as colunas monetárias."
    wb.save(caminho)
    return caminho


def celulas_fora_do_mapa(caminho: str) -> dict:
    """Snapshot de tudo que o preenchimento NAO pode alterar."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho)
    mapeadas = {"C", "D", "E", "G", "H", "I", "J", "L", "M"}
    snapshot: dict = {}
    for aba in wb.worksheets:
        for linha in aba.iter_rows():
            for celula in linha:
                coluna = get_column_letter(celula.column)
                if aba.title == "Conciliação" and coluna in mapeadas:
                    continue
                if celula.value is not None:
                    snapshot[f"{aba.title}!{celula.coordinate}"] = celula.value
        snapshot[f"{aba.title}!__larguras"] = {
            k: v.width for k, v in aba.column_dimensions.items()}
        snapshot[f"{aba.title}!__mesclagens"] = sorted(
            str(r) for r in aba.merged_cells.ranges)
    wb.close()
    return snapshot


# ----------------------------------------------------------------------
# Variacoes INVALIDAS (validacao do pacote, antes do openpyxl)

def gerar_nao_xlsx(caminho: str) -> str:
    """Extensao .xlsx num arquivo que nem sequer e ZIP."""
    with open(caminho, "wb") as saida:
        saida.write(b"isto nao e um pacote Open XML\n" * 10)
    return caminho


def gerar_xlsx_truncado(caminho: str, origem: str) -> str:
    """Metade de um XLSX valido: ZIP quebrado."""
    with open(origem, "rb") as entrada:
        dados = entrada.read()
    with open(caminho, "wb") as saida:
        saida.write(dados[: len(dados) // 2])
    return caminho


def gerar_xlsx_sem_estrutura_minima(caminho: str) -> str:
    """ZIP valido, mas sem [Content_Types].xml nem xl/workbook.xml."""
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("qualquer/coisa.txt", "conteudo")
    return caminho


def _copiar_com_extras(caminho: str, origem: str, extras: dict) -> str:
    """Copia um XLSX valido acrescentando entradas ao pacote."""
    shutil.copyfile(origem, caminho)
    with zipfile.ZipFile(caminho, "a", zipfile.ZIP_DEFLATED) as zf:
        for nome, conteudo in extras.items():
            zf.writestr(nome, conteudo)
    return caminho


def gerar_xlsx_com_macro(caminho: str, origem: str) -> str:
    return _copiar_com_extras(caminho, origem,
                              {"xl/vbaProject.bin": b"\x00MACRO\x00"})


def gerar_xlsx_com_vinculo_externo(caminho: str, origem: str) -> str:
    return _copiar_com_extras(caminho, origem, {
        "xl/externalLinks/externalLink1.xml":
            '<externalLink xmlns="http://schemas.openxmlformats.org/'
            'spreadsheetml/2006/main"/>'})


def gerar_xlsx_com_traversal(caminho: str, origem: str) -> str:
    return _copiar_com_extras(caminho, origem,
                              {"../fora_do_pacote.xml": "<a/>"})


def gerar_xlsx_muitas_entradas(caminho: str, origem: str,
                               quantidade: int = 5001) -> str:
    extras = {f"xl/lixo/{i}.xml": "<a/>" for i in range(quantidade)}
    return _copiar_com_extras(caminho, origem, extras)


def gerar_xlsx_razao_insegura(caminho: str, origem: str,
                              megabytes: int = 130) -> str:
    """Entrada que descompacta alem do limite de 120 MB e com razao > 120:1.

    Escrita em blocos de 1 MB: a fixture nao pode precisar de 130 MB de RAM
    so para provar que a validacao rejeita o arquivo.
    """
    shutil.copyfile(caminho if caminho == origem else origem, caminho)
    bloco = b"\x00" * (1024 * 1024)
    with zipfile.ZipFile(caminho, "a", zipfile.ZIP_DEFLATED) as zf:
        with zf.open("xl/bomba.xml", "w") as entrada:
            for _ in range(megabytes):
                entrada.write(bloco)
    return caminho


# ----------------------------------------------------------------------
# Lote completo pronto para os testes de importacao

def gerar_lote_sintetico(pasta: str) -> dict:
    """Gera o lote do cenario ponta a ponta do quickstart.

    Devolve um dicionario nome-logico -> caminho. Os nomes de arquivo sao
    propositalmente enganosos: o contrato diz que o NOME nunca determina CNPJ,
    competencia ou layout.
    """
    os.makedirs(pasta, exist_ok=True)
    caminhos: dict[str, str] = {}

    def _p(nome: str) -> str:
        return os.path.join(pasta, nome)

    caminhos["quadro_50_5"] = gerar_quadro_50_5(
        _p("relatorio_a.xlsx"))
    caminhos["quadro_50_3"] = gerar_quadro_50_3_dimp(
        _p("relatorio_b.xlsx"), competencia="202402")
    # Duplicata EXATA renomeada: mesmo SHA-256, nome diferente.
    caminhos["duplicata"] = _p("copia_do_relatorio_a.xlsx")
    shutil.copyfile(caminhos["quadro_50_5"], caminhos["duplicata"])
    # Retificacao: mesma chave (CNPJ+competencia), R$ 0,01 a mais -> conflito.
    caminhos["retificada"] = gerar_quadro_50_5(
        _p("relatorio_a_retificado.xlsx"),
        valores=valores_resumo(declarada_sem_st="30000.01"))
    caminhos["invalido"] = gerar_nao_xlsx(_p("relatorio_corrompido.xlsx"))
    return caminhos
