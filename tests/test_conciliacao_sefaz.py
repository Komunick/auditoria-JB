"""Parser dos relatorios da Malha Fiscal da SEFAZ/BA (nucleo puro).

Cobre os dois layouts (Quadro 50-5 e Quadro 50-3/DIMP), a leitura de numeros
em formato brasileiro e americano, CNPJ e competencia, a distincao entre
AUSENTE e ZERO, formula sem valor calculado, cabecalho deslocado, multiplas
abas (inclusive ocultas), as tolerancias de R$ 0,01, a proveniencia de cada
fato, os movimentos DIMP, as linhas de outro CNPJ/periodo, o layout
desconhecido, a validacao de seguranca do pacote XLSX antes do openpyxl e os
limites arquiteturais do modulo.

Fixtures sao geradas em tempo de execucao por `tests/fixtures/conciliacao`:
nenhum arquivo fiscal real entra no repositorio.
"""

from __future__ import annotations

import ast
import os
import shutil
import sys
import tempfile
from decimal import Decimal

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "src"))
sys.path.insert(0, os.path.join(RAIZ, "tests"))

from fixtures import conciliacao as fx  # noqa: E402

from auditoria_fiscal.core import conciliacao_sefaz as cs  # noqa: E402


def checar(cond, msg):
    if not cond:
        print(f"FALHOU - {msg}")
        raise SystemExit(1)


def erro_de(caminho: str, **kwargs) -> str:
    """Le o relatorio esperando falha fatal e devolve o codigo estavel."""
    try:
        cs.ler_relatorio(caminho, **kwargs)
    except cs.ErroConciliacao as exc:
        return exc.codigo
    return ""


def codigos(relatorio, severidade: str = "") -> set[str]:
    return {e.codigo for e in relatorio.excecoes
            if not severidade or e.severidade == severidade}


def fato(relatorio, metrica: str):
    achados = [f for f in relatorio.fatos if f.metrica == metrica]
    checar(len(achados) == 1,
           f"deveria haver exatamente 1 fato '{metrica}': {len(achados)}")
    return achados[0]


# ----------------------------------------------------------------------


def testar_numeros() -> None:
    """Texto vira Decimal so quando o separador decimal e' inequivoco."""
    for texto, esperado in (
        ("1234,56", "1234.56"), ("1.234,56", "1234.56"),
        ("1234.56", "1234.56"), ("1,234.56", "1234.56"),
        ("1.234.567,89", "1234567.89"), ("1,234,567.89", "1234567.89"),
        ("0,00", "0.00"), ("-1.234,56", "-1234.56"), ("  42  ", "42.00"),
        ("1.234.567", "1234567.00"),   # varios pontos: so pode ser milhar
        ("1,234,567", "1234567.00"),
    ):
        obtido = cs.para_decimal(texto)
        checar(obtido == Decimal(esperado),
               f"{texto!r} -> {obtido}, esperado {esperado}")

    # AMBIGUO: um unico separador com exatamente 3 digitos depois pode ser
    # milhar ou decimal. Inventar a interpretacao aqui erraria por 1000x.
    for texto in ("1.234", "1,234", "12.345", "1.2.3", "R$ 10,00", "dez",
                  "1,23,45"):
        try:
            valor = cs.para_decimal(texto)
        except cs.ErroConciliacao as exc:
            checar(exc.codigo == cs.VALOR_INVALIDO,
                   f"{texto!r} deveria dar VALOR_INVALIDO: {exc.codigo}")
            continue
        checar(False, f"{texto!r} deveria ser rejeitado, virou {valor}")

    # Celula numerica: o openpyxl devolve float. Decimal(str(v)) recupera o
    # valor gravado; Decimal(v) traria residuo binario.
    checar(cs.para_decimal(40000.01) == Decimal("40000.01"),
           f"float: {cs.para_decimal(40000.01)}")
    checar(cs.para_decimal(0) == Decimal("0.00"), "zero numerico e' valido")

    # Sub-centavo em campo monetario nao e' arredondado em silencio.
    try:
        cs.para_decimal(1234.5678)
        checar(False, "sub-centavo deveria ser rejeitado")
    except cs.ErroConciliacao as exc:
        checar(exc.codigo == cs.VALOR_INVALIDO, exc.codigo)

    # Vazio NAO e' zero (constituicao, principio I).
    for vazio in (None, "", "   "):
        checar(cs.para_decimal_opcional(vazio) is None,
               f"{vazio!r} deveria ser None, nunca 0")
    checar(cs.para_decimal_opcional("0,00") == Decimal("0.00"),
           "zero informado continua zero")


def testar_cnpj_e_competencia() -> None:
    checar(cs.normalizar_cnpj("11.222.333/0001-81") == fx.CNPJ_A,
           "CNPJ formatado deveria normalizar")
    for ruim in ("11.222.333/0001-80", "123", "", "00.000.000/0000-00",
                 "11111111111111"):
        try:
            cs.normalizar_cnpj(ruim)
            checar(False, f"CNPJ {ruim!r} deveria ser rejeitado")
        except cs.ErroConciliacao as exc:
            checar(exc.codigo == cs.CNPJ_INVALIDO, f"{ruim}: {exc.codigo}")

    checar(cs.normalizar_competencia("202401") == "202401", "competencia crua")
    checar(cs.normalizar_competencia("01/2024") == "202401", "mes/ano")
    for ruim in ("202413", "202400", "2024", "", "abcdef", "241"):
        try:
            cs.normalizar_competencia(ruim)
            checar(False, f"competencia {ruim!r} deveria ser rejeitada")
        except cs.ErroConciliacao as exc:
            checar(exc.codigo == cs.COMPETENCIA_INVALIDA, f"{ruim}: {exc.codigo}")


def testar_rotulos() -> None:
    """Acento, caixa e espaco repetido nao mudam o reconhecimento."""
    base = cs.normalizar_rotulo("Receita Total Informada")
    for variante in ("RECEITA TOTAL INFORMADA", "receita  total   informada",
                     "Receita Total Informada ", "Recéita Tótal Infórmada"):
        checar(cs.normalizar_rotulo(variante) == base,
               f"variante nao normalizou: {variante!r}")
    checar(cs.normalizar_rotulo("Instituição Financeira")
           == cs.normalizar_rotulo("INSTITUICAO FINANCEIRA"),
           "instituicao financeira com e sem acento")


def testar_quadro_50_5(pasta: str) -> None:
    caminho = fx.gerar_quadro_50_5(os.path.join(pasta, "q505.xlsx"))
    rel = cs.ler_relatorio(caminho, sha256="a" * 64, nome_original="q505.xlsx")

    checar(rel.layout == cs.LAYOUT_50_5, f"layout: {rel.layout}")
    checar(rel.cnpj == fx.CNPJ_A, f"cnpj: {rel.cnpj}")
    checar(rel.competencia == "202401", f"competencia: {rel.competencia}")
    checar(rel.receita_declarada == Decimal("40000.00"), rel.receita_declarada)
    checar(rel.declarada_com_st == Decimal("10000.00"), rel.declarada_com_st)
    checar(rel.declarada_sem_st == Decimal("30000.00"), rel.declarada_sem_st)
    checar(rel.receita_calculada == Decimal("41500.00"), rel.receita_calculada)
    checar(rel.calculada_com_st == Decimal("10500.00"), rel.calculada_com_st)
    checar(rel.calculada_sem_st == Decimal("31000.00"), rel.calculada_sem_st)

    # Diferenca e' SEMPRE calculado menos declarado.
    checar(rel.diferenca_receita == Decimal("1500.00"), rel.diferenca_receita)
    checar(rel.diferenca_com_st == Decimal("500.00"), rel.diferenca_com_st)
    checar(rel.diferenca_sem_st == Decimal("1000.00"), rel.diferenca_sem_st)

    # 50-5 nao traz DIMP: PIX e nao-PIX ficam NULOS, nunca zero, e o parser
    # avisa que a comparacao com DIMP nao e' possivel nesta fonte.
    checar(rel.total_pix is None, f"total_pix deveria ser None: {rel.total_pix}")
    checar(rel.total_nao_pix is None, f"total_nao_pix: {rel.total_nao_pix}")
    checar(not rel.movimentos, f"50-5 nao tem movimentos: {len(rel.movimentos)}")
    checar(cs.DIMP_AUSENTE_NO_LAYOUT in codigos(rel, "aviso"),
           f"deveria avisar a ausencia de DIMP: {codigos(rel)}")
    checar(not codigos(rel, "bloqueio"),
           f"50-5 limpo nao deveria ter bloqueio: {codigos(rel, 'bloqueio')}")

    # Proveniencia: cada fato de celula aponta arquivo, aba, celula e regra.
    f = fato(rel, "receita_declarada")
    checar(f.tipo_origem == "celula", f.tipo_origem)
    checar(f.aba == "Quadro 50-5", f.aba)
    checar(f.celula == "C7", f"celula: {f.celula}")
    checar(f.sha256 == "a" * 64, f.sha256)
    checar(f.versao_parser == cs.VERSAO_PARSER, f.versao_parser)
    checar(f.regra_parser, "todo fato precisa da regra que o extraiu")
    checar("informada" in cs.normalizar_rotulo(f.rotulo), f.rotulo)

    # Derivado: formula canonica e fatos de origem, nunca celula.
    d = fato(rel, "diferenca_receita")
    checar(d.tipo_origem == "formula", d.tipo_origem)
    checar(d.celula is None and d.aba is None, "derivado nao tem celula")
    checar(d.formula == "receita_calculada - receita_declarada", d.formula)
    checar(set(d.fatos_origem) == {"receita_calculada", "receita_declarada"},
           d.fatos_origem)

    checar(rel.aba_origem == "Quadro 50-5", rel.aba_origem)
    checar(rel.linha_resumo == 7, rel.linha_resumo)
    checar(rel.versao_parser == cs.VERSAO_PARSER, rel.versao_parser)


def testar_cabecalho_deslocado_e_abas(pasta: str) -> None:
    """Deslocar linhas e acrescentar abas nao pode quebrar o reconhecimento."""
    caminho = fx.gerar_quadro_50_5(
        os.path.join(pasta, "deslocado.xlsx"), linha_cabecalho=14,
        abas_extras=("Observações", "Anexo"), aba_oculta="Interno")
    rel = cs.ler_relatorio(caminho)
    checar(rel.layout == cs.LAYOUT_50_5, rel.layout)
    checar(rel.linha_resumo == 15, f"linha do resumo: {rel.linha_resumo}")
    checar(fato(rel, "receita_declarada").celula == "C15", "celula deslocada")

    # Aba oculta com bloco reconhecivel tornaria o relatorio ambiguo: cada
    # arquivo representa UM relatorio (contrato de entrada).
    duplo = fx.gerar_quadro_50_5(os.path.join(pasta, "duplo.xlsx"))
    from openpyxl import load_workbook
    wb = load_workbook(duplo)
    copia = wb.copy_worksheet(wb["Quadro 50-5"])
    copia.sheet_state = "hidden"
    wb.save(duplo)
    checar(erro_de(duplo) == cs.RELATORIO_AMBIGUO,
           "duas linhas-resumo reconhecidas deveriam dar RELATORIO_AMBIGUO")


def testar_quadro_50_3(pasta: str) -> None:
    caminho = fx.gerar_quadro_50_3_dimp(os.path.join(pasta, "q503.xlsx"))
    rel = cs.ler_relatorio(caminho, sha256="b" * 64)

    checar(rel.layout == cs.LAYOUT_50_3, f"layout: {rel.layout}")
    checar(len(rel.movimentos) == 3,
           f"tres linhas DIMP, inclusive a instituicao repetida: "
           f"{len(rel.movimentos)}")
    instituicoes = [m.instituicao for m in rel.movimentos]
    checar(instituicoes.count("BANCO SINTETICO S.A.") == 2,
           f"a linha repetida nao pode sumir: {instituicoes}")

    checar(rel.total_pix == fx.total_pix(), f"pix: {rel.total_pix}")
    checar(rel.total_nao_pix == fx.total_nao_pix(), rel.total_nao_pix)
    # nao-PIX e' total DIMP menos PIX; nao e' "cartao".
    checar(rel.total_nao_pix == fx.total_dimp() - fx.total_pix(),
           "nao-PIX = total DIMP - PIX")

    primeiro = rel.movimentos[0]
    checar(primeiro.linha_origem == 12, f"linha de origem: {primeiro.linha_origem}")
    checar(primeiro.total_dimp == primeiro.debito + primeiro.credito
           + primeiro.transferencia + primeiro.pix + primeiro.voucher
           + primeiro.outras, "componentes deveriam fechar com o total")
    checar(primeiro.celula_total == "J12", primeiro.celula_total)
    checar(primeiro.total_nao_pix == primeiro.total_dimp - primeiro.pix,
           "nao-PIX por movimento")

    # DIMP abaixo da receita calculada em mais de R$ 0,01: aviso, nao bloqueio.
    checar(cs.DIMP_ABAIXO_DA_CALCULADA in codigos(rel, "aviso"),
           f"esperava aviso de DIMP abaixo: {codigos(rel)}")
    checar(not codigos(rel, "bloqueio"), codigos(rel, "bloqueio"))

    fatos_pix = [f for f in rel.fatos if f.metrica == "total_pix"]
    checar(fatos_pix and fatos_pix[0].tipo_origem == "formula",
           "total_pix e' derivado da soma dos movimentos")


def testar_dimp_linhas_estranhas(pasta: str) -> None:
    """Linha de outro CNPJ/periodo sai com aviso; ambigua vira bloqueio."""
    outro = fx.movimento("BANCO DE OUTRA EMPRESA", pix="99.00")
    caminho = fx.gerar_quadro_50_3_dimp(
        os.path.join(pasta, "q503_estranhas.xlsx"),
        linhas_estranhas=[("202401", fx.CNPJ_B, outro),
                          ("202312", fx.CNPJ_A, outro)])
    rel = cs.ler_relatorio(caminho)
    checar(len(rel.movimentos) == 3,
           f"as duas linhas alheias nao entram: {len(rel.movimentos)}")
    checar(cs.DIMP_LINHA_DE_OUTRA_CHAVE in codigos(rel, "aviso"),
           f"a exclusao precisa ser avisada: {codigos(rel)}")
    excluidas = [e for e in rel.excecoes
                 if e.codigo == cs.DIMP_LINHA_DE_OUTRA_CHAVE]
    checar(excluidas and all("13" in e.mensagem or "14" in e.mensagem
                             or "linha" in e.mensagem.lower()
                             for e in excluidas),
           f"o aviso deveria localizar a linha: {[e.mensagem for e in excluidas]}")
    checar(rel.total_pix == fx.total_pix(),
           "o PIX nao pode incluir a linha de outra empresa")

    # Linha DIMP sem CNPJ no meio do bloco: nao da para decidir, entao bloqueia.
    sem_chave = fx.gerar_quadro_50_3_dimp(
        os.path.join(pasta, "q503_sem_chave.xlsx"),
        linhas_estranhas=[("", "", fx.movimento("BANCO SEM DONO", pix="5.00"))])
    rel2 = cs.ler_relatorio(sem_chave)
    checar(cs.DIMP_LINHA_AMBIGUA in codigos(rel2, "bloqueio"),
           f"linha sem chave deveria bloquear: {codigos(rel2)}")


def testar_fechamentos(pasta: str) -> None:
    """R$ 0,01 fecha; acima de R$ 0,01 nao."""
    # Exatamente 1 centavo de folga: dentro da tolerancia.
    no_limite = fx.gerar_quadro_50_5(
        os.path.join(pasta, "limite.xlsx"),
        valores={**fx.valores_resumo(),
                 "declarada_sem_st": Decimal("30000.01")})
    rel = cs.ler_relatorio(no_limite)
    checar(cs.FECHAMENTO_RECEITA not in codigos(rel),
           f"R$ 0,01 deveria fechar: {codigos(rel)}")

    # Dois centavos: nao fecha e vira bloqueio (o valor existe, mas alguem
    # precisa decidir antes de aprovar).
    fora = fx.gerar_quadro_50_5(
        os.path.join(pasta, "fora.xlsx"),
        valores={**fx.valores_resumo(),
                 "declarada_sem_st": Decimal("30000.02")})
    rel = cs.ler_relatorio(fora)
    checar(cs.FECHAMENTO_RECEITA in codigos(rel, "bloqueio"),
           f"R$ 0,02 nao deveria fechar: {codigos(rel)}")

    # Movimento DIMP que nao fecha com o proprio total.
    torto = dict(fx.movimento("BANCO TORTO", pix="100.00"))
    torto["total_dimp"] = Decimal("100.05")
    caminho = fx.gerar_quadro_50_3_dimp(os.path.join(pasta, "dimp_torto.xlsx"),
                                        movimentos=(torto,))
    checar(cs.FECHAMENTO_DIMP in codigos(cs.ler_relatorio(caminho), "bloqueio"),
           "componentes que nao somam o total deveriam bloquear")

    # DIMP acima da receita calculada em mais de R$ 0,01: bloqueio.
    alto = fx.movimento("BANCO ALTO", pix="99000.00")
    caminho = fx.gerar_quadro_50_3_dimp(os.path.join(pasta, "dimp_alto.xlsx"),
                                        movimentos=(alto,))
    rel = cs.ler_relatorio(caminho)
    checar(cs.DIMP_ACIMA_DA_CALCULADA in codigos(rel, "bloqueio"),
           f"DIMP acima da calculada deveria bloquear: {codigos(rel)}")


def testar_campos_e_formulas(pasta: str) -> None:
    from openpyxl import load_workbook

    # Campo obrigatorio VAZIO: erro explicito, nunca zero silencioso.
    caminho = fx.gerar_quadro_50_5(os.path.join(pasta, "vazio.xlsx"))
    wb = load_workbook(caminho)
    wb["Quadro 50-5"]["C7"] = None
    wb.save(caminho)
    checar(erro_de(caminho) == cs.CAMPO_OBRIGATORIO_AUSENTE,
           "campo obrigatorio vazio deveria falhar")

    # ZERO informado e' valido e nao pode ser confundido com ausencia.
    zerado = fx.gerar_quadro_50_5(
        os.path.join(pasta, "zerado.xlsx"),
        valores=fx.valores_resumo(declarada_com_st="0.00",
                                  declarada_sem_st="0.00"))
    rel = cs.ler_relatorio(zerado)
    checar(rel.receita_declarada == Decimal("0.00"),
           f"zero informado: {rel.receita_declarada}")

    # Formula sem valor calculado armazenado: o parser le com data_only e nao
    # pode aceitar a formula como se fosse numero.
    formula = fx.gerar_quadro_50_5(os.path.join(pasta, "formula.xlsx"))
    wb = load_workbook(formula)
    wb["Quadro 50-5"]["G7"] = "=I7+J7"
    wb.save(formula)
    checar(erro_de(formula) == cs.FORMULA_SEM_VALOR,
           "formula sem cache deveria falhar")


def testar_cnpj_divergente_e_layout(pasta: str) -> None:
    from openpyxl import load_workbook

    caminho = fx.gerar_quadro_50_5(os.path.join(pasta, "divergente.xlsx"))
    wb = load_workbook(caminho)
    wb["Quadro 50-5"]["A4"] = f"CNPJ: {fx.formatar_cnpj(fx.CNPJ_B)}"
    wb.save(caminho)
    checar(erro_de(caminho) == cs.CNPJ_DIVERGENTE,
           "cabecalho e linha com CNPJ diferente deveriam falhar")

    # Sem nenhum bloco reconhecivel.
    from openpyxl import Workbook
    vazio = os.path.join(pasta, "desconhecido.xlsx")
    wb = Workbook()
    wb.active["A1"] = "Planilha de outra coisa"
    wb.save(vazio)
    checar(erro_de(vazio) == cs.LAYOUT_NAO_RECONHECIDO,
           "layout desconhecido deveria falhar explicitamente")

    # Um relatorio que se anuncia como Quadro 50-3 e chega sem bloco DIMP
    # valido NAO pode ser reinterpretado como 50-5: seria trocar um bloqueio
    # (falta a DIMP obrigatoria) por um aviso benigno (este layout nao tem
    # DIMP), e a competencia seria aprovavel sem ninguem perceber a falta.
    sem_dimp = fx.gerar_quadro_50_3_dimp(os.path.join(pasta, "sem_dimp.xlsx"),
                                         incluir_bloco_dimp=False)
    rel = cs.ler_relatorio(sem_dimp)
    checar(rel.layout == cs.LAYOUT_50_3,
           f"o layout anunciado tem de ser preservado: {rel.layout}")
    checar(cs.DIMP_OBRIGATORIA_AUSENTE in codigos(rel, "bloqueio"),
           f"50-3 sem DIMP deveria bloquear: {codigos(rel)}")
    checar(cs.DIMP_AUSENTE_NO_LAYOUT not in codigos(rel),
           "nao pode virar o aviso benigno do 50-5")
    checar(rel.total_pix is None and rel.total_nao_pix is None,
           "sem DIMP nao ha PIX; ausencia nao vira zero")

    # O mesmo vale quando o cabecalho do bloco DIMP e' destruido.
    quebrado = fx.gerar_quadro_50_3_dimp(os.path.join(pasta, "dimp_quebrado.xlsx"))
    wb = load_workbook(quebrado)
    aba = wb["Quadro 50-3"]
    for col in "ABCDEFGHIJ":
        aba[f"{col}11"] = None          # apaga o cabecalho do bloco DIMP
    wb.save(quebrado)
    rel = cs.ler_relatorio(quebrado)
    checar(cs.DIMP_OBRIGATORIA_AUSENTE in codigos(rel, "bloqueio"),
           f"50-3 com DIMP quebrada deveria bloquear: {codigos(rel)}")


def testar_pacote_inseguro(pasta: str) -> None:
    """Tudo isso e' recusado ANTES de o openpyxl abrir o arquivo."""
    bom = fx.gerar_quadro_50_5(os.path.join(pasta, "bom.xlsx"))

    casos = [
        (fx.gerar_nao_xlsx(os.path.join(pasta, "falso.xlsx")), cs.XLSX_INVALIDO),
        (fx.gerar_xlsx_truncado(os.path.join(pasta, "trunc.xlsx"), bom),
         cs.XLSX_INVALIDO),
        (fx.gerar_xlsx_sem_estrutura_minima(os.path.join(pasta, "sem.xlsx")),
         cs.XLSX_INVALIDO),
        (fx.gerar_xlsx_com_macro(os.path.join(pasta, "macro.xlsx"), bom),
         cs.XLSX_MACRO_NAO_PERMITIDA),
        (fx.gerar_xlsx_com_vinculo_externo(os.path.join(pasta, "link.xlsx"), bom),
         cs.XLSX_VINCULO_EXTERNO),
        (fx.gerar_xlsx_com_traversal(os.path.join(pasta, "trav.xlsx"), bom),
         cs.XLSX_COMPACTACAO_INSEGURA),
        (fx.gerar_xlsx_muitas_entradas(os.path.join(pasta, "muitas.xlsx"), bom),
         cs.XLSX_COMPACTACAO_INSEGURA),
        (fx.gerar_xlsx_razao_insegura(os.path.join(pasta, "bomba.xlsx"), bom),
         cs.XLSX_COMPACTACAO_INSEGURA),
    ]
    for caminho, esperado in casos:
        nome = os.path.basename(caminho)
        try:
            cs.validar_pacote_xlsx(caminho)
            checar(False, f"{nome} deveria ser recusado")
        except cs.ErroConciliacao as exc:
            checar(exc.codigo == esperado,
                   f"{nome}: {exc.codigo}, esperado {esperado}")
        # ler_relatorio tambem valida: a porta de entrada nunca fica aberta.
        checar(erro_de(caminho) == esperado, f"ler_relatorio({nome})")

    # Extensao errada e' recusada antes de qualquer leitura.
    txt = os.path.join(pasta, "planilha.txt")
    shutil.copyfile(bom, txt)
    checar(erro_de(txt) == cs.XLSX_INVALIDO, "so .xlsx e' aceito")

    # Teto de bytes: o parser respeita o limite que a camada web informa.
    tamanho = os.path.getsize(bom)
    try:
        cs.validar_pacote_xlsx(bom, max_bytes=tamanho - 1)
        checar(False, "acima do teto deveria ser recusado")
    except cs.ErroConciliacao as exc:
        checar(exc.codigo == cs.XLSX_ACIMA_DO_LIMITE, exc.codigo)
    cs.validar_pacote_xlsx(bom, max_bytes=tamanho)   # exatamente no teto passa


def testar_limites_arquiteturais() -> None:
    """O nucleo fiscal nao pode conhecer a interface (constituicao II).

    Vale para o parser, o store e o exportador: se um deles importar FastAPI
    ou `auditoria_fiscal.web`, o dominio deixa de ser testavel sem servidor e
    a regra fiscal passa a depender de estado de requisicao.
    """
    proibidos = ("fastapi", "starlette", "flask", "pydantic", "uvicorn",
                 "PySide6")
    modulos = [
        os.path.join(RAIZ, "src", "auditoria_fiscal", "core",
                     "conciliacao_sefaz.py"),
        os.path.join(RAIZ, "src", "auditoria_fiscal", "ferramentas",
                     "conciliacao_store.py"),
        os.path.join(RAIZ, "src", "auditoria_fiscal", "ferramentas",
                     "relatorio_conciliacao.py"),
    ]
    for caminho in modulos:
        if not os.path.isfile(caminho):
            continue                     # modulo ainda nao existe nesta fase
        nome = os.path.basename(caminho)
        with open(caminho, encoding="utf-8") as arq:
            arvore = ast.parse(arq.read())
        importados: list[str] = []
        for no in ast.walk(arvore):
            if isinstance(no, ast.Import):
                importados += [a.name for a in no.names]
            elif isinstance(no, ast.ImportFrom):
                importados.append("." * no.level + (no.module or ""))
        for modulo in importados:
            raiz = modulo.split(".")[0]
            checar(raiz not in proibidos,
                   f"{nome} nao pode importar {modulo}")
            checar("auditoria_fiscal.web" not in modulo and
                   not modulo.startswith("..web"),
                   f"{nome} nao pode importar a camada web ({modulo})")


def main() -> int:
    pasta = tempfile.mkdtemp(prefix="conc_sefaz_")
    try:
        testar_numeros()
        testar_cnpj_e_competencia()
        testar_rotulos()
        testar_quadro_50_5(pasta)
        testar_cabecalho_deslocado_e_abas(pasta)
        testar_quadro_50_3(pasta)
        testar_dimp_linhas_estranhas(pasta)
        testar_fechamentos(pasta)
        testar_campos_e_formulas(pasta)
        testar_cnpj_divergente_e_layout(pasta)
        testar_pacote_inseguro(pasta)
        testar_limites_arquiteturais()
    finally:
        shutil.rmtree(pasta, ignore_errors=True)

    print("OK - parser SEFAZ (dois layouts, numeros BR/US, CNPJ e competencia, "
          "vazio versus zero, formula sem cache, cabecalho deslocado, abas "
          "extras/ocultas, tolerancia de R$ 0,01, fechamentos, proveniencia, "
          "movimentos DIMP e linhas alheias, pacote inseguro e limites "
          "arquiteturais) passou.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
