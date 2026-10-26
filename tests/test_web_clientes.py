"""Versao web — Clientes, cofre de senhas e vencimentos (spec 012).

Exercita: 401/403 com trilha, slugs fora do padrao de usuario novo, cadastro,
cofre (negado nas duas trilhas, desbloqueio com a propria senha, revelacao
auditada sem cache, limite de tentativas), cofre desligado sem chave,
importacao da planilha com previa sem senha, coleta, certificado por arquivo
sem guardar a senha, vencimentos, exportacao sem senha e varredura da pasta
de dados atras de senha em texto.
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
from datetime import datetime, timedelta, timezone

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "src"))
sys.path.insert(0, os.path.join(RAIZ, "tests"))

_TMP = tempfile.mkdtemp(prefix="auditoria_web_cli_")
os.environ["AUDITORIA_WEB_DADOS"] = _TMP

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from auditoria_fiscal.core import cofre as cf  # noqa: E402
from auditoria_fiscal.core.documentos_br import dv_cnpj  # noqa: E402
from auditoria_fiscal.web.servidor import criar_app  # noqa: E402

import test_planilha_clientes as tp  # noqa: E402

os.environ[cf.VAR_CHAVE] = cf.gerar_chave()

ADMIN = {"usuario": "dono", "nome": "Dono", "senha": "segredo1"}
ATENDENTE = {"usuario": "ana", "nome": "Ana", "senha": "segredo2"}
CONSULTA = {"usuario": "bruno", "nome": "Bruno", "senha": "segredo3"}
SENHA_CLIENTE = "Portal#Sintetico 9"
SENHA_PFX = "pfx-sintetica-77"


def cnpj(base: str) -> str:
    return base + dv_cnpj(base)


def checar(cond, msg):
    if not cond:
        print(f"FALHOU - {msg}")
        raise SystemExit(1)


def entrar(app, credenciais: dict) -> TestClient:
    cliente = TestClient(app)
    r = cliente.post("/api/login", json={"usuario": credenciais["usuario"],
                                         "senha": credenciais["senha"]})
    checar(r.status_code == 200, f"login {credenciais['usuario']}: {r.text}")
    return cliente


def eventos(adm, **filtros) -> list[dict]:
    r = adm.get("/api/admin/historico", params={"limite": 500, **filtros})
    checar(r.status_code == 200, r.status_code)
    return r.json()["itens"]


def varrer(segredos: list[str]) -> list[str]:
    """Arquivos da pasta de dados que contem alguma senha em texto."""
    achados = []
    for raiz, _, arquivos in os.walk(_TMP):
        for nome in arquivos:
            caminho = os.path.join(raiz, nome)
            with open(caminho, "rb") as arquivo:
                conteudo = arquivo.read()
            for segredo in segredos:
                if segredo.encode() in conteudo or \
                        segredo.encode("utf-16-le") in conteudo:
                    achados.append(f"{nome}: {segredo}")
    return achados


def pfx_sintetico(documento: str) -> bytes:
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import ec
    from cryptography.hazmat.primitives.serialization import pkcs12
    from cryptography.x509.oid import NameOID

    chave = ec.generate_private_key(ec.SECP256R1())
    fim = datetime.now(timezone.utc) + timedelta(days=12)
    cert = (x509.CertificateBuilder()
            .subject_name(x509.Name([x509.NameAttribute(
                NameOID.COMMON_NAME, f"EMPRESA WEB LTDA:{documento}")]))
            .issuer_name(x509.Name([x509.NameAttribute(NameOID.COMMON_NAME,
                                                       "AC SINTETICA")]))
            .public_key(chave.public_key()).serial_number(0x5151)
            .not_valid_before(fim - timedelta(days=365)).not_valid_after(fim)
            .sign(chave, hashes.SHA256()))
    return pkcs12.serialize_key_and_certificates(
        b"x", chave, cert, None,
        serialization.BestAvailableEncryption(SENHA_PFX.encode()))


def main() -> int:
    app = criar_app()
    adm = TestClient(app)
    checar(adm.post("/api/bootstrap", json=ADMIN).status_code == 200, "bootstrap")

    # ------------------------------------------------------------------
    # Portao de acesso

    anonimo = TestClient(app)
    for caminho in ("/api/clientes", "/api/clientes/catalogo",
                    "/api/clientes/avisos", "/api/clientes/vencimentos/lista"):
        checar(anonimo.get(caminho).status_code == 401, f"{caminho} sem login")

    catalogo = adm.get("/api/admin/usuarios").json()
    slugs = {i["slug"] for g in catalogo["catalogo"] for i in g["itens"]}
    novos = {"aba.clientes", "clientes.editar", "clientes.importar",
             "clientes.exportar", "cofre.ver", "cofre.editar",
             "vencimentos.configurar"}
    checar(not (novos - slugs), f"slugs ausentes: {novos - slugs}")
    checar(not (novos & set(catalogo["padrao_novo"])),
           "nenhum slug da aba 9 pode estar no padrao de usuario novo")

    r = adm.post("/api/admin/usuarios", json={
        **CONSULTA, "admin": False, "permissoes": ["aba.conferencia"]})
    checar(r.status_code == 200, r.text[:160])
    bruno_id = r.json()["id"]
    bruno = entrar(app, CONSULTA)
    checar(bruno.get("/api/clientes").status_code == 403, "sem a aba")
    r = bruno.post("/api/clientes", json={"nome": "Tentativa"})
    checar(r.status_code == 403, f"criar sem a aba: {r.status_code}")
    adm.put(f"/api/admin/usuarios/{bruno_id}/permissoes",
            json={"permissoes": ["aba.clientes"]})

    r = adm.post("/api/admin/usuarios", json={
        **ATENDENTE, "admin": False,
        "permissoes": ["aba.clientes", "clientes.editar", "cofre.ver",
                       "clientes.importar"]})
    checar(r.status_code == 200, r.text[:160])
    ana = entrar(app, ATENDENTE)

    # ------------------------------------------------------------------
    # Cadastro

    doc = cnpj("044444440001")
    checar(bruno.post("/api/clientes", json={"nome": "X"}).status_code == 403,
           "quem so consulta nao cadastra")
    r = ana.post("/api/clientes", json={"nome": "Cliente Web", "documento": doc,
                                        "ie": "NÃO TEM"})
    checar(r.status_code == 200, f"criar: {r.text[:200]}")
    cid = r.json()["id"]
    r = ana.post("/api/clientes", json={"nome": "Outro", "documento": doc})
    checar(r.status_code == 409 and r.json()["detail"]["code"]
           == "DOCUMENTO_DUPLICADO", r.text[:200])
    r = ana.post("/api/clientes", json={"nome": "X", "documento": "123"})
    checar(r.status_code == 422, f"documento invalido: {r.status_code}")

    d = bruno.get(f"/api/clientes/{cid}").json()
    checar(d["documento"].endswith(doc[-2:]) and d["ie_situacao"] == "nao_tem", d)
    r = ana.put(f"/api/clientes/{cid}", json={"nome": "Cliente Web Ltda",
                                             "documento": doc, "revision": 99})
    checar(r.status_code == 409, f"revision velha: {r.status_code}")
    r = ana.put(f"/api/clientes/{cid}", json={"nome": "Cliente Web Ltda",
                                             "documento": doc,
                                             "revision": d["revision"]})
    checar(r.status_code == 200, r.text[:200])
    r = ana.post(f"/api/clientes/{cid}/situacao", json={"situacao": "inativo"})
    checar(r.status_code == 422, "mudar situacao sem motivo")

    lista = bruno.get("/api/clientes", params={"texto": "web"}).json()
    checar(lista["total"] == 1, lista)

    # ------------------------------------------------------------------
    # Cofre

    r = ana.put(f"/api/clientes/{cid}/senhas",
                json={"campo": "prefeitura", "valor": SENHA_CLIENTE})
    checar(r.status_code == 403, "ana ve mas nao grava (sem cofre.editar)")
    r = adm.put(f"/api/clientes/{cid}/senhas",
                json={"campo": "prefeitura", "valor": SENHA_CLIENTE})
    checar(r.status_code == 200, f"admin grava: {r.text[:200]}")
    sid = r.json()["id"]
    d = bruno.get(f"/api/clientes/{cid}").json()
    checar(SENHA_CLIENTE not in json.dumps(d), "o detalhe nunca traz a senha")
    checar(d["senhas"][0]["presente"], d["senhas"])

    r = bruno.post(f"/api/clientes/senhas/{sid}/revelar")
    checar(r.status_code == 403 and SENHA_CLIENTE not in r.text, "sem cofre.ver")
    r = ana.post(f"/api/clientes/senhas/{sid}/revelar")
    checar(r.status_code == 428 and r.json()["detail"]["code"]
           == "COFRE_BLOQUEADO", f"sem desbloqueio: {r.status_code}")
    r = ana.post("/api/clientes/cofre/desbloquear", json={"senha": "errada"})
    checar(r.status_code == 422, f"senha errada nao pode dar 401: {r.status_code}")
    r = bruno.post("/api/clientes/cofre/desbloquear",
                   json={"senha": CONSULTA["senha"]})
    checar(r.status_code == 403, "sem cofre.ver nem desbloqueia")
    r = ana.post("/api/clientes/cofre/desbloquear",
                 json={"senha": ATENDENTE["senha"]})
    checar(r.status_code == 200 and r.json()["desbloqueado_ate"], r.text)
    r = ana.post(f"/api/clientes/senhas/{sid}/revelar")
    checar(r.status_code == 200 and r.json()["valor"] == SENHA_CLIENTE, r.text)
    checar(r.headers.get("cache-control") == "no-store", r.headers)
    # O desbloqueio e' da SESSAO: outro login da mesma pessoa comeca bloqueado.
    ana2 = entrar(app, ATENDENTE)
    checar(ana2.post(f"/api/clientes/senhas/{sid}/revelar").status_code == 428,
           "desbloqueio nao passa para outra sessao")

    acessos = adm.get("/api/clientes/cofre/acessos").json()["itens"]
    resultados = {(a["usuario"], a["resultado"]) for a in acessos
                  if a["acao"] == "revelar"}
    checar(("ana", "ok") in resultados and ("bruno", "negado") in resultados,
           f"a trilha do cofre tem a revelacao e a negativa: {resultados}")
    checar(ana.get("/api/clientes/cofre/acessos").status_code == 403,
           "relatorio de acessos exige admin.historico")

    # Cinco senhas erradas bloqueiam o desbloqueio.
    for _ in range(5):
        ana2.post("/api/clientes/cofre/desbloquear", json={"senha": "x"})
    r = ana2.post("/api/clientes/cofre/desbloquear",
                  json={"senha": ATENDENTE["senha"]})
    checar(r.status_code == 429, f"tentativas demais: {r.status_code}")

    # Sem chave: o cadastro funciona, o cofre nao.
    chave = os.environ.pop(cf.VAR_CHAVE)
    checar(not ana.get("/api/clientes/catalogo").json()["cofre_disponivel"],
           "catalogo avisa que o cofre esta desligado")
    r = ana.post(f"/api/clientes/senhas/{sid}/revelar")
    checar(r.status_code == 503, f"sem chave: {r.status_code}")
    checar(ana.get(f"/api/clientes/{cid}").status_code == 200, "le sem chave")
    os.environ[cf.VAR_CHAVE] = chave

    # ------------------------------------------------------------------
    # Importacao da planilha

    xlsx = os.path.join(tempfile.mkdtemp(prefix="cli_xlsx_"), "senhas.xlsx")
    tp.planilha(xlsx)
    with open(xlsx, "rb") as arquivo:
        conteudo = arquivo.read()
    r = bruno.post("/api/clientes/importar/planilha",
                   files={"arquivo": ("senhas.xlsx", conteudo)})
    checar(r.status_code == 403, "consulta nao importa")
    r = ana.post("/api/clientes/importar/planilha",
                 files={"arquivo": ("senhas.xlsx", conteudo)})
    checar(r.status_code == 200, f"previa: {r.text[:300]}")
    previa = r.json()
    for senha in tp.SENHAS:
        checar(senha not in r.text, f"a previa nao pode trazer {senha!r}")
    checar(previa["contagens"]["incluir"] == 6, previa["contagens"])
    alfa = next(l for l in previa["linhas"] if l["linha"] == 3)
    checar(alfa["segredos"]["sefaz"] == "presente" and alfa["resultado"] == "novo",
           alfa)
    checar(varrer(tp.SENHAS) == [], "o upload nao pode ficar no disco")

    r = bruno.post(f"/api/clientes/importar/planilha/{previa['token']}/confirmar",
                   json={})
    checar(r.status_code == 403, "outro usuario nao confirma")
    r = ana.post(f"/api/clientes/importar/planilha/{previa['token']}/confirmar",
                 json={"decisoes": {"9": "ignorar"}})
    checar(r.status_code == 200, f"confirmar: {r.text[:300]}")
    checar(r.json()["criados"] == 5 and r.json()["ignorados"] == 3, r.json())
    r = ana.post(f"/api/clientes/importar/planilha/{previa['token']}/confirmar",
                 json={})
    checar(r.status_code == 404, "a previa so vale uma vez")

    # ------------------------------------------------------------------
    # Coleta e certificado por arquivo

    coleta = {"formato": "auditoria-jb/coleta-certificados", "versao": 1,
              "gerada_em": "2026-10-26T09:00:00", "pasta": "CERT",
              "itens": [{"pasta": "CLIENTE WEB", "arquivo": "c.pfx",
                         "hash_arquivo": "c" * 64, "titular_nome": "CLIENTE",
                         "titular_documento": doc, "emissor": "AC X",
                         "inicio": "2025-01-01", "fim": "2030-01-01",
                         "serie": "C0", "tipo": "A1", "info": {},
                         "problemas": [], "senha": "nunca-deveria-vir"}]}
    r = ana.post("/api/clientes/importar/coleta", files={
        "arquivo": ("coleta.json", json.dumps(coleta).encode())})
    checar(r.status_code == 200 and r.json()["contagens"] == {"casou": 1},
           r.text[:300])
    checar("nunca-deveria-vir" not in r.text, "senha na coleta e' descartada")
    r = ana.post(f"/api/clientes/importar/coleta/{r.json()['token']}/confirmar")
    checar(r.status_code == 200 and r.json()["novos"] == 1, r.text[:200])
    r = ana.post("/api/clientes/importar/coleta",
                 files={"arquivo": ("x.json", b"{nao e json")})
    checar(r.status_code == 422, f"coleta invalida: {r.status_code}")

    doc2 = cnpj("055555550001")
    outro = ana.post("/api/clientes", json={"nome": "Empresa Web",
                                            "documento": doc2}).json()["id"]
    pfx = pfx_sintetico(doc2)
    r = ana.post(f"/api/clientes/{outro}/certificados/arquivo",
                 files={"arquivo": ("e.pfx", pfx)}, data={"senha": "errada"})
    checar(r.status_code == 422 and r.json()["detail"]["code"]
           == "CERTIFICADO_SENHA_INCORRETA", r.text[:200])
    r = ana.post(f"/api/clientes/{outro}/certificados/arquivo",
                 files={"arquivo": ("e.pfx", pfx)},
                 data={"senha": SENHA_PFX, "local_guarda": "Pasta da rede"})
    checar(r.status_code == 200, f"certificado por arquivo: {r.text[:200]}")
    cert_id = r.json()["id"]

    # ------------------------------------------------------------------
    # Vencimentos e aviso

    venc = bruno.get("/api/clientes/vencimentos/lista").json()
    meu = next(v for v in venc["itens"] if v["id"] == cert_id)
    checar(meu["faixa"] == "ate_15" and meu["cliente_nome"] == "Empresa Web", meu)
    avisos = bruno.get("/api/clientes/avisos").json()
    checar(avisos["aviso"] >= 1, avisos)
    r = ana.post(f"/api/clientes/certificados/{cert_id}/renovacao",
                 json={"em_andamento": True})
    checar(r.status_code == 200, r.text[:200])
    checar(bruno.get("/api/clientes/avisos").json()["aviso"]
           == avisos["aviso"] - 1, "renovacao em andamento sai do aviso")
    r = ana.put("/api/clientes/vencimentos/faixas", json={"faixas": [7]})
    checar(r.status_code == 403, "configurar prazos exige permissao propria")
    r = adm.put("/api/clientes/vencimentos/faixas", json={"faixas": [30, 15, 20]})
    checar(r.status_code == 200 and r.json()["faixas"] == [15, 20, 30], r.text)

    # ------------------------------------------------------------------
    # Exportacao sem senha

    checar(ana.post("/api/clientes/exportar").status_code == 403,
           "exportar exige clientes.exportar")
    r = adm.post("/api/clientes/exportar")
    checar(r.status_code == 200 and r.content[:2] == b"PK", r.status_code)
    livro = load_workbook(io.BytesIO(r.content))
    texto = " ".join(str(c.value) for aba in livro.worksheets
                     for linha in aba.iter_rows() for c in linha if c.value)
    for senha in [SENHA_CLIENTE, *tp.SENHAS]:
        checar(senha not in texto, f"a exportacao nao pode trazer {senha!r}")
    checar("Senha Prefeitura" in texto, "mas diz quais senhas existem")

    # ------------------------------------------------------------------
    # Nenhuma senha em texto em arquivo nenhum da pasta de dados

    achados = varrer([SENHA_CLIENTE, SENHA_PFX, *tp.SENHAS])
    checar(achados == [], f"senha em texto na pasta de dados: {achados}")

    # ------------------------------------------------------------------
    # Trilha de uso: ok e negado, sem valor

    itens = eventos(adm)
    acoes = {(e["acao"], e["resultado"]) for e in itens}
    for esperado in (("cofre.revelar", "ok"), ("cofre.revelar", "negado"),
                     ("cofre.desbloquear", "ok"), ("cofre.gravar", "ok"),
                     ("clientes.importar", "ok"), ("clientes.cadastrar", "ok"),
                     ("clientes.exportar", "ok")):
        checar(esperado in acoes, f"{esperado} deveria estar na trilha")
    checar(SENHA_CLIENTE not in json.dumps(itens), "trilha de uso sem senha")

    print("OK - clientes web (401/403 com trilha, slugs fora do padrao, "
          "cadastro, cofre com desbloqueio por sessao e revelacao auditada, "
          "negado nas duas trilhas, cofre desligado sem chave, importacao com "
          "previa sem senha e de uso unico, coleta sem senha, certificado por "
          "arquivo, vencimentos e aviso, exportacao e pasta de dados sem "
          "senha) passou.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
