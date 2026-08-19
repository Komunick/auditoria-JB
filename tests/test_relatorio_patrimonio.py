"""Saidas do Patrimonio: relacao em Excel e termo em PDF (modulo puro).

Cobre colunas e valores da relacao, texto perigoso neutralizado, dinheiro como
celula numerica, metadados, falha sem arquivo parcial e a geracao do termo de
responsabilidade.
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "src"))

from openpyxl import load_workbook  # noqa: E402

from auditoria_fiscal.ferramentas import relatorio_patrimonio as rp  # noqa: E402


def checar(cond, msg):
    if not cond:
        print(f"FALHOU - {msg}")
        raise SystemExit(1)


def bem_exemplo(**ajustes) -> dict:
    base = {
        "id": 1, "etiqueta": "JBF-000001", "tipo": "Notebook",
        "descricao": "Notebook Dell Latitude 5440", "marca": "Dell",
        "modelo": "5440", "numero_serie": "SN-001", "conservacao": "novo",
        "situacao": "em_uso", "valor_aquisicao": 450000,
        "data_aquisicao": "2026-03-15", "nota_fiscal": "NF 1234",
        "responsavel_nome": "Maria Silva", "local_nome": "Sala Fiscal",
        "observacao": "",
    }
    base.update(ajustes)
    return base


def testar_relacao(pasta: str) -> None:
    destino = os.path.join(pasta, "relacao.xlsx")
    # Um bem com descricao MALICIOSA e outro sem valor informado.
    perigoso = bem_exemplo(id=2, etiqueta="JBF-000002",
                           descricao="=HYPERLINK(\"http://x\",\"clique\")",
                           responsavel_nome="-Fulano", valor_aquisicao=None,
                           situacao="disponivel", numero_serie="")
    rp.gerar_relacao(destino, bens=[bem_exemplo(), perigoso],
                     gerado_por="ana", filtros={"situacao": "em_uso"})

    livro = load_workbook(destino)
    checar(livro.sheetnames == ["Bens", "Metadados"], livro.sheetnames)
    aba = livro["Bens"]
    checar(aba.freeze_panes == "A2", "primeira linha congelada")
    checar(aba.auto_filter.ref, "autofiltro no cabecalho")

    cabecalhos = [c.value for c in aba[1]]
    checar(cabecalhos[0] == "Etiqueta" and "Responsável" in cabecalhos,
           cabecalhos)

    checar(aba["A2"].value == "JBF-000001", aba["A2"].value)
    checar(aba["H2"].value == "Em uso", aba["H2"].value)
    checar(aba["G2"].value == "Novo", aba["G2"].value)

    # Dinheiro e' celula NUMERICA formatada, nao texto.
    checar(float(aba["K2"].value) == 4500.00, aba["K2"].value)
    checar("R$" in (aba["K2"].number_format or ""), aba["K2"].number_format)
    # Valor ausente fica VAZIO, nunca R$ 0,00.
    checar(aba["K3"].value is None, f"valor ausente: {aba['K3'].value!r}")

    # Texto perigoso neutralizado, conteudo preservado.
    checar(aba["C3"].value.startswith("'="), f"descricao: {aba['C3'].value!r}")
    checar(aba["I3"].value.startswith("'-"), f"responsavel: {aba['I3'].value!r}")
    for linha in aba.iter_rows():
        for celula in linha:
            if isinstance(celula.value, str):
                checar(not celula.value.startswith("="),
                       f"{celula.coordinate} virou formula: {celula.value!r}")

    meta = {l[0]: l[1] for l in livro["Metadados"].iter_rows(
        min_row=2, values_only=True)}
    checar(meta["Gerado por"] == "ana", meta)
    checar(meta["Total de bens"] == 2, meta)
    checar("situacao=em_uso" in meta["Filtros aplicados"], meta)
    livro.close()

    # Soma so o que tem valor.
    livro = load_workbook(destino)
    aba_meta = livro["Metadados"]
    total = [c for linha in aba_meta.iter_rows() for c in linha
             if isinstance(c.value, (int, float)) and c.column == 2]
    checar(any(float(c.value) == 4500.00 for c in total),
           f"valor total deveria somar so o informado: "
           f"{[c.value for c in total]}")
    livro.close()


def testar_falha_sem_parcial(pasta: str) -> None:
    destino = os.path.join(pasta, "quebrado.xlsx")
    ruim = bem_exemplo(valor_aquisicao=object())
    try:
        rp.gerar_relacao(destino, bens=[ruim], gerado_por="x", filtros={})
        checar(False, "deveria ter falhado")
    except Exception:
        pass
    checar(not os.path.exists(destino), "falha nao pode deixar arquivo")
    checar(not os.path.exists(destino + ".parcial"), "nem o temporario")


def testar_termo(pasta: str) -> None:
    destino = os.path.join(pasta, "termo.pdf")
    rp.gerar_termo(
        destino,
        bem=bem_exemplo(descricao="Notebook Dell — Latitude “5440”"),
        responsavel={"nome": "Maria Silva", "setor": "Fiscal",
                     "desde": "2026-03-20T12:00:00+00:00"},
        emitido_por="ana")
    checar(os.path.isfile(destino), "o termo deveria existir")
    with open(destino, "rb") as arq:
        conteudo = arq.read()
    checar(conteudo[:4] == b"%PDF", "deveria ser um PDF")
    checar(len(conteudo) > 800, f"PDF pequeno demais: {len(conteudo)} bytes")
    checar(not os.path.exists(destino + ".parcial"), "sem temporario sobrando")

    # Acentos e aspas tipograficas nao podem quebrar a geracao: a fonte core
    # do PDF e' latin-1, e um glifo fora da tabela derrubaria o arquivo todo.
    outro = os.path.join(pasta, "termo_acentos.pdf")
    rp.gerar_termo(
        outro,
        bem=bem_exemplo(descricao="Cadeira ergonômica — ação/coração ✓"),
        responsavel={"nome": "João Conceição", "setor": "Contábil",
                     "desde": "2026-01-05T09:00:00+00:00"},
        emitido_por="ana")
    checar(os.path.getsize(outro) > 800, "termo com acentos deveria gerar")


def main() -> int:
    pasta = tempfile.mkdtemp(prefix="patrimonio_relatorio_")
    try:
        testar_relacao(pasta)
        testar_falha_sem_parcial(pasta)
        testar_termo(pasta)
    finally:
        shutil.rmtree(pasta, ignore_errors=True)
    print("OK - saidas do patrimonio (relacao com valores numericos, valor "
          "ausente vazio, texto perigoso neutralizado, metadados, falha sem "
          "arquivo parcial e termo em PDF tolerante a acentos) passou.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
