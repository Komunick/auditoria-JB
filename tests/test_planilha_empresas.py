"""Spec 017: base Fiscal sintetica, sem cadastro nem dados reais."""
import hashlib
import json
import random
import re
import sys
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from auditoria_fiscal.core import planilha_empresas as pe
from auditoria_fiscal.core.clientes import ErroClientes
from auditoria_fiscal.core.documentos_br import dv_cnpj

CABECALHO = [None, "RESPONSAVEL", "EMPRESAS", "RESPONSAVEL", "ZERADAS", "REGIME DA EMPRESA"]
MASCARA = '00"."000"."000"/"0000"-"00'


def cnpj(seed, *, zero=False):
    rng = random.Random(seed)
    base = ("0" if zero else str(rng.randint(1, 9))) + str(rng.randint(1, 9))
    base += "".join(str(rng.randrange(10)) for _ in range(10))
    return base + dv_cnpj(base)


def livro(linhas=None, *, titulo="Carteira sintetica"):
    wb = Workbook()
    ws = wb.active
    ws.title = titulo
    ws.append(["Base de teste"])
    ws.append(CABECALHO)
    for linha in linhas or []:
        ws.append(linha)
    return wb, ws


def linha(documento=None, *, nome="Empresa sintetica", regime="Simples Nacional", responsavel="Pessoa sintetica", contato="Contato sintetico"):
    return [None, responsavel, nome, contato, documento, regime]


def ler(wb, pasta, nome="base.xlsx"):
    caminho = Path(pasta) / nome
    wb.save(caminho)
    wb.close()
    return pe.ler_planilha_empresas(caminho).resumo()


def recusa(caminho, trecho=None, **kwargs):
    try:
        pe.ler_planilha_empresas(caminho, **kwargs)
    except ErroClientes as exc:
        assert exc.codigo == "PLANILHA_INVALIDA"
        if trecho:
            assert trecho.casefold() in exc.mensagem.casefold(), exc.mensagem
        return
    raise AssertionError("Planilha invalida foi aceita")


def reescrever(caminho, mudar):
    with zipfile.ZipFile(caminho) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    mudar(partes)
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, conteudo in partes.items():
            z.writestr(nome, conteudo)


def test_layouts_origem_duplicidades_e_ocultas(pasta):
    documento = cnpj(1701)
    wb, ws = livro([linha(documento), linha(cnpj(1702), regime=" presumído ")])
    ws.cell(2, 2).value = " responsÁvel "
    ws.cell(2, 6).value = "regime   da empresa"
    ws.row_dimensions[4].hidden = True
    ws.merge_cells("D3:D5")  # A linha 5 vazia nao vira empresa por propagacao.
    zeradas = wb.create_sheet(" zéradas ")
    zeradas.sheet_state = "hidden"
    zeradas.append(["Somente sinteticos"])
    zeradas.append([None, " empresa ", "CNPJ"])
    zeradas.append([None, "Nome divergente preservado", documento])
    wb.create_sheet("Vazia")
    resumo = ler(wb, pasta)
    assert resumo["formato"] == "base_fiscal_v1"
    assert resumo["hash_arquivo"] == hashlib.sha256((Path(pasta) / "base.xlsx").read_bytes()).hexdigest()
    assert resumo["contagens"] == dict(linhas=3, documentos_validos=3, cnpjs_distintos=2,
        documentos_pendentes=0, grupos_duplicados=1, linhas_duplicadas=2, sem_movimento=1,
        sem_responsavel=1, sem_regime=1, com_pendencias=2)
    a, b, c = resumo["linhas"]
    assert [r["id"] for r in (a, b, c)] == ["1:3", "1:4", "2:3"]
    assert a["sem_movimento"] is None and b["sem_movimento"] is None and c["sem_movimento"] is True
    assert a["documento"] == c["documento"] and a["nome"] != c["nome"]
    assert a["duplicado_em"] == [c["id"]] and c["duplicado_em"] == [a["id"]]
    assert a["duplicado"] and c["duplicado"] and not b["duplicado"]
    assert b["regime"] == "Lucro Presumido" and b["regime_informado"] == " presumído "
    assert b["oculta"] and not a["oculta"] and c["aba_oculta"]
    assert b["contato_nome"] == "Contato sintetico"
    assert b["origem"]["contato_nome"] == dict(celula="D4", celula_origem="D3",
        valor_informado="Contato sintetico", formato="General", intervalo_mesclado="D3:D5", propagado=True, formula=False)
    assert a["origem"]["documento"]["celula"] == "E3"
    assert a["origem"]["responsavel_fiscal"]["celula"] == "B3"
    assert "responsavel_fiscal" not in c["origem"] and len(resumo["avisos"]) == 1
    assert resumo["abas"] == [dict(nome="Carteira sintetica", tipo="principal", linha_cabecalho=2, total=2, oculta=False),
                             dict(nome=" zéradas ", tipo="zeradas", linha_cabecalho=2, total=1, oculta=True)]
    json.dumps(resumo, allow_nan=False)


def test_documentos_estritos_e_mascara(pasta):
    completo, zero = cnpj(1703), cnpj(1704, zero=True)
    valores = [completo, int(completo), int(zero), str(int(zero)), int(zero), True, float(int(zero)) + .5,
               int("9" + completo), "", "0" * 14, "abc" + completo, zero, int(zero), int(zero), -int(completo)]
    wb, ws = livro([linha(v) for v in valores])
    ws["E5"].number_format = MASCARA
    ws["E7"].number_format = "General"
    ws["E15"].number_format = "00000000000000"
    ws["E16"].number_format = "00.000.000/0000-00"
    dados = ler(wb, pasta)["linhas"]
    assert [d["documento_status"] for d in dados] == ["valido", "valido", "zero_inicial_recuperado", "invalido", "invalido",
        "invalido", "invalido", "invalido", "ausente", "invalido", "invalido", "valido", "invalido", "zero_inicial_recuperado", "invalido"]
    assert dados[2]["documento"] == zero and dados[2]["documento_informado"] == str(int(zero))
    assert dados[2]["origem"]["documento"]["formato"] == MASCARA
    assert all(d["documento"] == "" for d in dados if d["documento_status"] in {"invalido", "ausente"})
    assert not dados[2]["pendencias"] or all("duplic" in p.casefold() for p in dados[2]["pendencias"])


def test_formulas_sem_cache_nem_texto(pasta):
    wb, ws = livro([linha("=1+SECRET_FORMULA()", nome="=SECRET_NAME()", regime="=SECRET_REGIME()", responsavel="=SECRET_RESP()", contato="=SECRET_CONTACT()"),
                   linha(cnpj(1705), contato=None)])
    ws.merge_cells("D3:D4")
    caminho = Path(pasta) / "formulas.xlsx"
    wb.save(caminho)
    wb.close()
    reescrever(caminho, lambda partes: partes.update({"xl/worksheets/sheet1.xml": partes["xl/worksheets/sheet1.xml"].replace(b"<v></v>", b"<v>123456789</v>")}))
    resumo = pe.ler_planilha_empresas(caminho).resumo()
    a, b = resumo["linhas"]
    assert a["documento_status"] == "formula"
    assert a["documento"] == a["nome"] == a["regime"] == a["responsavel_fiscal_informado"] == a["contato_nome"] == ""
    assert a["documento_informado"] == a["regime_informado"] == "[fórmula não lida]"
    assert b["contato_nome"] == "" and b["origem"]["contato_nome"]["formula"]
    assert all(o["valor_informado"] == "[fórmula não lida]" for o in a["origem"].values())
    assert "SECRET" not in json.dumps(resumo) and "123456789" not in json.dumps(resumo)


def test_pendencias_sem_inferir(pasta):
    wb, ws = livro([linha(None, nome=None, regime="MEI", responsavel=""), linha(cnpj(1706), regime="lucro real")])
    resumo = ler(wb, pasta)
    a, b = resumo["linhas"]
    assert a["nome"] == a["documento"] == a["regime"] == a["responsavel_fiscal_informado"] == ""
    assert a["regime_informado"] == "MEI" and len(a["pendencias"]) == 4
    assert b["regime"] == "Lucro Real" and not b["pendencias"]
    assert resumo["contagens"]["com_pendencias"] == 1


def test_layouts_recusados_e_vazio(pasta):
    for caso in ("desconhecida", "coluna_extra", "coluna_extra_zero", "ambiguo", "zeradas_falsa", "principal_em_zeradas", "cabecalho_tardio"):
        wb, ws = livro([linha(cnpj(1710))])
        if caso == "desconhecida":
            wb.create_sheet("Desconhecida").append(["conteudo"])
        elif caso.startswith("coluna_extra"):
            ws["G3"] = 0 if caso.endswith("zero") else "campo nao mapeado"
        elif caso == "ambiguo":
            ws.append(CABECALHO)
        elif caso == "zeradas_falsa":
            ws.title = "Carteira"
            for c in ws[2]:
                c.value = None
            ws["B2"], ws["C2"] = "EMPRESA", "CNPJ"
        elif caso == "principal_em_zeradas":
            ws.title = "ZERADAS"
        else:
            ws.insert_rows(1, 10)
        caminho = Path(pasta) / (caso + ".xlsx")
        wb.save(caminho)
        wb.close()
        recusa(caminho)
    wb, _ = livro()
    assert ler(wb, pasta)["contagens"]["linhas"] == 0
    wb = Workbook()
    caminho = Path(pasta) / "so-vazia.xlsx"
    wb.save(caminho)
    wb.close()
    recusa(caminho)


def test_mesclagens_invalidas(pasta):
    for intervalo in ("B3:B4", "C3:C4", "E3:E4", "F3:F4", "D2:D3", "C3:D3", "D1:D2", "A1:F1"):
        wb, ws = livro([linha(cnpj(1711)), linha(cnpj(1712))])
        ws.merge_cells(intervalo)
        caminho = Path(pasta) / "mesclagem.xlsx"
        wb.save(caminho)
        wb.close()
        recusa(caminho)


def test_limites_antes_openpyxl_e_dimensao_falsa(pasta):
    wb, ws = livro([linha(cnpj(1713)), linha(cnpj(1714))])
    caminho = Path(pasta) / "limites.xlsx"
    wb.save(caminho)
    wb.close()
    original = caminho.read_bytes()
    alteracoes = [lambda b: b.replace(b'r="4"', b'r="10021"'),
                 lambda b: b.replace(b'r="E4"', b'r="AG4"'),
                 lambda b: b.replace(b'</worksheet>', b'<mergeCells count="1"><mergeCell ref="D3:D1048576"/></mergeCells></worksheet>'),
                 lambda b: b.replace(b'</worksheet>', b'<cols><col min="1" max="16384"/></cols></worksheet>')]
    for alterar in alteracoes:
        caminho.write_bytes(original)
        reescrever(caminho, lambda p: p.update({"xl/worksheets/sheet1.xml": alterar(p["xl/worksheets/sheet1.xml"])}))
        with patch.object(pe, "load_workbook", side_effect=AssertionError("Parser de alto nivel abriu antes dos limites")):
            recusa(caminho)
    caminho.write_bytes(original)
    reescrever(caminho, lambda p: p.update({"xl/worksheets/sheet1.xml": re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', p["xl/worksheets/sheet1.xml"])}))
    assert len(pe.ler_planilha_empresas(caminho).resumo()["linhas"]) == 2
    wb, _ = livro()
    for i in range(20):
        wb.create_sheet("Vazia" + str(i))
    wb.save(caminho)
    wb.close()
    with patch.object(pe, "load_workbook", side_effect=AssertionError("Abas acima do limite")):
        recusa(caminho)
    wb, ws = livro([linha(cnpj(1715))])
    ws["E10020"] = cnpj(1716)
    assert len(ler(wb, pasta)["linhas"]) == 2


def test_limite_util_global(pasta):
    wb, _ = livro([linha(cnpj(1717)), linha(cnpj(1718))])
    ws = wb.create_sheet("ZERADAS")
    ws.append([None, "EMPRESA", "CNPJ"])
    ws.append([None, "Empresa sintetica Z", cnpj(1719)])
    caminho = Path(pasta) / "limite-util.xlsx"
    wb.save(caminho)
    wb.close()
    with patch.object(pe, "MAX_LINHAS", 2):
        recusa(caminho)


def test_pacote_xml_e_tamanho(pasta):
    wb, _ = livro([linha(cnpj(1720))])
    caminho = Path(pasta) / "pacote.xlsx"
    wb.save(caminho)
    wb.close()
    recusa(caminho, max_bytes=1)
    original = caminho.read_bytes()
    for entrada in ("xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml", "../fora.xml"):
        caminho.write_bytes(original)
        with zipfile.ZipFile(caminho, "a") as z:
            z.writestr(entrada, b"invalido")
        recusa(caminho)
    caminho.write_bytes(original)
    reescrever(caminho, lambda p: p.update({"xl/worksheets/sheet1.xml": b'<!DOCTYPE worksheet [<!ENTITY segredo SYSTEM "file:///arquivo-inexistente">]>' + p["xl/worksheets/sheet1.xml"]}))
    with patch.object(pe, "load_workbook", side_effect=AssertionError("DTD chegou ao openpyxl")):
        recusa(caminho)
    caminho.write_bytes(b"nao e xlsx")
    recusa(caminho)
    caminho_xls = Path(pasta) / "arquivo.xls"
    caminho_xls.write_bytes(original)
    recusa(caminho_xls)


def test_mascara_escapada_e_numero_muito_longo(pasta):
    zero, completo = cnpj(1721, zero=True), cnpj(1722)
    wb, ws = livro([linha(int(zero)), linha(int(completo)), linha(42)])
    ws["E3"].number_format = '00"."000"."000"/"0000\\-00'
    ws["E4"].number_format = "0"
    caminho = Path(pasta) / "numero-longo.xlsx"
    wb.save(caminho)
    wb.close()
    reescrever(caminho, lambda p: p.update({"xl/worksheets/sheet1.xml": p["xl/worksheets/sheet1.xml"].replace(b'<v>42</v>', ('<v>' + str(10 ** 310) + '</v>').encode())}))
    dados = pe.ler_planilha_empresas(caminho).resumo()["linhas"]
    assert dados[0]["documento_status"] == "zero_inicial_recuperado" and dados[0]["documento"] == zero
    assert dados[1]["documento_status"] == "valido" and dados[1]["documento"] == completo
    assert dados[2]["documento_status"] == "invalido" and not dados[2]["documento"]


def test_duplicidades_limitadas_antes_da_saida(pasta):
    documento = cnpj(1723)
    wb, _ = livro([linha(documento), linha(documento), linha(documento)])
    caminho = Path(pasta) / "duplicidades.xlsx"
    wb.save(caminho)
    wb.close()
    with patch.object(pe, "MAX_REFERENCIAS_DUPLICADAS", 5):
        recusa(caminho, "duplicidades")
    with patch.object(pe, "MAX_REFERENCIAS_DUPLICADAS", 6):
        resumo = pe.ler_planilha_empresas(caminho).resumo()
        assert resumo["contagens"]["linhas_duplicadas"] == 3
        assert sum(len(l["duplicado_em"]) for l in resumo["linhas"]) == 6


def test_xml_malformado_nao_vira_erro_interno(pasta):
    wb, _ = livro([linha(cnpj(1724))])
    caminho = Path(pasta) / "estilo-invalido.xlsx"
    wb.save(caminho)
    wb.close()
    reescrever(caminho, lambda p: p.update({"xl/worksheets/sheet1.xml": p["xl/worksheets/sheet1.xml"].replace(b'<c r="E3"', b'<c s="99999" r="E3"')}))
    recusa(caminho)


def test_mesclagem_nao_descarta_celula_preenchida(pasta):
    wb, _ = livro([linha(cnpj(1725)), linha(cnpj(1726), contato="Outro contato preservado")])
    caminho = Path(pasta) / "mesclagem-conflitante.xlsx"
    wb.save(caminho)
    wb.close()
    reescrever(caminho, lambda p: p.update({"xl/worksheets/sheet1.xml": p["xl/worksheets/sheet1.xml"].replace(b'</worksheet>', b'<mergeCells count="1"><mergeCell ref="D3:D4"/></mergeCells></worksheet>')}))
    recusa(caminho)


def test_pontuacao_de_cnpj_sem_aceitar_prefixos(pasta):
    documento = cnpj(1727)
    formatos = [documento[:2] + "." + documento[2:5] + "." + documento[5:8] + "." + documento[8:12] + "-" + documento[12:],
                documento[:12] + "-" + documento[12:], documento[:2] + "." + documento[2:5] + "." + documento[5:8] + "/" + documento[8:],
                " ".join(documento), "`" + documento, "´" + documento, "x" + documento, "+" + documento,
                " ".join(documento[:-1]), documento.replace(documento[0], chr(0xFF10 + int(documento[0])), 1)]
    wb, _ = livro([linha(v) for v in formatos])
    dados = ler(wb, pasta)["linhas"]
    assert all(d["documento"] == documento for d in dados[:4])
    assert all(d["documento_status"] == "invalido" and not d["documento"] for d in dados[4:])
    assert [d["documento_informado"] for d in dados] == formatos


def test_preflight_segue_relacao_e_recusa_linha_fora_da_estrutura(pasta):
    wb, ws = livro([linha(cnpj(1728))])
    ws.delete_rows(1)  # Cabecalho na primeira linha.
    caminho = Path(pasta) / "referencia.xlsx"
    wb.save(caminho)
    wb.close()
    def renomear(partes):
        partes["xl/worksheets/folha.dados"] = partes.pop("xl/worksheets/sheet1.xml")
        for nome in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
            partes[nome] = partes[nome].replace(b"sheet1.xml", b"folha.dados")
    reescrever(caminho, renomear)
    assert len(pe.ler_planilha_empresas(caminho).resumo()["linhas"]) == 1, "A extensao interna nao pode pular o preflight e truncar dados"
    wb, _ = livro([linha(cnpj(1729))])
    wb.save(caminho)
    wb.close()
    reescrever(caminho, lambda p: p.update({"xl/worksheets/sheet1.xml": p["xl/worksheets/sheet1.xml"].replace(b'</worksheet>', b'<row r="10021"><c r="B10021" t="inlineStr"><is><t>fora</t></is></c></row></worksheet>')}))
    with patch.object(pe, "load_workbook", side_effect=AssertionError("Linha fora da estrutura chegou ao openpyxl")):
        recusa(caminho)


def main():
    testes = [valor for nome, valor in globals().items() if nome.startswith("test_")]
    with tempfile.TemporaryDirectory(prefix="base_fiscal_sintetica_") as pasta:
        for teste in testes:
            teste(pasta)
    print(f"OK - base Fiscal: {len(testes)} cenarios de layout, origem, documentos, limites e pacote seguro")


if __name__ == "__main__":
    main()
