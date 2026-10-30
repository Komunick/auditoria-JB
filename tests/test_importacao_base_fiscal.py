"""Carga Fiscal sintética: revisão, transação, concorrência e recibo imutável."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
from threading import Barrier
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import Workbook, load_workbook
from auditoria_fiscal.core import correcoes_base_fiscal as cbf
from auditoria_fiscal.core.clientes import ErroClientes
from auditoria_fiscal.core.documentos_br import dv_cnpj
from auditoria_fiscal.core.planilha_empresas import ler_planilha_empresas
from auditoria_fiscal.ferramentas.clientes_store import ClientesStore
from auditoria_fiscal.ferramentas.previa_cadastro_fiscal import consultar_cadastro
from auditoria_fiscal.ferramentas.relatorio_base_fiscal import gerar_relatorio_base_fiscal
from auditoria_fiscal.ferramentas.importacao_base_fiscal import (
    assinatura_previa, enriquecer_origens, importar_base_fiscal, validar_apta,
)


DOCS = [b + dv_cnpj(b) for b in ("019876540001", "529876540001", "639876540001", "749876540001")]
AUTOR = {"usuario_id": 7, "usuario_login": "importador_sintetico"}


def preparar(pasta):
    fonte = pasta / "fonte.xlsx"
    livro = Workbook()
    aba = livro.active
    aba.title = "Fiscal sintético"
    aba.append([None, "RESPONSAVEL", "EMPRESAS", "RESPONSAVEL", "CNPJ", "REGIME DA EMPRESA"])
    for i, nome in enumerate(("Empresa nova A", "Empresa nova B", "Empresa existente")):
        aba.append([None, "Pessoa sintética", nome, f"Contato {i}", DOCS[i], "Lucro Real"])
    aba.append([None, None, "Ocorrência desconsiderada", None, "pendente", None])
    zeradas = livro.create_sheet("ZERADAS")
    zeradas.append([None, "EMPRESA", "CNPJ"])
    zeradas.append([None, "Empresa zerada", DOCS[3]])
    livro.save(fonte)
    livro.close()
    leitura = ler_planilha_empresas(fonte)
    relatorio = load_workbook(BytesIO(gerar_relatorio_base_fiscal(leitura, nome_arquivo="fonte.xlsx")))
    relatorio["Correções"]["L5"] = "ignorar"
    relatorio["Correções"]["M5"] = "Decisão sintética de não incluir esta ocorrência"
    relatorio["Correções"]["K2"] = "Contato revisado"
    revisao = pasta / "correcoes.xlsx"
    relatorio.save(revisao)
    relatorio.close()
    base = enriquecer_origens(cbf.preparar_previa(leitura, revisao), leitura)
    return base, leitura


def banco(pasta, nome="clientes.db"):
    store = ClientesStore(str(pasta / nome))
    store.criar_cliente(dict(nome="Empresa existente", documento=DOCS[2], contato_nome="Contato 2"), **AUTOR)
    store.criar_cliente(dict(nome="Empresa zerada", documento=DOCS[3], contato_nome="Contato preservado"), **AUTOR)
    return store


def previa(store, base):
    snap = consultar_cadastro(store.db_path, [l["efetivo"]["documento"] for l in base["linhas"]])
    return cbf.comparar_cadastro(base, snap["clientes"])


def importar(store, base, assinatura):
    return importar_base_fiscal(store, base, assinatura, nome_arquivo="fonte.xlsx",
                               nome_correcoes="correcoes.xlsx", **AUTOR)


def codigo(esperado, acao):
    try:
        acao()
    except ErroClientes as exc:
        assert exc.codigo == esperado, (exc.codigo, exc.mensagem)
    else:
        raise AssertionError(f"Esperava recusa {esperado}")


def test_assinatura_e_bloqueios(pasta, base, leitura):
    assert base["linhas"][0]["origem"]["documento"]["celula"] == "E2"
    assert base["linhas"][0]["origem"]["nome"]["valor_informado"] == "Empresa nova A"
    assinatura = assinatura_previa(base)
    assert len(assinatura) == 64
    r = deepcopy(base)
    r.update(cadastro_consultado_em="2099-01-01", nome_arquivo="outro.xlsx", cadastro_disponivel=False)
    assert assinatura_previa(r) == assinatura
    r["linhas"][0]["efetivo"]["nome"] = "Proposta diferente"
    assert assinatura_previa(r) != assinatura
    assert base["linhas"][0]["efetivo"]["nome"] == "Empresa nova A"
    assert enriquecer_origens(cbf.preparar_previa(leitura, pasta / "correcoes.xlsx"), leitura) == base
    validar_apta(base)
    for estado in ("pendente", "conflito"):
        r = deepcopy(base)
        r["linhas"][0]["status"] = estado
        r["contagens"] = {"incluir": 999, "pendente": 0, "conflito": 0}
        codigo("BASE_FISCAL_BLOQUEADA", lambda: validar_apta(r))
    r = deepcopy(base)
    for linha in r["linhas"]:
        linha["status"] = "existente" if linha["status"] != "ignorar" else "ignorar"
    codigo("BASE_FISCAL_SEM_INCLUSAO", lambda: validar_apta(r))
    store = banco(pasta, "bloqueios.db")
    try:
        r = deepcopy(base)
        r["linhas"][0]["status"] = "pendente"
        antes = list(store._conn.iterdump())
        codigo("BASE_FISCAL_BLOQUEADA", lambda: importar(store, r, assinatura_previa(previa(store, r))))
        assert list(store._conn.iterdump()) == antes
    finally:
        store.fechar()


def test_carga_recibo_e_proveniencia(pasta, base):
    store = banco(pasta)
    try:
        store._conn.execute("PRAGMA user_version=3")
        p = previa(store, base)
        assert [l["status"] for l in p["linhas"]] == ["incluir", "incluir", "existente", "ignorar", "existente"]
        existentes = [dict(l) for l in store._conn.execute("SELECT * FROM cliente ORDER BY id")]
        copia = deepcopy(base)
        r = importar(store, base, assinatura_previa(p))
        assert r["formato"] == "base_fiscal_resultado_v1" and r["repetida"] is False
        assert r["contagens"] == {"incluidas": 2, "existentes": 2, "ignoradas": 1}
        assert r["hash_arquivo"] == base["hash_arquivo"] and r["hash_correcoes"] == base["hash_correcoes"]
        assert r["importacao_id"] > 0 and r["criado_em"]
        assert base == copia
        assert existentes == [dict(l) for l in store._conn.execute("SELECT * FROM cliente WHERE id<=2 ORDER BY id")]
        novos = [dict(l) for l in store._conn.execute("SELECT * FROM cliente WHERE id>2 ORDER BY id")]
        assert len(novos) == 2 and novos[0]["contato_nome"] == "Contato revisado"
        assert all(l["situacao"] == "ativo" for l in novos)
        origem = json.loads(novos[0]["origem"])
        assert base["hash_arquivo"] in json.dumps(origem) and base["hash_correcoes"] in json.dumps(origem)
        assert "E2" in json.dumps(origem) and "Correções!K2" in json.dumps(origem, ensure_ascii=False)
        assert store.contar("segredo") == store.contar("certificado") == store.contar("licenca") == 0
        tabela = store._conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='importacao_base_fiscal'").fetchone()
        assert tabela is not None
        recibo = dict(store._conn.execute("SELECT * FROM importacao_base_fiscal").fetchone())
        conteudo = json.dumps(recibo, ensure_ascii=False)
        for valor in ("Ocorrência desconsiderada", "Decisão sintética", "Empresa existente", "Contato revisado", "Pessoa sintética"):
            assert valor in conteudo, valor
        eventos = [dict(e) for e in store._conn.execute("SELECT * FROM alteracao WHERE entidade='importacao_base_fiscal'")]
        assert len(eventos) == 1 and eventos[0]["usuario_id"] == 7
        antes = list(store._conn.iterdump())
        r2 = importar(store, base, assinatura_previa(p))
        assert r2 == {**r, "repetida": True} and list(store._conn.iterdump()) == antes
        store._conn.execute("UPDATE cliente SET nome='Alteração posterior',revision=revision+1 WHERE id=3")
        antes = list(store._conn.iterdump())
        assert importar(store, base, assinatura_previa(p)) == r2
        assert list(store._conn.iterdump()) == antes
        for sql in ("UPDATE importacao_base_fiscal SET hash_arquivo='x'", "DELETE FROM importacao_base_fiscal"):
            try:
                store._conn.execute(sql)
            except sqlite3.DatabaseError:
                pass
            else:
                raise AssertionError("Recibo precisa ser imutável")
    finally:
        store.fechar()
    aberto = ClientesStore(str(pasta / "clientes.db"))
    try:
        assert aberto._conn.execute("PRAGMA user_version").fetchone()[0] == 3
        assert aberto._conn.execute("SELECT count(*) FROM importacao_base_fiscal").fetchone()[0] == 1
    finally:
        aberto.fechar()


def test_snapshot_desatualizado(pasta, base):
    mudancas = [
        lambda s: s._conn.execute("UPDATE cliente SET revision=revision+1 WHERE documento=?", (DOCS[2],)),
        lambda s: s._conn.execute("UPDATE cliente SET situacao='inativo',revision=revision+1 WHERE documento=?", (DOCS[2],)),
        lambda s: s.criar_cliente(dict(nome="Empresa nova A", documento=DOCS[0], contato_nome="Contato revisado"), **AUTOR),
        lambda s: s._conn.execute("UPDATE cliente SET contato_nome='Outro contato',revision=revision+1 WHERE documento=?", (DOCS[3],)),
    ]
    for i, mudar in enumerate(mudancas):
        store = banco(pasta, f"snapshot-{i}.db")
        try:
            assinatura = assinatura_previa(previa(store, base))
            mudar(store)
            antes = list(store._conn.iterdump())
            codigo("BASE_FISCAL_DESATUALIZADA", lambda: importar(store, base, assinatura))
            assert list(store._conn.iterdump()) == antes
        finally:
            store.fechar()


def test_rollback(pasta, base):
    for tipo in ("segundo_cliente", "auditoria", "recibo"):
        store = banco(pasta, f"rollback-{tipo}.db")
        try:
            assinatura = assinatura_previa(previa(store, base))
            original = store._inserir_cliente
            chamadas = []
            def inserir(*args, **kwargs):
                chamadas.append(True)
                if len(chamadas) == 2:
                    raise RuntimeError("Falha sintética no segundo cliente")
                return original(*args, **kwargs)
            if tipo != "segundo_cliente":
                tabela = "alteracao" if tipo == "auditoria" else "importacao_base_fiscal"
                condicao = " WHEN NEW.entidade='importacao_base_fiscal'" if tipo == "auditoria" else ""
                store._conn.execute(f"CREATE TRIGGER falha_sintetica BEFORE INSERT ON {tabela}{condicao} BEGIN SELECT RAISE(ABORT,'falha sintética'); END")
            antes = list(store._conn.iterdump())
            try:
                if tipo == "segundo_cliente":
                    with patch.object(store, "_inserir_cliente", side_effect=inserir):
                        importar(store, base, assinatura)
                else:
                    importar(store, base, assinatura)
            except (RuntimeError, sqlite3.DatabaseError):
                pass
            else:
                raise AssertionError("A falha injetada deveria abortar")
            assert list(store._conn.iterdump()) == antes, tipo
            assert store._conn.in_transaction is False
        finally:
            store.fechar()


def test_concorrencia(pasta, base):
    store = banco(pasta, "concorrencia.db")
    assinatura = assinatura_previa(previa(store, base))
    caminho = store.db_path
    store.fechar()
    barreira = Barrier(2)
    def executar():
        atual = ClientesStore(caminho)
        try:
            barreira.wait(timeout=10)
            return importar(atual, base, assinatura)
        finally:
            atual.fechar()
    with ThreadPoolExecutor(max_workers=2) as executor:
        resultados = list(executor.map(lambda _: executar(), range(2)))
    assert sorted(r["repetida"] for r in resultados) == [False, True]
    assert resultados[0]["importacao_id"] == resultados[1]["importacao_id"]
    conn = sqlite3.connect(caminho)
    try:
        assert conn.execute("SELECT count(*) FROM cliente").fetchone()[0] == 4
        assert conn.execute("SELECT count(*) FROM importacao_base_fiscal").fetchone()[0] == 1
        assert conn.execute("SELECT count(*) FROM alteracao WHERE entidade='importacao_base_fiscal'").fetchone()[0] == 1
        assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert conn.execute("PRAGMA foreign_key_check").fetchall() == []
    finally:
        conn.close()


def test_revalidacao_sob_trava(pasta, base):
    store = banco(pasta, "revalidacao.db")
    try:
        assinatura = assinatura_previa(previa(store, base))
        chamadas = []
        def recusada():
            chamadas.append(store._conn.in_transaction)
            raise RuntimeError("Autoridade sintética revogada antes da gravação")
        def confirmar_revalidando():
            return importar_base_fiscal(store, base, assinatura, nome_arquivo="fonte.xlsx",
                                         nome_correcoes="correcoes.xlsx", revalidar=recusada, **AUTOR)
        antes = list(store._conn.iterdump())
        try:
            confirmar_revalidando()
        except RuntimeError:
            pass
        else:
            raise AssertionError("Revalidação recusada não pode confirmar")
        assert chamadas == [True], "Revalidar só depois de obter a trava de escrita"
        assert list(store._conn.iterdump()) == antes
        importar(store, base, assinatura)
        antes = list(store._conn.iterdump())
        try:
            confirmar_revalidando()
        except RuntimeError:
            pass
        else:
            raise AssertionError("Consultar recibo também exige autoridade atual")
        assert chamadas == [True, True]
        assert list(store._conn.iterdump()) == antes
    finally:
        store.fechar()


def test_constructor_fecha_conexao(pasta):
    original = sqlite3.connect
    conexoes = []
    def conectar(*args, **kwargs):
        conn = original(*args, **kwargs)
        conexoes.append(conn)
        return conn
    with patch("auditoria_fiscal.ferramentas.clientes_store.sqlite3.connect", side_effect=conectar):
        with patch.object(ClientesStore, "_migrar", side_effect=sqlite3.OperationalError("Falha sintética de migração")):
            try:
                ClientesStore(str(pasta / "falha-abertura.db"))
            except sqlite3.OperationalError:
                pass
            else:
                raise AssertionError("Falha de migração precisa propagar")
    assert len(conexoes) == 1
    try:
        conexoes[0].execute("SELECT 1")
    except sqlite3.ProgrammingError:
        pass
    else:
        conexoes[0].close()
        raise AssertionError("Falha no construtor deixou conexão aberta")


def main():
    with tempfile.TemporaryDirectory(prefix="carga_fiscal_021_sintetica_") as tmp:
        pasta = Path(tmp)
        base, leitura = preparar(pasta)
        test_assinatura_e_bloqueios(pasta, base, leitura)
        test_carga_recibo_e_proveniencia(pasta, base)
        test_snapshot_desatualizado(pasta, base)
        test_rollback(pasta, base)
        test_concorrencia(pasta, base)
        test_revalidacao_sob_trava(pasta, base)
        test_constructor_fecha_conexao(pasta)
    print("OK - carga Fiscal: snapshot, origem, recibo imutável, rollback, repetição e concorrência")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
