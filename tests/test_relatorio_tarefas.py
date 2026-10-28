import io
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from openpyxl import load_workbook
from auditoria_fiscal.ferramentas.relatorio_tarefas import gerar_relatorio

def main():
    dados=[dict(id=1,titulo='=HYPERLINK("invalido")',cliente_nome='+Empresa',
        competencia='2026-10',vencimento='2026-11-20',situacao='concluida',
        concluida_por=2,conferida_por=3,fora_do_prazo=True)]
    wb=load_workbook(io.BytesIO(gerar_relatorio(dados)))
    assert wb.active.max_row==2
    assert all(c.data_type!='f' for row in wb.active for c in row)
    assert any(c.value=="'+Empresa" for row in wb.active for c in row)
    assert any(c.value=='Sim' for row in wb.active for c in row)
    wb.close()
    dados = [dict(id=2, titulo='Rotina sintética', competencia='2026-10', vencimento='',
                  ocorrencia=1, fora_do_prazo=False, situacao='concluida'),
             dict(id=3, titulo='Período sintético', competencia='2026-10', vencimento='2026-11-20',
                  ocorrencia=2, origem={'intervalo': '11–20'})]
    wb = load_workbook(io.BytesIO(gerar_relatorio(dados)))
    cabecalho = {c.value: c.column for c in wb.active[1]}
    assert wb.active.cell(2, cabecalho['Vencimento']).value == 'Sem vencimento'
    assert wb.active.cell(2, cabecalho['Fora do prazo']).value == 'Não se aplica'
    assert str(wb.active.cell(3, cabecalho['Ocorrência']).value) == '2'
    assert wb.active.cell(3, cabecalho['Intervalo']).value == '11–20'
    wb.close()
    print('OK - exportação de tarefas com autoria e texto seguro')
    return 0
if __name__=='__main__': raise SystemExit(main())
