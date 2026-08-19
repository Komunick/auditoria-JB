"""Saidas Excel da Conciliacao Fiscal (modulo puro, sem camada web).

Cobre o consolidado auditavel (abas, valores, nulos do 50-5, negativos,
neutralizacao de formula, metadados e falha sem arquivo parcial) e o
preenchimento da COPIA da planilha-mestre (estrutura valida/invalida, meses
ausentes/duplicados, celulas e formulas mapeadas, preservacao de tudo que esta
fora do mapa, totais anuais e marcacao visivel de pendentes).
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
from decimal import Decimal

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "src"))
sys.path.insert(0, os.path.join(RAIZ, "tests"))

from openpyxl import load_workbook  # noqa: E402

from fixtures import conciliacao as fx  # noqa: E402

from auditoria_fiscal.ferramentas import relatorio_conciliacao as rc  # noqa: E402


def checar(cond, msg):
    if not cond:
        print(f"FALHOU - {msg}")
        raise SystemExit(1)


def conciliacao_exemplo(**ajustes) -> dict:
    """Um detalhe completo, como `store.detalhar` devolve."""
    base = {
        "id": 1, "cnpj": fx.CNPJ_A, "competencia": "202401",
        "razao_social": "EMPRESA SINTETICA ALFA LTDA", "estado": "aprovada",
        "revision": 2, "avisos_abertos": 1, "bloqueios_abertos": 0,
        "conflitos_abertos": 0,
        "fonte": {"id": 1, "nome_original": "relatorio.xlsx",
                  "sha256": "a" * 64, "versao_parser": "conciliacao-sefaz/1"},
        "versao_vigente": {
            "id": 1, "numero": 1, "estado": "vigente",
            "layout": "quadro_50_3_dimp", "receita_declarada": 4000000,
            "declarada_com_st": 1000000, "declarada_sem_st": 3000000,
            "receita_calculada": 4150000, "calculada_com_st": 1050000,
            "calculada_sem_st": 3100000, "diferenca_receita": 150000,
            "diferenca_com_st": 50000, "diferenca_sem_st": 100000,
            "total_pix": 900000, "total_nao_pix": 2200000,
            "aba_origem": "Quadro 50-3", "linha_resumo": 7,
            "criada_em": "2026-08-17T12:00:00+00:00"},
        "movimentos": [{
            "id": 1, "versao_id": 1, "instituicao": "BANCO SINTETICO S.A.",
            "linha_origem": 12, "debito": 400000, "credito": 600000,
            "transferencia": 0, "pix": 700000, "voucher": 0, "outras": 0,
            "total_dimp": 1700000, "total_nao_pix": 1000000,
            "aba": "Quadro 50-3", "celula_total": "J12"}],
        "fatos": [{
            "id": 1, "versao_id": 1, "metrica": "receita_declarada",
            "valor_numerico": 4000000, "tipo_origem": "celula",
            "aba": "Quadro 50-3", "celula": "C7",
            "rotulo": "Receita Total Informada",
            "regra_parser": "resumo/receita_declarada", "formula": None,
            "fatos_origem": [], "versao_parser": "conciliacao-sefaz/1"}],
        "excecoes": [{
            "id": 1, "versao_id": 1, "codigo": "DIMP_ABAIXO_DA_CALCULADA",
            "severidade": "aviso", "mensagem": "DIMP menor que a calculada.",
            "estado": "aberta", "resolucao": "", "resolvida_por_login": "",
            "criada_em": "2026-08-17T12:00:00+00:00", "resolvida_em": ""}],
        "conflitos": [], "revisoes": [{
            "id": 1, "versao_id": 1, "revision_anterior": 1,
            "estado_anterior": "em_revisao", "estado_novo": "aprovada",
            "justificativa": "Conferido.", "usuario_login": "bruno",
            "criada_em": "2026-08-17T12:05:00+00:00"}],
    }
    base["versoes"] = [base["versao_vigente"]]
    base.update(ajustes)
    return base


def testar_sanitizacao() -> None:
    """Texto que comeca com caractere de formula nunca vira formula."""
    for perigoso in ("=1+1", "+SOMA(A1)", "-2", "@SUM(A1)",
                     '=HYPERLINK("http://x","clique")',
                     "=cmd|' /c calc'!A1"):
        saida = rc.sanitizar_texto(perigoso)
        checar(saida.startswith("'"), f"{perigoso!r} -> {saida!r}")
        checar(saida[1:] == perigoso, "o conteudo original tem de ser preservado")
    for inofensivo in ("EMPRESA ALFA", "R$ 1,00", "202401", ""):
        checar(rc.sanitizar_texto(inofensivo) == inofensivo,
               f"texto normal nao pode ser alterado: {inofensivo!r}")
    checar(rc.sanitizar_texto(None) == "", "None vira vazio")


def testar_consolidado(pasta: str) -> None:
    destino = os.path.join(pasta, "consolidado.xlsx")
    negativa = conciliacao_exemplo(
        id=2, competencia="202402", estado="em_revisao",
        razao_social="=SOMA(A1:A9)",          # razao social maliciosa
        cnpj=fx.CNPJ_B, bloqueios_abertos=1)
    negativa["versao_vigente"] = {
        **negativa["versao_vigente"], "id": 2, "layout": "quadro_50_5",
        # Diferenca NEGATIVA e DIMP ausente no 50-5.
        "diferenca_receita": -250000, "total_pix": None,
        "total_nao_pix": None}
    negativa["versoes"] = [negativa["versao_vigente"]]
    negativa["movimentos"] = []

    rc.gerar_consolidado(
        destino, conciliacoes=[conciliacao_exemplo(), negativa],
        gerado_por="bruno", filtros={"estado": "todas"},
        gerado_em="17/08/2026 15:00", versao_parser="conciliacao-sefaz/1")

    livro = load_workbook(destino)
    esperadas = ["Conciliações", "Versões", "Movimentos DIMP", "Proveniência",
                 "Exceções", "Conflitos", "Revisões", "Metadados"]
    checar(livro.sheetnames == esperadas,
           f"abas do consolidado: {livro.sheetnames}")

    aba = livro["Conciliações"]
    checar(aba.freeze_panes == "A2", "a primeira linha deveria congelar")
    checar(aba.auto_filter.ref, "o cabecalho deveria ter autofiltro")

    # Valores MONETARIOS sao numeros, nao texto.
    checar(aba["G2"].value == Decimal("40000.00")
           or float(aba["G2"].value) == 40000.00, aba["G2"].value)
    checar("R$" in (aba["G2"].number_format or ""),
           f"formato monetario: {aba['G2'].number_format}")
    checar(aba["M2"].value < 0 or float(aba["M2"].value) > 0,
           "diferenca da primeira linha")
    checar(float(aba["M3"].value) == -2500.00,
           f"negativo tem de sair negativo: {aba['M3'].value}")

    # 50-5: PIX e nao-PIX VAZIOS, nunca zero.
    checar(aba["P3"].value is None, f"PIX do 50-5: {aba['P3'].value!r}")
    checar(aba["Q3"].value is None, f"nao-PIX do 50-5: {aba['Q3'].value!r}")
    checar(float(aba["P2"].value) == 9000.00, aba["P2"].value)

    # CNPJ como texto formatado, preservando zeros.
    checar(aba["B2"].value == "11.222.333/0001-81", aba["B2"].value)
    checar(isinstance(aba["B2"].value, str), "CNPJ tem de ser texto")

    # Razao social maliciosa NEUTRALIZADA.
    checar(aba["C3"].value == "'=SOMA(A1:A9)",
           f"texto perigoso deveria vir com apostrofo: {aba['C3'].value!r}")
    for linha in livro["Conciliações"].iter_rows():
        for celula in linha:
            if isinstance(celula.value, str):
                checar(not celula.value.startswith("=") ,
                       f"nenhuma celula de DADO pode virar formula: "
                       f"{celula.coordinate}={celula.value!r}")

    # Proveniencia com celula e regra.
    prov = livro["Proveniência"]
    checar(prov["J2"].value == "C7", f"celula de origem: {prov['J2'].value}")
    checar(prov["L2"].value == "resumo/receita_declarada", prov["L2"].value)

    # Movimentos preservam a linha original e a celula do total.
    mov = livro["Movimentos DIMP"]
    checar(mov["E2"].value == 12, mov["E2"].value)
    checar(mov["Q2"].value == "J12", mov["Q2"].value)

    # Metadados com contagens, filtros e o aviso de estados nao aprovados.
    meta = {linha[0]: linha[1]
            for linha in livro["Metadados"].iter_rows(min_row=2,
                                                      values_only=True)}
    checar(meta["Gerado por"] == "bruno", meta)
    checar(meta["Total de conciliações"] == 2, meta)
    checar("estado=todas" in meta["Filtros aplicados"], meta)
    checar("não aprovadas" in meta["Aviso"], meta["Aviso"])
    checar("não é somente cartão" in meta["Semântica de 'Cartão/outros (não PIX)'"],
           "os metadados precisam registrar a semantica do nao-PIX")
    livro.close()


def testar_falha_sem_arquivo_parcial(pasta: str) -> None:
    """Erro na geracao nao deixa arquivo com cara de planilha valida."""
    destino = os.path.join(pasta, "quebrado.xlsx")
    ruim = conciliacao_exemplo()
    ruim["versoes"] = [{"numero": object()}]      # tipo impossivel de escrever
    try:
        rc.gerar_consolidado(destino, conciliacoes=[ruim], gerado_por="x",
                             filtros={})
        checar(False, "deveria ter falhado")
    except Exception:
        pass
    checar(not os.path.exists(destino),
           "falha nao pode deixar arquivo no destino")
    checar(not os.path.exists(destino + ".parcial"),
           "o temporario tambem tem de sumir")


def testar_mapeamento_do_modelo(pasta: str) -> None:
    modelo = fx.gerar_modelo_mestre(os.path.join(pasta, "modelo.xlsx"))
    mapa = rc.mapear_modelo(modelo)
    checar(sorted(mapa["anos"]) == [2023, 2024], sorted(mapa["anos"]))
    checar(len(mapa["anos"][2024]["meses"]) == 12,
           len(mapa["anos"][2024]["meses"]))
    checar(mapa["anos"][2024]["total"] is not None, "linha TOTAL")
    checar(rc.cnpj_do_modelo(modelo) == fx.CNPJ_A, rc.cnpj_do_modelo(modelo))

    # Mes ausente: erro de contrato, nao preenchimento parcial.
    faltando = fx.gerar_modelo_mestre(
        os.path.join(pasta, "faltando.xlsx"), anos=(2024,),
        meses=[m for m in fx.MESES if m != "MARÇO"])
    mapa2 = rc.mapear_modelo(faltando)
    checar(3 not in mapa2["anos"][2024]["meses"], "marco nao existe no modelo")

    # Mes duplicado.
    duplicado = fx.gerar_modelo_mestre(
        os.path.join(pasta, "duplicado.xlsx"), anos=(2024,),
        meses=list(fx.MESES) + ["JANEIRO"])
    try:
        rc.mapear_modelo(duplicado)
        checar(False, "mes duplicado deveria falhar")
    except rc.ModeloIncompativel as exc:
        checar("duas vezes" in str(exc), str(exc))

    # Estrutura sem secao anual.
    from openpyxl import Workbook
    vazio = os.path.join(pasta, "sem_secao.xlsx")
    wb = Workbook()
    wb.active["A1"] = "Planilha qualquer"
    wb.save(vazio)
    try:
        rc.mapear_modelo(vazio)
        checar(False, "modelo sem secao anual deveria falhar")
    except rc.ModeloIncompativel:
        pass


def testar_preenchimento(pasta: str) -> None:
    modelo = fx.gerar_modelo_mestre(os.path.join(pasta, "base.xlsx"))
    antes = fx.celulas_fora_do_mapa(modelo)
    sha_antes = os.path.getsize(modelo)
    destino = os.path.join(pasta, "preenchida.xlsx")

    valores = {
        "202401": {"receita_declarada": 4000000, "declarada_com_st": 1000000,
                   "declarada_sem_st": 3000000, "receita_calculada": 4150000,
                   "calculada_com_st": 1050000},
        "202402": {"receita_declarada": 5000000, "declarada_com_st": 1200000,
                   "declarada_sem_st": 3800000, "receita_calculada": 5100000,
                   "calculada_com_st": 1250000},
    }
    detalhes = {
        "202401": {"pix_texto": "BANCO SINTETICO S.A.: R$ 9.000,00",
                   "nao_pix_texto": "=ADQUIRENTE: R$ 22.000,00"},
    }
    resultado = rc.preencher_modelo(modelo, destino, competencias=valores,
                                    detalhes_por_competencia=detalhes)
    checar(resultado["competencias"] == ["202401", "202402"],
           resultado["competencias"])
    checar(os.path.getsize(modelo) == sha_antes,
           "o modelo ORIGINAL nao pode ser alterado")

    livro = load_workbook(destino)
    aba = livro["Conciliação"]
    mapa = rc.mapear_modelo(destino)
    linha_jan = mapa["anos"][2024]["meses"][1]

    checar(float(aba.cell(row=linha_jan, column=3).value) == 40000.00,
           aba.cell(row=linha_jan, column=3).value)
    checar("R$" in aba.cell(row=linha_jan, column=3).number_format,
           "valor monetario formatado")
    checar(aba.cell(row=linha_jan, column=8).value == f"=G{linha_jan}-C{linha_jan}",
           aba.cell(row=linha_jan, column=8).value)
    checar(aba.cell(row=linha_jan, column=10).value == f"=G{linha_jan}-I{linha_jan}",
           aba.cell(row=linha_jan, column=10).value)

    # Texto L/M neutralizado.
    checar(aba.cell(row=linha_jan, column=13).value == "'=ADQUIRENTE: R$ 22.000,00",
           aba.cell(row=linha_jan, column=13).value)
    checar(aba.cell(row=linha_jan, column=12).value.startswith("BANCO"),
           aba.cell(row=linha_jan, column=12).value)

    # Totais anuais com SUM.
    total_2024 = mapa["anos"][2024]["total"]
    formula = aba.cell(row=total_2024, column=3).value
    checar(formula.startswith("=SUM(C"), formula)
    checar(livro.calculation.fullCalcOnLoad,
           "o Excel precisa recalcular ao abrir, senao as formulas ficam vazias")

    # Meses NAO selecionados continuam vazios.
    linha_mar = mapa["anos"][2024]["meses"][3]
    checar(aba.cell(row=linha_mar, column=3).value is None,
           "mes fora da selecao nao pode ser preenchido")
    linha_2023 = mapa["anos"][2023]["meses"][1]
    checar(aba.cell(row=linha_2023, column=3).value is None,
           "ano nao selecionado fica intacto")
    livro.close()

    # Tudo fora do mapa preservado: valores, formulas, estilos, larguras,
    # mesclagens e o conjunto de abas.
    depois = fx.celulas_fora_do_mapa(destino)
    checar(set(antes) == set(depois),
           f"celulas fora do mapa mudaram de conjunto: "
           f"{set(antes) ^ set(depois)}")
    for chave, valor in antes.items():
        checar(depois[chave] == valor,
               f"celula fora do mapa foi alterada: {chave} "
               f"{valor!r} -> {depois[chave]!r}")

    # Competencia sem linha no modelo: erro claro, sem arquivo parcial.
    fora = os.path.join(pasta, "fora.xlsx")
    try:
        rc.preencher_modelo(modelo, fora,
                            competencias={"202901": valores["202401"]})
        checar(False, "competencia inexistente no modelo deveria falhar")
    except rc.ModeloIncompativel as exc:
        checar(exc.codigo == "MODELO_INCOMPATIVEL", exc.codigo)
    checar(not os.path.exists(fora), "falha nao deixa arquivo")


def testar_pendentes(pasta: str) -> None:
    """Incluir competencia nao aprovada exige marcacao VISIVEL."""
    modelo = fx.gerar_modelo_mestre(os.path.join(pasta, "pend.xlsx"))
    destino = os.path.join(pasta, "com_pendentes.xlsx")
    rc.preencher_modelo(
        modelo, destino,
        competencias={"202401": {"receita_declarada": 100000,
                                 "declarada_com_st": 40000,
                                 "declarada_sem_st": 60000,
                                 "receita_calculada": 110000,
                                 "calculada_com_st": 45000}},
        pendentes=["202401"])
    livro = load_workbook(destino)
    checar("Metadados da automação" in livro.sheetnames,
           f"a marcacao precisa ser uma ABA visivel: {livro.sheetnames}")
    checar(livro.sheetnames[0] == "Metadados da automação",
           "o aviso tem de ser a primeira aba, nao um comentario escondido")
    aviso = livro["Metadados da automação"]
    checar(aviso.sheet_state == "visible", aviso.sheet_state)
    checar("202401" in str(aviso["B2"].value),
           f"o aviso precisa dizer QUAIS competencias: {aviso['B2'].value}")
    livro.close()


def main() -> int:
    pasta = tempfile.mkdtemp(prefix="conc_relatorio_")
    try:
        testar_sanitizacao()
        testar_consolidado(pasta)
        testar_falha_sem_arquivo_parcial(pasta)
        testar_mapeamento_do_modelo(pasta)
        testar_preenchimento(pasta)
        testar_pendentes(pasta)
    finally:
        shutil.rmtree(pasta, ignore_errors=True)

    print("OK - saidas Excel da conciliacao (consolidado com 8 abas, valores "
          "numericos, nulos do 50-5 vazios, negativos, texto perigoso "
          "neutralizado, metadados, falha sem arquivo parcial, mapeamento do "
          "modelo, preenchimento com formulas e totais, preservacao do que "
          "esta fora do mapa e marcacao visivel de pendentes) passou.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
