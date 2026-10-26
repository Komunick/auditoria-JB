"""Leitura da planilha de senhas atual (spec 012, US3).

A planilha sintetica imita a ESTRUTURA da real (lida em 22/10/2026 sem
exibir valores): cabecalho na linha 2, nome na coluna A sem rotulo, CNPJ
numerico, situacao na cor da linha, legenda misturada nas linhas de cliente,
colunas sem cabecalho com conteudo, linha oculta, CNPJ repetido, data
impossivel e "NAO TEM"/"DEFINITIVO" no lugar de valor. Nomes, documentos e
senhas sao inventados.
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
from datetime import datetime

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "src"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402

from auditoria_fiscal.core import clientes as cl  # noqa: E402
from auditoria_fiscal.core import planilha_clientes as pc  # noqa: E402
from auditoria_fiscal.core.documentos_br import dv_cnpj  # noqa: E402


def cnpj(base12: str) -> str:
    return base12 + dv_cnpj(base12)


C1 = cnpj("011111111000")        # com zero a esquerda
C2 = cnpj("222222220001")
C3 = cnpj("333333330001")
C4 = cnpj("444444440001")
C5 = cnpj("555555550001")
SENHAS = ["Sefaz@Alfa1", "pref-alfa-9", "cod-alfa-1234", "Sefaz@Beta2",
          "SenhaColunaN", "SenhaColunaP"]

CABECALHO = ["", "CNPJ", "INSCRIÇÃO ESTADUAL", "INSCRIÇÃO MUNICIPAL",
             "SENHA SEFAZ ESTADUAL", "SENHA PREFEITURA", "IPTU", "TVL",
             "VENC TVL", "VENC ALVARA", "CÓDIGO ACESSO", "PUBLICIDADE",
             "ALVARÁ SAÚDE", "", "LEGENDA", ""]

VERMELHO, LARANJA, CINZA, ROSA, VERDE = (
    "FFFF0000", "FFFF9900", "FF999999", "FFF9CB9C", "FFB2BDD2")


def pintar(aba, linha, cor):
    aba.cell(linha, 1).fill = PatternFill("solid", fgColor=cor)


def planilha(caminho: str) -> None:
    livro = Workbook()
    aba = livro.active
    aba.title = "Plan1"
    aba.cell(1, 1, " ")
    for coluna, titulo in enumerate(CABECALHO, 1):
        if titulo:
            aba.cell(2, coluna, titulo)
    linhas = [
        # 3: ativo, CNPJ numerico com zero comido, tudo preenchido
        ["Alfa Comercio Ltda", float(int(C1)), "123456789", 12345678.0,
         SENHAS[0], SENHAS[1], 1234567.0, 87654321.0, datetime(2026, 12, 1),
         datetime(2027, 3, 1), SENHAS[2], "NÃO TEM", datetime(2026, 11, 5)],
        # 4: baixado (vermelho), legenda na coluna O, coluna N sem cabecalho
        ["Beta Servicos", float(int(C2)), "NÃO TEM", "PENTENDE", SENHAS[3],
         "PENDENTE", None, None, "DEFINITIVO", datetime(2878, 7, 13), None,
         None, "VENCIDO", SENHAS[4], "BAIXADAS"],
        # 5: inativo (laranja), coluna P com RENOVAR e outra "senha"
        ["Gama Padaria", C3, "AGUARDANDO CERTIFICADO", None, "N", None,
         None, None, None, None, None, None, None, None, None, "RENOVAR"],
        # 6: Lauro de Freitas (rosa), cor da legenda
        ["Delta Mercado", C4, "12.345.678", "IM PENDENTE", None, None, None,
         None, "DEFFINITIVO", None, None, None, None, None, None, SENHAS[5]],
        # 7: cinza = nao e' mais cliente; linha OCULTA
        ["Epsilon Ltda", C5, None, None, None, None],
        # 8: CNPJ repetido do Alfa (primeira ocorrencia na linha 3)
        ["Alfa Filial Repetida", C1, None, None, None, None],
        # 9: CNPJ invalido e cor fora da legenda
        ["Zeta Sem Documento", "12345678000100", None, None, None, None],
        # 10: linha sem nome e sem nada: e' ignorada sem barulho
        [None, None, None],
        # 11: so documento, sem nome: vira pendencia de nome
        [None, cnpj("666666660001")],
    ]
    for deslocamento, valores in enumerate(linhas):
        for coluna, valor in enumerate(valores, 1):
            if valor is not None:
                aba.cell(3 + deslocamento, coluna, valor)
    pintar(aba, 3, "FFFFFFFF")
    pintar(aba, 4, VERMELHO)
    pintar(aba, 5, LARANJA)
    pintar(aba, 6, ROSA)
    pintar(aba, 7, CINZA)
    pintar(aba, 9, VERDE)
    aba.row_dimensions[7].hidden = True
    livro.save(caminho)


def checar(cond, msg):
    if not cond:
        print(f"FALHOU - {msg}")
        raise SystemExit(1)


def main() -> int:
    pasta = tempfile.mkdtemp(prefix="planilha_clientes_")
    caminho = os.path.join(pasta, "senhas.xlsx")
    planilha(caminho)

    leitura = pc.ler_planilha_clientes(caminho)
    checar(leitura.linha_cabecalho == 2, leitura.linha_cabecalho)
    checar(len(leitura.hash_arquivo) == 64, "hash SHA-256 do arquivo")
    por_linha = {l.linha: l for l in leitura.linhas}
    checar(sorted(por_linha) == [3, 4, 5, 6, 7, 8, 9, 11],
           f"linhas lidas: {sorted(por_linha)} (a 10, vazia, sai sem barulho)")

    # Colunas sem cabecalho (N e P) sao ignoradas, mas aparecem no relatorio.
    checar(leitura.colunas_sem_cabecalho == ["N", "P"],
           leitura.colunas_sem_cabecalho)
    # A legenda divide as linhas com os clientes: a coluna inteira e' ignorada.
    checar(leitura.colunas_ignoradas == ["O"], leitura.colunas_ignoradas)

    alfa = por_linha[3]
    checar(alfa.nome == "Alfa Comercio Ltda", alfa.nome)
    checar(alfa.documento == C1, f"zero a esquerda recuperado: {alfa.documento}")
    checar(alfa.situacao == cl.SITUACAO_ATIVO, alfa.situacao)
    checar((alfa.ie, alfa.ie_situacao) == ("123456789", ""), "IE")
    checar((alfa.im, alfa.im_situacao) == ("12345678", ""), "IM sem '.0'")
    checar(alfa.iptu == "1234567", alfa.iptu)
    checar(alfa.segredos["sefaz"] == (SENHAS[0], ""), "senha SEFAZ")
    checar(alfa.segredos["prefeitura"] == (SENHAS[1], ""), "senha Prefeitura")
    checar(alfa.segredos["codigo_acesso"] == (SENHAS[2], ""), "codigo")
    checar(alfa.licencas["tvl"]["numero"] == "87654321", alfa.licencas["tvl"])
    checar(alfa.licencas["tvl"]["vencimento"] == "2026-12-01",
           alfa.licencas["tvl"])
    checar(alfa.licencas["alvara"]["vencimento"] == "2027-03-01", "alvara")
    checar(alfa.licencas["publicidade"]["situacao"] == cl.CAMPO_NAO_TEM,
           alfa.licencas["publicidade"])
    checar(alfa.licencas["alvara_saude"]["vencimento"] == "2026-11-05",
           "alvara saude")
    checar(alfa.acao == pc.ACAO_INCLUIR and not alfa.pendencias,
           f"linha limpa: {alfa.pendencias}")

    beta = por_linha[4]
    checar(beta.situacao == cl.SITUACAO_BAIXADO, "vermelho = baixado")
    checar(beta.ie_situacao == cl.CAMPO_NAO_TEM, beta.ie_situacao)
    checar(beta.im_situacao == cl.CAMPO_PENDENTE, "PENTENDE = pendente")
    checar(beta.segredos["prefeitura"] == ("", cl.CAMPO_PENDENTE), "pendente")
    checar(beta.licencas["tvl"]["situacao"] == cl.CAMPO_DEFINITIVO, "DEFINITIVO")
    checar(beta.licencas["alvara"]["situacao"] == cl.CAMPO_DATA_INVALIDA
           and beta.licencas["alvara"]["vencimento"] == ""
           and "2878" in beta.licencas["alvara"]["texto_original"],
           f"ano 2878 vira data invalida: {beta.licencas['alvara']}")
    checar(any("2878" in p or "data" in p.lower() for p in beta.pendencias),
           f"data invalida e' pendencia: {beta.pendencias}")
    checar(beta.licencas["alvara_saude"]["situacao"] == cl.CAMPO_VENCIDO,
           "alvara saude vencido")

    gama = por_linha[5]
    checar(gama.situacao == cl.SITUACAO_INATIVO, "laranja = inativa")
    checar(gama.ie_situacao == cl.CAMPO_AGUARDANDO_CERTIFICADO, gama.ie_situacao)
    checar(gama.segredos["sefaz"] == ("", cl.CAMPO_NAO_TEM), "N = nao tem")

    delta = por_linha[6]
    checar(delta.situacao == cl.SITUACAO_ATIVO
           and delta.municipio == "Lauro de Freitas",
           "rosa e' municipio, nao situacao")
    checar(delta.ie == "12.345.678", "IE formatada fica como veio")
    checar(delta.licencas["tvl"]["situacao"] == cl.CAMPO_DEFINITIVO,
           "DEFFINITIVO (com erro) = definitivo")

    epsilon = por_linha[7]
    checar(epsilon.situacao == cl.SITUACAO_NAO_CLIENTE, "cinza")
    checar(epsilon.oculta and any("oculta" in p.lower()
                                  for p in epsilon.pendencias),
           "linha oculta vira pendencia")
    checar(epsilon.acao == pc.ACAO_INCLUIR, "oculta entra, mas sinalizada")

    repetida = por_linha[8]
    checar(repetida.acao == pc.ACAO_IGNORAR and "3" in repetida.motivo,
           f"CNPJ repetido: a segunda ocorrencia vem para ignorar "
           f"({repetida.acao}, {repetida.motivo})")

    zeta = por_linha[9]
    checar(zeta.documento == "" and zeta.documento_informado == "12345678000100",
           "documento invalido nao e' inventado")
    checar(any("CNPJ" in p or "documento" in p.lower() for p in zeta.pendencias),
           zeta.pendencias)
    checar(any("cor" in p.lower() for p in zeta.pendencias),
           f"cor fora da legenda: {zeta.pendencias}")

    sem_nome = por_linha[11]
    checar(sem_nome.acao == pc.ACAO_IGNORAR and "nome" in sem_nome.motivo.lower(),
           "sem nome nao entra")

    # ------------------------------------------------------------------
    # O resumo (o que vai para a tela) NUNCA leva senha.

    resumo = json.dumps(leitura.resumo(), ensure_ascii=False)
    for senha in SENHAS:
        checar(senha not in resumo, f"a previa nao pode conter {senha!r}")
    checar('"sefaz": "presente"' in resumo, "a previa diz que a senha existe")
    contagens = leitura.resumo()["contagens"]
    checar(contagens["incluir"] == 6 and contagens["ignorar"] == 2, contagens)
    checar(contagens["com_pendencia"] >= 3, contagens)

    # ------------------------------------------------------------------
    # Arquivo hostil e' recusado antes do openpyxl.

    falso = os.path.join(pasta, "falso.xlsx")
    with open(falso, "wb") as arquivo:
        arquivo.write(b"nao e zip")
    try:
        pc.ler_planilha_clientes(falso)
        checar(False, "arquivo que nao e' xlsx deveria ser recusado")
    except cl.ErroClientes as exc:
        checar(exc.codigo == pc.PLANILHA_INVALIDA, exc.codigo)

    sem_cnpj = os.path.join(pasta, "sem_cnpj.xlsx")
    livro = Workbook()
    livro.active.cell(1, 1, "Nome")
    livro.save(sem_cnpj)
    try:
        pc.ler_planilha_clientes(sem_cnpj)
        checar(False, "planilha sem coluna CNPJ deveria ser recusada")
    except cl.ErroClientes as exc:
        checar(exc.codigo == pc.CABECALHO_NAO_ENCONTRADO, exc.codigo)

    print("OK - planilha de clientes (cabecalho na linha 2, CNPJ numerico, "
          "cor vira situacao, rosa vira municipio, linha oculta, colunas sem "
          "cabecalho ignoradas, CNPJ repetido, data impossivel, previa sem "
          "senha, arquivo hostil recusado) passou.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
