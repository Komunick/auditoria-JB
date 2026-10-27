"""Fluxo HTTP sintético e isolamento por pessoa, setor, período e prévia."""
import io
import json
import os
import sys
import tempfile
import threading
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
os.environ['AUDITORIA_WEB_DADOS']=tempfile.mkdtemp(prefix='web_tarefas_')
from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook
from auditoria_fiscal.web import rotas_tarefas as rt, auth, permissoes
from auditoria_fiscal.web.servidor import criar_app
from auditoria_fiscal.ferramentas.tarefas_store import TarefasStore


def regressao_migracao_auth_concorrente():
    """Dois primeiros acessos ao schema antigo nao disputam o ALTER de email."""
    caminho_original = auth.caminho_db_usuarios
    conectar_original = sqlite3.connect
    with tempfile.TemporaryDirectory(prefix='auth_schema_antigo_') as pasta:
        banco = str(Path(pasta) / 'usuarios.db')
        auth.caminho_db_usuarios = lambda: banco
        try:
            conn = auth.conexao()
            try:
                conn.execute("INSERT INTO usuario(usuario,nome,senha_hash,salt,admin,ativo,criado_em) VALUES('legado_sintetico','Legado Sintetico','hash-sintetico','salt-sintetico',1,1,'2026-10-01')")
                conn.commit()
                conn.execute('ALTER TABLE usuario DROP COLUMN email')
                conn.commit()
                antes = tuple(conn.execute('SELECT * FROM usuario').fetchone())
            finally:
                conn.close()
            barreira = threading.Barrier(2)
            class ConexaoConcorrente(sqlite3.Connection):
                def execute(self, sql, *args, **kwargs):
                    resultado = super().execute(sql, *args, **kwargs)
                    if sql == 'PRAGMA table_info(usuario)':
                        linhas = list(resultado)
                        # Na versao vulneravel as duas leituras veem a coluna ausente.
                        # Com BEGIN IMMEDIATE, o segundo escritor espera antes da leitura.
                        if not self.in_transaction and not any(r['name'] == 'email' for r in linhas):
                            barreira.wait(timeout=5)
                        return linhas
                    return resultado
            def conectar(*args, **kwargs):
                kwargs['factory'] = ConexaoConcorrente
                return conectar_original(*args, **kwargs)
            def migrar(_):
                try:
                    conexao = auth.conexao()
                    conexao.close()
                    return None
                except Exception as exc:
                    return type(exc).__name__ + ': ' + str(exc)
            auth.sqlite3.connect = conectar
            with ThreadPoolExecutor(max_workers=2) as executor:
                erros = list(executor.map(migrar, range(2)))
            assert erros == [None, None], erros
            auth.sqlite3.connect = conectar_original
            conn = auth.conexao()
            try:
                depois = tuple(conn.execute('SELECT * FROM usuario').fetchone())
                assert depois[:-1] == antes and depois[-1] == '', 'Migracao deve preservar identidades e credenciais'
                assert [r['name'] for r in conn.execute('PRAGMA table_info(usuario)')].count('email') == 1
            finally:
                conn.close()
        finally:
            auth.sqlite3.connect = conectar_original
            auth.caminho_db_usuarios = caminho_original


def regressao_catalogo_e_concorrencia(adm, gestor, colab, pessoas, setor, outro_setor, cid, cid2):
    api = '/api/tarefas'
    # Ser membro de B permite ver o catalogo, nao as excecoes privadas de terceiros.
    outro_id = pessoas['outro'][0]
    gestor_id = pessoas['gestor'][0]
    ok(adm.put(f'{api}/setores/{outro_setor}', json=dict(nome='Setor Separado',
        responsavel_usuario_id=outro_id, membros=[outro_id, gestor_id])))
    obrigacao = ok(adm.post(api+'/obrigacoes', json=dict(setor_id=outro_setor,
        nome='Obrigacao Setor Separado', regra_vencimento=dict(tipo='dia_mes_seguinte',dia=20))))
    ajuste = ok(adm.post(api+'/ajustes', json=dict(cliente_id=cid2, obrigacao_id=obrigacao['id'],
        tipo='excluir', motivo='Motivo privado do outro setor')))
    catalogo = ok(gestor.get(api+'/catalogo'))
    assert obrigacao['id'] in {o['id'] for o in catalogo['obrigacoes']}
    assert ajuste['id'] not in {a['id'] for a in catalogo['ajustes']}, 'Membro de outro setor nao deve receber excecoes de toda a carteira'
    obrigacao_propria = next(o['id'] for o in catalogo['obrigacoes'] if o['setor_id'] == setor)
    ajuste_proprio = ok(gestor.post(api+'/ajustes', json=dict(cliente_id=cid,
        obrigacao_id=obrigacao_propria,tipo='excluir',motivo='Excecao sintetica a corrigir')))
    correcao = dict(cliente_id=cid,obrigacao_id=obrigacao_propria,tipo='incluir',
        motivo='Correcao justificada da excecao',revision=ajuste_proprio['revision'])
    corrigido = ok(gestor.put(f"{api}/ajustes/{ajuste_proprio['id']}",json=correcao))
    assert corrigido['tipo'] == 'incluir' and corrigido['revision'] == ajuste_proprio['revision'] + 1
    ok(gestor.put(f"{api}/ajustes/{ajuste_proprio['id']}",json=correcao),409)
    ok(gestor.put(f"{api}/ajustes/{ajuste['id']}",json=dict(cliente_id=cid2,
        obrigacao_id=obrigacao['id'],tipo='incluir',motivo='Mudanca fora do setor',revision=ajuste['revision'])),403)

    tarefa = ok(adm.post(api+'/avulsas', json=dict(cliente_id=cid,setor_id=setor,
        titulo='Conferencia concorrente sintetica',inicio='2026-10-01',vencimento='2026-11-20',
        responsavel_usuario_id=pessoas['colab'][0],exige_conferencia=True)))
    for estado in ('em_andamento','concluida'):
        tarefa = ok(colab.post(f"{api}/{tarefa['id']}/situacao",json=dict(situacao=estado,revision=tarefa['revision'])))
    preparado = threading.Event()
    disparar = threading.Event()
    tentando = threading.Event()
    terminou = threading.Event()
    erros = []
    original = TarefasStore.mudar_situacao
    def revogar_em_outra_conexao():
        store = None
        try:
            store = rt.criar_store()
            preparado.set()
            if not disparar.wait(timeout=10):
                raise AssertionError('A requisicao nao chegou ao ponto de concorrencia')
            tentando.set()
            store.salvar_setor(dict(nome='Fiscal Teste',responsavel_usuario_id=outro_id,
                membros=[pessoas[n][0] for n in ('colab','outro','gestor','substituto')]),
                id=setor,usuario_id=1,usuario_login='admin_teste')
        except Exception as exc:
            erros.append(exc)
        finally:
            if store is not None:
                store.fechar()
            terminou.set()
    def mudar_com_disputa(store, *args, **kwargs):
        disparar.set()
        assert tentando.wait(timeout=5), 'O segundo escritor nao iniciou'
        assert not terminou.wait(timeout=.2), 'A chefia mudou entre a autorizacao e a gravacao'
        return original(store, *args, **kwargs)
    worker = threading.Thread(target=revogar_em_outra_conexao,daemon=True)
    worker.start()
    try:
        assert preparado.wait(timeout=5), erros
        TarefasStore.mudar_situacao = mudar_com_disputa
        confirmada = ok(gestor.post(f"{api}/{tarefa['id']}/situacao",
            json=dict(situacao='conferida',revision=tarefa['revision'])))
        assert confirmada['conferida_por'] == gestor_id
    finally:
        TarefasStore.mudar_situacao = original
        disparar.set()
        worker.join(timeout=10)
    assert not worker.is_alive() and not erros, erros
    assert terminou.is_set()
    ok(gestor.post(f"{api}/{tarefa['id']}/situacao",
        json=dict(situacao='conferida',revision=confirmada['revision'])),403)


def ok(r, status=200):
    assert r.status_code==status,(r.status_code,r.text)
    return r.json() if status==200 else None


def main():
    regressao_migracao_auth_concorrente()
    rt._hoje=lambda:date(2026,10,27)
    app=criar_app(); adm=TestClient(app)
    ok(adm.post('/api/bootstrap',json=dict(usuario='admin_teste',nome='Admin Teste',senha='sintetica-123')))
    api='/api/tarefas'
    ok(TestClient(app).get(api+'/catalogo'),401)
    assert not set(('aba.tarefas','tarefas.executar','tarefas.setor','tarefas.todas','tarefas.configurar','setores.administrar','tarefas.exportar')) & set(permissoes.PADRAO_NOVO_USUARIO)
    pessoas={}
    for nome,slugs in [('colab',['aba.tarefas','tarefas.executar']),
                       ('outro',['aba.tarefas','tarefas.executar']),
                       ('gestor',['aba.tarefas','tarefas.executar','tarefas.setor','tarefas.configurar','tarefas.exportar']),
                       ('substituto',['aba.tarefas','tarefas.setor','tarefas.executar']),
                       ('diretoria',['aba.tarefas','tarefas.todas','tarefas.exportar']),('bloqueado',[])]:
        u=ok(adm.post('/api/admin/usuarios',json=dict(usuario=nome,nome='Pessoa '+nome,senha='sintetica-123',email=nome+'@example.test',permissoes=slugs)))
        c=TestClient(app);ok(c.post('/api/login',json=dict(usuario=nome,senha='sintetica-123')))
        pessoas[nome]=(u['id'],c)
        assert u['email']==nome+'@example.test'
    colab=pessoas['colab'][1]; gestor=pessoas['gestor'][1]; sub=pessoas['substituto'][1]
    ok(pessoas['bloqueado'][1].get(api+'/catalogo'),403)
    setor=ok(adm.post(api+'/setores',json=dict(nome='Fiscal Teste',responsavel_usuario_id=pessoas['gestor'][0],membros=[pessoas[n][0] for n in ['colab','outro','gestor','substituto']])) )['id']
    outro_setor=ok(adm.post(api+'/setores',json=dict(nome='Setor Separado',responsavel_usuario_id=pessoas['outro'][0],membros=[pessoas['outro'][0]])))['id']
    cid=ok(adm.post('/api/clientes',json=dict(nome='Empresa Sintética A',contato_nome='Contato Teste',contato_email='contato@example.test')))['id']
    cid2=ok(adm.post('/api/clientes',json=dict(nome='Empresa Sintética B')))['id']
    assert ok(adm.get(f'/api/clientes/{cid}'))['contato_email']=='contato@example.test'
    ok(gestor.put(f'{api}/clientes/{cid}/carteira',json=dict(setor_id=setor,responsavel_usuario_id=pessoas['colab'][0])))
    ok(adm.put(f'{api}/clientes/{cid2}/carteira',json=dict(setor_id=outro_setor,responsavel_usuario_id=pessoas['outro'][0])))
    car=ok(gestor.post(api+'/caracteristicas',json=dict(grupo='Teste',nome='Aplicável',exclusiva_no_grupo=False)))['id']
    ok(gestor.put(f'{api}/clientes/{cid}/caracteristicas',json=dict(caracteristica_ids=[car])))
    ob=dict(setor_id=setor,nome='Obrigação Sintética',periodicidade='mensal',competencias=[],regra_vencimento=dict(tipo='dia_mes_seguinte',dia=20,sem_expediente='manter'),condicao=dict(todas=[car]),exige_conferencia=True)
    oid=ok(gestor.post(api+'/obrigacoes',json=ob))['id']
    ok(colab.post(api+'/obrigacoes',json=ob),403)
    ok(gestor.post(api+'/obrigacoes',json=dict(ob,setor_id=outro_setor)),403)
    previa=ok(gestor.post(api+'/geracao/previa',json=dict(competencia='2026-10',setor_id=setor)))
    assert previa['total']==1,previa
    ok(colab.post(api+'/geracao/confirmar',json=dict(token=previa['token'])),403)
    ok(gestor.post(api+'/geracao/confirmar',json=dict(token=previa['token'])))
    ok(gestor.post(api+'/geracao/confirmar',json=dict(token=previa['token'])),404)
    tarefa=ok(colab.get(api))['itens'][0]; tid=tarefa['id']
    assert tarefa['pode_executar'] and not tarefa['pode_gerir']
    outra=ok(adm.post(api+'/avulsas',json=dict(cliente_id=cid2,setor_id=outro_setor,titulo='Outra carteira',vencimento='2026-10-28',responsavel_usuario_id=pessoas['outro'][0])))
    assert ok(colab.get(api))['total']==1
    ok(colab.get(f"{api}/{outra['id']}"),403)
    ok(colab.post(f"{api}/{outra['id']}/situacao",json=dict(situacao='em_andamento',revision=1)),403)
    assert ok(pessoas['diretoria'][1].get(api,params={'visao':'setor'}))['total']==2
    ok(pessoas['diretoria'][1].post(f'{api}/{tid}/situacao',json=dict(situacao='em_andamento',revision=1)),403)
    tarefa=ok(colab.post(f'{api}/{tid}/situacao',json=dict(situacao='em_andamento',revision=tarefa['revision'])))
    ok(colab.post(f'{api}/{tid}/comentarios',json=dict(comentario='Dado privado no domínio',revision=1)),409)
    tarefa=ok(colab.post(f'{api}/{tid}/situacao',json=dict(situacao='concluida',revision=tarefa['revision'])))
    ok(colab.post(f'{api}/{tid}/situacao',json=dict(situacao='conferida',revision=tarefa['revision'])),403)
    ok(adm.post(api+'/substituicoes',json=dict(titular_usuario_id=pessoas['gestor'][0],substituto_usuario_id=pessoas['substituto'][0],setor_id=setor,inicio='2026-10-27',fim='2026-10-28',motivo='Ausência sintética')))
    ok(gestor.get(api,params={'visao':'setor','setor_id':setor}),403)
    assert ok(sub.get(api,params={'visao':'setor','setor_id':setor}))['total']==1
    tarefa=ok(sub.post(f'{api}/{tid}/situacao',json=dict(situacao='conferida',revision=tarefa['revision'])))
    assert tarefa['conferida_por']==pessoas['substituto'][0]
    assert tarefa['conferida_substituindo']==pessoas['gestor'][0]
    rt._hoje=lambda:date(2026,10,29)
    ok(sub.get(api,params={'visao':'setor','setor_id':setor}),403)
    assert ok(gestor.get(api,params={'visao':'setor','setor_id':setor}))['total']==1
    # Prévia de carteira vinculada ao dono e consumida uma vez.
    wb=Workbook(); ws=wb.active;ws.append(['Carteira sintética']);ws.append(['Empresas','Contato','E-mail','Observação']);ws.append(['Empresa Sintética A','Contato importado','novo@example.test','Sem movimento'])
    arquivo=io.BytesIO();wb.save(arquivo)
    p=ok(gestor.post(api+'/importacao/previa',data={'setor_id':str(setor),'responsavel_usuario_id':str(pessoas['colab'][0])},files={'arquivo':('carteira.xlsx',arquivo.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}))
    decisao=dict(token=p['token'],decisoes=[dict(linha=3,cliente_id=cid,caracteristica_ids=[car])])
    ok(adm.post(api+'/importacao/confirmar',json=decisao),404)
    ok(gestor.post(api+'/importacao/confirmar',json=decisao))
    ok(gestor.post(api+'/importacao/confirmar',json=decisao),404)
    assert ok(adm.get(f'/api/clientes/{cid}'))['contato_email']=='novo@example.test'
    r=gestor.get(api+'/exportar',params={'visao':'setor','setor_id':setor});assert r.status_code==200,r.text
    exportado=load_workbook(io.BytesIO(r.content));assert exportado.active.max_row==2
    ok(colab.get(api+'/exportar'),403)
    avisos=ok(colab.get(api+'/avisos'));assert avisos['email_ativo'] is False
    trilha=ok(adm.get('/api/admin/historico',params={'limite':500}))['itens']
    eventos=[e for e in trilha if e.get('acao','').startswith(('tarefas.','setores.'))]
    assert eventos
    assert not any(x in json.dumps(eventos,ensure_ascii=False) for x in ['Empresa Sintética','novo@example.test','Dado privado no domínio'])
    regressao_catalogo_e_concorrencia(adm,gestor,colab,pessoas,setor,outro_setor,cid,cid2)
    print('OK - tarefas web: autorização, substituição, prévias, autoria, carteira, exportação e trilha')
    return 0
if __name__=='__main__':raise SystemExit(main())
