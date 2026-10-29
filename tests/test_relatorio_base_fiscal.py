"""Spec 019: relatório fiel, texto literal e propostas separadas da origem."""
from copy import deepcopy
from contextlib import suppress
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import zipfile
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from auditoria_fiscal.core.clientes import ErroClientes
from auditoria_fiscal.core.documentos_br import dv_cnpj
from auditoria_fiscal.core.planilha_empresas import ler_planilha_empresas, LeituraEmpresas, LinhaEmpresa
from auditoria_fiscal.ferramentas.relatorio_base_fiscal import gerar_relatorio_base_fiscal


def fonte_sintetica(pasta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Carteira sintética"
    ws.append(["Base sintética"])
    ws.append([None, "RESPONSAVEL", "EMPRESAS", "RESPONSAVEL", "CNPJ", "REGIME DA EMPRESA"])
    zero = "012345670001" + dv_cnpj("012345670001")
    outro = "519876540001" + dv_cnpj("519876540001")
    ws.append([None, "@Pessoa sintética", '=HYPERLINK("https://invalid.example")', "+Contato sintético", int(zero), "Simples Nacional"])
    ws["C3"].data_type = "s"
    ws["E3"].number_format = '00"."000"."000"/"0000"-"00'
    ws.append([None, None, "-Outra empresa sintética", None, "=SOURCE_SECRET()", "A conferir"])
    ws.row_dimensions[4].hidden = True
    ws.merge_cells("D3:D4")
    ws.append([None, "Pessoa 2", "\tEmpresa com espaços  ", "=CONTACT_SECRET()", outro, " presumido "])
    zeradas = wb.create_sheet("ZERADAS")
    zeradas.sheet_state = "hidden"
    zeradas.append([None, "EMPRESA", "CNPJ"])
    zeradas.append([None, "Nome alternativo sintético", zero])
    wb.create_sheet("Vazia")
    caminho = Path(pasta) / "base-sintetica.xlsx"
    wb.save(caminho)
    wb.close()
    return ler_planilha_empresas(caminho), zero


def tabelar(ws):
    linhas = list(ws.iter_rows(values_only=True))
    return [dict(zip(linhas[0], linha)) for linha in linhas[1:]]


def test_fidelidade_e_propostas(pasta):
    leitura, zero = fonte_sintetica(pasta)
    antes = deepcopy(leitura)
    dados = gerar_relatorio_base_fiscal(leitura, nome_arquivo="=base-sintetica.xlsx")
    assert isinstance(dados, bytes) and dados.startswith(b"PK")
    assert leitura == antes, "Exportação alterou a leitura original"
    wb = load_workbook(BytesIO(dados), data_only=False)
    try:
        assert wb.sheetnames == ["Resumo", "Conferência", "Duplicidades", "Correções", "Origens"]
        conferidos = tabelar(wb["Conferência"])
        assert len(conferidos) == len(leitura.linhas) == 4
        assert [l["ID de origem"] for l in conferidos] == [l.id for l in leitura.linhas]
        assert conferidos[0]["CNPJ normalizado"] == zero
        assert conferidos[0]["CNPJ informado"] == str(int(zero))
        assert conferidos[0]["Situação do CNPJ"] == "zero_inicial_recuperado"
        assert conferidos[0]["Sem movimento"] == "Não informado"
        assert conferidos[1]["Linha oculta na fonte"] == "Sim"
        assert conferidos[1]["CNPJ informado"] == "[fórmula não lida]"
        assert conferidos[3]["Sem movimento"] == "Sim — permanece na carteira"
        assert conferidos[3]["Aba oculta na fonte"] == "Sim"
        for campo in ("Regime informado", "Regime reconhecido", "Responsável Fiscal informado", "Contato da empresa"):
            assert conferidos[3][campo] == "Coluna ausente no layout"
        assert conferidos[3]["Pendências de leitura"] == "; ".join(leitura.linhas[3].pendencias)

        repetidos = tabelar(wb["Duplicidades"])
        assert len(repetidos) == 2
        assert {r["ID de origem"] for r in repetidos} == {leitura.linhas[0].id, leitura.linhas[3].id}
        assert {r["CNPJ normalizado"] for r in repetidos} == {zero}
        propostas = tabelar(wb["Correções"])
        assert len(propostas) == 4
        for origem, linha in zip(leitura.linhas, propostas):
            assert linha["ID de origem"] == origem.id
            assert linha["Empresa original"] == origem.nome
            assert (linha["CNPJ normalizado original"] or "") == origem.documento
            for campo in ("CNPJ proposto", "Empresa proposta", "Regime proposto", "Responsável Fiscal proposto", "Contato proposto", "Decisão sobre duplicidade", "Observações da revisão"):
                assert linha[campo] is None, "Decisão humana não pode ser inferida"

        origens = tabelar(wb["Origens"])
        assert len(origens) == sum(len(l.origem) for l in leitura.linhas)
        contato = next(o for o in origens if o["ID de origem"] == "1:4" and o["Campo"] == "contato_nome")
        assert contato["Célula"] == "D4" and contato["Célula de origem"] == "D3"
        assert contato["Intervalo mesclado"] == "D3:D4" and contato["Propagado"] == "Sim"
        assert contato["Valor informado"] == "+Contato sintético"
        formula = next(o for o in origens if o["ID de origem"] == "1:4" and o["Campo"] == "documento")
        assert formula["Valor informado"] == "[fórmula não lida]" and formula["Fórmula não lida"] == "Sim"
        resumo = {linha["Item"]: linha["Valor"] for linha in tabelar(wb["Resumo"])}
        assert resumo["SHA-256 da fonte"] == leitura.hash_arquivo
        assert resumo["Arquivo de origem"] == "=base-sintetica.xlsx"
        assert resumo["Registros no arquivo"] == 4
        assert resumo["Grupos de CNPJ repetido"] == 1
        from auditoria_fiscal.ferramentas.relatorio_base_fiscal import _ROTULOS_CONTAGENS
        for chave, rotulo in _ROTULOS_CONTAGENS:
            assert resumo[rotulo] == leitura.resumo()["contagens"][chave]
        assert "não é aceito" in resumo["Uso do relatório"]
        assert "Aba vazia ignorada: Vazia." in str(list(wb["Resumo"].values))
        for ws in wb:
            assert ws.freeze_panes == "A2"
            assert ws.auto_filter.ref
            for row in ws:
                for celula in row:
                    assert celula.data_type != "f"
                    assert celula.hyperlink is None
                    if isinstance(celula.value, str):
                        assert celula.data_type == "s" and celula.number_format == "@"
        for nome in ("Conferência", "Correções"):
            assert wb[nome].max_row == len(leitura.linhas) + 1
    finally:
        wb.close()
    with zipfile.ZipFile(BytesIO(dados)) as pacote:
        assert not any("externalLinks" in n or "vbaProject" in n for n in pacote.namelist())
        xml = b"".join(pacote.read(n) for n in pacote.namelist() if n.endswith(".xml"))
        assert b"SOURCE_SECRET" not in xml and b"CONTACT_SECRET" not in xml
        assert b"<f>" not in xml


def test_representacao_sem_truncar(pasta):
    leitura, _ = fonte_sintetica(pasta)
    for valor in ("z" * 32768, "SEGREDO\x00", "SEGREDO\ud800", "SEGREDO\ufffe", "😀" * 16384):
        copia = deepcopy(leitura)
        copia.linhas[0].nome = valor
        try:
            gerar_relatorio_base_fiscal(copia, nome_arquivo="base.xlsx")
        except ErroClientes as exc:
            assert exc.codigo == "PLANILHA_INVALIDA"
            assert "1:3" in exc.mensagem and "Empresa" in exc.mensagem
            assert "SEGREDO" not in exc.mensagem and valor not in exc.mensagem
        else:
            raise AssertionError("Valor não representável foi truncado ou aceito")
    copia = deepcopy(leitura)
    copia.linhas[1].pendencias = ["p" * 20000, "q" * 20000]
    try:
        gerar_relatorio_base_fiscal(copia, nome_arquivo="base.xlsx")
    except ErroClientes as exc:
        assert "1:4" in exc.mensagem and "Pendências" in exc.mensagem
    else:
        raise AssertionError("Concatenação não pode ultrapassar limite de célula")
    copia.linhas[1].pendencias = []
    copia.linhas[0].nome = "z" * 32767
    wb = load_workbook(BytesIO(gerar_relatorio_base_fiscal(copia, nome_arquivo="base.xlsx")))
    try:
        assert wb["Conferência"]["D2"].value == "z" * 32767
    finally:
        wb.close()


def test_vazio_e_duplicidades_lineares():
    vazio = LeituraEmpresas("a" * 64, [dict(nome="Carteira", tipo="principal", linha_cabecalho=1, total=0, oculta=False)], [], [])
    wb = load_workbook(BytesIO(gerar_relatorio_base_fiscal(vazio, nome_arquivo="vazia.xlsx")))
    try:
        for nome in ("Conferência", "Duplicidades", "Correções", "Origens"):
            assert wb[nome].max_row == 1
    finally:
        wb.close()
    documento = "012345670001" + dv_cnpj("012345670001")
    repetidos = [LinhaEmpresa(id=f"1:{i+2}", aba="Carteira", linha=i+2, nome=f"Empresa sintética {i}",
                             documento=documento, documento_status="valido", duplicado=True) for i in range(316)]
    for linha in repetidos:
        linha.duplicado_em = [outro.id for outro in repetidos if outro.id != linha.id]
    vazio.linhas = repetidos
    vazio.abas[0]["total"] = len(repetidos)
    wb = load_workbook(BytesIO(gerar_relatorio_base_fiscal(vazio, nome_arquivo="duplicadas.xlsx")))
    try:
        assert wb["Duplicidades"].max_row == 317
        assert wb["Conferência"].max_row == wb["Correções"].max_row == 317
        assert {r["ID de origem"] for r in tabelar(wb["Duplicidades"])} == {l.id for l in repetidos}
    finally:
        wb.close()


def test_limpeza_em_falha_de_escrita(pasta):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    from openpyxl.worksheet._write_only import WriteOnlyWorksheet
    from openpyxl.writer.excel import ExcelWriter
    leitura, _ = fonte_sintetica(pasta)
    antes = set(ALL_TEMP_FILES)
    with patch.object(ExcelWriter, "write_data", side_effect=OSError("Falha sintética de escrita")):
        try:
            gerar_relatorio_base_fiscal(leitura, nome_arquivo="base.xlsx")
        except OSError as exc:
            assert "Falha sintética" in str(exc)
        else:
            raise AssertionError("Falha de escrita não pode devolver conteúdo parcial")
    assert set(ALL_TEMP_FILES) == antes, "XML temporário permaneceu depois da falha"
    append = WriteOnlyWorksheet.append
    chamadas = 0
    def falhar_no_corpo(aba, linha):
        nonlocal chamadas
        chamadas += 1
        if chamadas == 2:
            raise OSError("Falha sintética durante as linhas")
        return append(aba, linha)
    with patch.object(WriteOnlyWorksheet, "append", falhar_no_corpo):
        try:
            gerar_relatorio_base_fiscal(leitura, nome_arquivo="base.xlsx")
        except OSError as exc:
            assert "durante as linhas" in str(exc)
        else:
            raise AssertionError("Falha durante as linhas não pode gerar relatório parcial")
    assert set(ALL_TEMP_FILES) == antes, "XML temporário permaneceu depois da falha de append"


def test_limpeza_quando_fechamento_falha(pasta):
    """Falha de rodapé não pode impedir a limpeza das cinco abas abertas."""
    from openpyxl.worksheet._writer import ALL_TEMP_FILES, WorksheetWriter
    from auditoria_fiscal.ferramentas import relatorio_base_fiscal as rb
    leitura, _ = fonte_sintetica(pasta)
    antes = set(ALL_TEMP_FILES)
    livros, pacotes = [], []
    fechar_writer = WorksheetWriter.close
    fechar_chamadas = 0
    def criar_livro(*args, **kwargs):
        wb = Workbook(*args, **kwargs)
        livros.append(wb)
        return wb
    def criar_pacote(*args, **kwargs):
        pacote = zipfile.ZipFile(*args, **kwargs)
        pacotes.append(pacote)
        return pacote
    def fechar_com_erro_secundario(writer):
        nonlocal fechar_chamadas
        fechar_chamadas += 1
        fechar_writer(writer)
        if fechar_chamadas == 1:
            raise OSError("Erro secundário no fechamento")
    try:
        with patch.object(rb, "Workbook", criar_livro), patch.object(rb, "ZipFile", criar_pacote), patch.object(
                WorksheetWriter, "close", fechar_com_erro_secundario), patch.object(
                WorksheetWriter, "write_tail", side_effect=OSError("Disco cheio sintético")):
            try:
                gerar_relatorio_base_fiscal(leitura, nome_arquivo="base.xlsx")
            except OSError as exc:
                assert str(exc) == "Disco cheio sintético", "A limpeza mascarou a falha original"
            else:
                raise AssertionError("Falha no rodapé não pode devolver relatório")
        residuais = set(ALL_TEMP_FILES) - antes
        assert not residuais, f"Falha no rodapé deixou {len(residuais)} XMLs temporários"
        assert fechar_chamadas == 5, "Erro secundário interrompeu a limpeza das demais abas"
        assert all(not Path(aba._writer.out).exists() for wb in livros for aba in wb.worksheets)
        assert pacotes and all(p.fp is None for p in pacotes), "ZIP não foi fechado após falha"
    finally:
        # A primeira execução RED também precisa deixar o ambiente limpo.
        for wb in livros:
            for aba in wb.worksheets:
                if aba._rows is not None:
                    with suppress(Exception):
                        aba._rows.close()
                if aba._writer is not None:
                    with suppress(Exception):
                        aba._writer.close()
                    with suppress(Exception):
                        aba._writer.cleanup()
            wb.close()


def main():
    with tempfile.TemporaryDirectory(prefix="relatorio_base_fiscal_sintetico_") as pasta:
        test_fidelidade_e_propostas(pasta)
        test_representacao_sem_truncar(pasta)
        test_vazio_e_duplicidades_lineares()
        test_limpeza_em_falha_de_escrita(pasta)
        test_limpeza_quando_fechamento_falha(pasta)
    print("OK - relatório Base Fiscal: fidelidade, propostas vazias, origem, texto literal e limites")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
