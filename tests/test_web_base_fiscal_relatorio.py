"""Relatorio Fiscal: acesso cumulativo, fonte verificada e nenhuma carga cadastral."""
from pathlib import Path
import gc
import hashlib
import io
import os
import sys
import tempfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
_DADOS = tempfile.TemporaryDirectory(prefix='web_relatorio_fiscal_sintetico_')
os.environ['AUDITORIA_WEB_DADOS'] = _DADOS.name

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from auditoria_fiscal.core.documentos_br import dv_cnpj
from auditoria_fiscal.web.servidor import criar_app
from auditoria_fiscal.web import rotas_clientes as rc


def fonte_sintetica():
    livro = Workbook()
    aba = livro.active
    aba.title = 'Fiscal sintético'
    aba.append([None, 'RESPONSAVEL', 'EMPRESAS', 'RESPONSAVEL', 'CNPJ', 'REGIME DA EMPRESA'])
    cnpj = '019876540001' + dv_cnpj('019876540001')
    aba.append([None, 'Responsável sintético', 'Empresa sintética HTTP', 'Contato sintético', cnpj, 'Lucro Real'])
    aba.row_dimensions[2].hidden = True
    zeradas = livro.create_sheet('ZERADAS')
    zeradas.append([None, 'EMPRESA', 'CNPJ'])
    zeradas.append([None, 'Duplicada sintética', cnpj])
    zeradas.sheet_state = 'hidden'
    destino = io.BytesIO()
    livro.save(destino)
    livro.close()
    return destino.getvalue(), cnpj


def main():
    app = criar_app()
    adm = TestClient(app)
    assert adm.post('/api/bootstrap', json=dict(usuario='dono_relatorio', nome='Perfil sintético',
                    senha='teste-relatorio-019')).status_code == 200
    rota = '/api/clientes/base-fiscal/relatorio'
    conteudo, cnpj = fonte_sintetica()
    arquivos = {'arquivo': ('fonte-sintetica.xlsx', conteudo, 'application/octet-stream')}
    esperado = hashlib.sha256(conteudo).hexdigest()
    dados = {'hash_esperado': esperado}
    # Antes da implementacao este contrato falha: endpoint inexistente (404).
    assert TestClient(app).post(rota, files=arquivos, data=dados).status_code == 401

    perfis = {}
    todas = ['aba.clientes', 'clientes.importar', 'clientes.exportar']
    for i, permissoes in enumerate([todas[1:], todas[::2], todas[:2], todas]):
        nome = f'revisor_fiscal_{i}'
        resposta = adm.post('/api/admin/usuarios', json=dict(usuario=nome, nome='Perfil sintético',
                           senha='teste-relatorio-019', permissoes=permissoes))
        assert resposta.status_code == 200, resposta.text
        cliente = TestClient(app)
        assert cliente.post('/api/login', json=dict(usuario=nome, senha='teste-relatorio-019')).status_code == 200
        perfis[i] = cliente

    originais = rc.pe.ler_planilha_empresas, rc._receber, rc._store
    chamadas, caminhos, uploads = [], [], []
    def ler(*args, **kwargs):
        chamadas.append(True)
        return originais[0](*args, **kwargs)
    async def receber(arquivo, destino, limite):
        caminhos.append(Path(destino))
        uploads.append(arquivo)
        return await originais[1](arquivo, destino, limite)
    def store_proibido(*args, **kwargs):
        raise AssertionError('Relatorio de origem nao deve abrir store cadastral')
    rc.pe.ler_planilha_empresas, rc._receber, rc._store = ler, receber, store_proibido
    previas = dict(rc._previas)
    try:
        for i in range(3):
            negado = perfis[i].post(rota, files=arquivos, data=dados)
            assert negado.status_code == 403, negado.text
        assert not chamadas and not caminhos, 'Nao analisar a fonte sem todas as permissoes'

        previa = perfis[3].post('/api/clientes/base-fiscal/previa', files=arquivos).json()
        assert previa['hash_arquivo'] == esperado
        for cliente in (perfis[3], adm):
            resposta = cliente.post(rota, files=arquivos, data=dados)
            assert resposta.status_code == 200, resposta.text
            assert resposta.headers['content-type'].split(';')[0] == rc.MIME_XLSX
            assert resposta.headers['cache-control'] == 'no-store'
            assert resposta.headers['content-disposition'] == 'attachment; filename="conferencia-base-fiscal.xlsx"'
            with io.BytesIO(resposta.content) as arquivo:
                livro = load_workbook(arquivo)
                assert livro.sheetnames == ['Resumo', 'Conferência', 'Duplicidades', 'Correções', 'Origens']
                conferencia = livro['Conferência']
                assert conferencia.max_row == previa['contagens']['linhas'] + 1
                assert [conferencia.cell(i, 1).value for i in range(2, 4)] == [l['id'] for l in previa['linhas']]
                assert livro['Duplicidades'].max_row == 3
                textos = [c for aba in livro for linha in aba for c in linha]
                assert any(c.value == esperado for c in textos)
                assert all(c.data_type != 'f' for c in textos)
                assert all(c.data_type == 's' for c in textos if c.value == cnpj)
                livro.close()

        n = len(chamadas)
        for invalido in ({}, {'hash_esperado': 'a'}, {'hash_esperado': 'G'*64},
                         {'hash_esperado': 'A'*64}, {'hash_esperado': 'a'*65}):
            resposta = adm.post(rota, files=arquivos, data=invalido)
            assert resposta.status_code == 422, resposta.text
        assert len(chamadas) == n, 'Hash malformado deve ser recusado antes do parser'

        resposta = adm.post(rota, files=arquivos, data={'hash_esperado': '0'*64})
        assert resposta.status_code == 409, resposta.text
        assert resposta.json()['detail']['code'] == 'BASE_FISCAL_DIVERGENTE'
        assert esperado not in resposta.text and 'Empresa sintética' not in resposta.text
        for nome, corpo in [('base.xls', conteudo), ('base.xlsx', b'nao e xlsx')]:
            resposta = adm.post(rota, files={'arquivo': (nome, corpo)}, data=dados)
            assert resposta.status_code == 422, resposta.text
        limite = rc.pe.MAX_PLANILHA_BYTES
        rc.pe.MAX_PLANILHA_BYTES = 32
        try:
            assert adm.post(rota, files=arquivos, data=dados).status_code == 413
        finally:
            rc.pe.MAX_PLANILHA_BYTES = limite

        # Falhas de leitura e de geracao nao podem deixar upload nem criar previa.
        def falhar(*args, **kwargs):
            raise RuntimeError('Falha sintética')
        cliente_falha = TestClient(app, raise_server_exceptions=False)
        cliente_falha.cookies.update(adm.cookies)
        rc.pe.ler_planilha_empresas = falhar
        assert cliente_falha.post(rota, files=arquivos, data=dados).status_code == 500
        rc.pe.ler_planilha_empresas = ler
        gerar = rc.gerar_relatorio_base_fiscal
        try:
            rc.gerar_relatorio_base_fiscal = falhar
            assert cliente_falha.post(rota, files=arquivos, data=dados).status_code == 500
            def irrepresentavel(*args, **kwargs):
                raise rc.cl.ErroClientes('PLANILHA_INVALIDA', 'Campo excede o limite de célula.')
            rc.gerar_relatorio_base_fiscal = irrepresentavel
            resposta = adm.post(rota, files=arquivos, data=dados)
            assert resposta.status_code == 422, resposta.text
        finally:
            rc.gerar_relatorio_base_fiscal = gerar
        assert all(a.file.closed for a in uploads)
        assert not any(p.exists() or p.parent.exists() for p in caminhos)
        assert rc._previas == previas
        assert not (Path(_DADOS.name) / 'clientes.db').exists()
    finally:
        rc.pe.ler_planilha_empresas, rc._receber, rc._store = originais

    eventos = adm.get('/api/admin/historico', params={'limite': 500}).json()['itens']
    exportacoes = [e for e in eventos if e['acao'] == 'clientes.exportar']
    assert {200, 403, 409, 413, 422, 500}.issubset({e['http_status'] for e in exportacoes})
    assert any(e['detalhe'] == 'relatorio base fiscal: 2 linha(s)' for e in exportacoes)
    for segredo in (cnpj, esperado, 'Empresa sintética HTTP', 'Responsável sintético', 'fonte-sintetica.xlsx'):
        assert segredo not in str(exportacoes), 'Conteudo da fonte nao pertence a auditoria geral'
    for cliente in [adm, cliente_falha, *perfis.values()]:
        cliente.close()
    # O contexto transacional sqlite3 não fecha a conexão; a coleta libera
    # as conexões de autenticação do legado antes de remover a fixture Windows.
    gc.collect()
    _DADOS.cleanup()
    print('OK - relatorio Fiscal HTTP: tres permissoes, hash, XLSX completo, limpeza, sem cadastro e auditoria minima')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
