"""Confirmação Fiscal: autorização, sessão, arquivos, concorrência e retry local."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import timedelta
from pathlib import Path
import gc
import hashlib
import io
import os
import socket
import sqlite3
import sys
import tempfile
import threading

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
_DADOS = tempfile.TemporaryDirectory(prefix='web_confirmacao_fiscal_sintetica_')
os.environ['AUDITORIA_WEB_DADOS'] = _DADOS.name

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from auditoria_fiscal.core.documentos_br import dv_cnpj
from auditoria_fiscal.core.planilha_empresas import ler_planilha_empresas
from auditoria_fiscal.ferramentas.relatorio_base_fiscal import gerar_relatorio_base_fiscal
from auditoria_fiscal.web.servidor import criar_app
from auditoria_fiscal.web import rotas_clientes as rc

PREPARAR = '/api/clientes/base-fiscal/importacao/preparar'
PREFIXO = '/api/clientes/base-fiscal/importacao/'
SLUG = 'clientes.importar_base_fiscal'
PERMISSOES = ['aba.clientes', 'clientes.importar', SLUG]


def fonte_relatorio(numero=1):
    livro = Workbook()
    aba = livro.active
    aba.title = 'Fiscal sintético'
    aba.append([None, 'RESPONSAVEL', 'EMPRESAS', 'RESPONSAVEL', 'CNPJ', 'REGIME DA EMPRESA'])
    base = f'01987654{numero:04d}'
    cnpj = base + dv_cnpj(base)
    aba.append([None, 'Equipe sintética', f'Empresa sintética carga {numero}',
                'Contato sintético', cnpj, 'Lucro Real'])
    saida = io.BytesIO()
    livro.save(saida)
    livro.close()
    fonte = saida.getvalue()
    with tempfile.TemporaryDirectory() as pasta:
        p = Path(pasta) / 'fonte.xlsx'
        p.write_bytes(fonte)
        relatorio = gerar_relatorio_base_fiscal(ler_planilha_empresas(p), nome_arquivo='fonte.xlsx')
    return fonte, relatorio, cnpj


def editar(conteudo, celula, valor):
    livro = load_workbook(io.BytesIO(conteudo))
    livro['Correções'][celula] = valor
    saida = io.BytesIO()
    livro.save(saida)
    livro.close()
    return saida.getvalue()


def main():
    app = criar_app()
    clientes = []
    def cliente():
        c = TestClient(app)
        clientes.append(c)
        return c
    adm = cliente()
    assert adm.post('/api/bootstrap', json=dict(usuario='dono_carga', nome='Perfil sintético',
                    senha='teste-carga-021')).status_code == 200
    fonte, revisao, cnpj = fonte_relatorio()
    def arquivos(f=fonte, r=revisao):
        return {'arquivo': ('fonte-sintetica.xlsx', f), 'correcoes': ('revisao-sintetica.xlsx', r)}
    def hashes(f=fonte, r=revisao):
        return {'hash_esperado': hashlib.sha256(f).hexdigest(),
                'hash_correcoes_esperado': hashlib.sha256(r).hexdigest()}
    def preparar(c=adm, f=fonte, r=revisao):
        resposta = c.post(PREPARAR, files=arquivos(f, r), data=hashes(f, r))
        assert resposta.status_code == 200, resposta.text
        assert resposta.headers['cache-control'] == 'no-store'
        p = resposta.json()
        assert p['formato'] == 'base_fiscal_confirmacao_v1' and p['expira_em_minutos'] == 30
        assert p['previa']['somente_leitura'] is True and p['previa']['contagens']['incluir'] == 1
        return p['token']
    def confirmar(token, c=adm, f=fonte, r=revisao, confirmacao='IMPORTAR EMPRESAS'):
        return c.post(PREFIXO + token + '/confirmar', files=arquivos(f, r),
                      data={'confirmacao': confirmacao})

    # RED: as três rotas ainda não existem; autenticação deve anteceder leitura/carga.
    anonimo = cliente()
    negada = anonimo.post(PREPARAR, files=arquivos(), data=hashes())
    assert negada.status_code == 401 and negada.headers['cache-control'] == 'no-store'
    assert confirmar('inexistente', anonimo).status_code == 401
    assert anonimo.delete(PREFIXO + 'inexistente').status_code == 401
    catalogo = adm.get('/api/admin/usuarios').json()
    assert SLUG in {i['slug'] for g in catalogo['catalogo'] for i in g['itens']}
    assert SLUG not in catalogo['padrao_novo']
    perfis = []
    ids = []
    for i, permissoes in enumerate([PERMISSOES[:2], PERMISSOES[1:],
                                    [PERMISSOES[0], SLUG], PERMISSOES]):
        login = f'operador_carga_{i}'
        r = adm.post('/api/admin/usuarios', json=dict(usuario=login, nome='Perfil sintético',
                     senha='teste-carga-021', permissoes=permissoes))
        assert r.status_code == 200, r.text
        ids.append(r.json()['id'])
        c = cliente()
        assert c.post('/api/login', json=dict(usuario=login, senha='teste-carga-021')).status_code == 200
        perfis.append(c)

    originais = rc._receber, rc.cs.ClientesStore, rc._agora, socket.create_connection
    caminhos, uploads, stores = [], [], []
    async def receber(arquivo, destino, limite):
        caminhos.append(Path(destino)); uploads.append(arquivo)
        return await originais[0](arquivo, destino, limite)
    def abrir_store(*args, **kwargs):
        stores.append(True)
        return originais[1](*args, **kwargs)
    def sem_rede(*args, **kwargs):
        raise AssertionError('A importação local não deve abrir conexão externa')
    rc._receber, rc.cs.ClientesStore, socket.create_connection = receber, abrir_store, sem_rede
    banco = Path(_DADOS.name) / 'clientes.db'
    try:
        for c in perfis[:3]:
            assert c.post(PREPARAR, files=arquivos(), data=hashes()).status_code == 403
            assert confirmar('inexistente', c).status_code == 403
            assert c.delete(PREFIXO + 'inexistente').status_code == 403
        assert not caminhos and not stores

        token = preparar(perfis[3])
        assert not banco.exists() and not stores
        metadados = vars(rc._preparacoes_base_fiscal[token])
        assert not any(isinstance(v, (bytes, dict, list)) for v in metadados.values())
        assert cnpj not in str(metadados) and 'Empresa sintética' not in str(metadados)
        assert confirmar(token).status_code == 404, 'Outro usuário não pode usar token'
        outra_sessao = cliente()
        assert outra_sessao.post('/api/login', json=dict(usuario='operador_carga_3', senha='teste-carga-021')).status_code == 200
        assert confirmar(token, outra_sessao).status_code == 404
        assert outra_sessao.delete(PREFIXO + token).status_code == 404
        for texto in ('', 'true', 'IMPORTAR', 'importar empresas'):
            assert confirmar(token, perfis[3], confirmacao=texto).status_code == 422
        assert not banco.exists() and not stores

        # Permissão revogada após a preparação não pode confirmar.
        assert adm.put(f'/api/admin/usuarios/{ids[3]}/permissoes', json={'permissoes': PERMISSOES[:2]}).status_code == 200
        assert confirmar(token, perfis[3]).status_code == 403
        assert adm.put(f'/api/admin/usuarios/{ids[3]}/permissoes', json={'permissoes': PERMISSOES}).status_code == 200
        assert perfis[3].delete(PREFIXO + token).status_code == 200
        assert confirmar(token, perfis[3]).status_code == 410

        token = preparar()
        rc._agora = lambda: originais[2]() + timedelta(minutes=31)
        assert confirmar(token).status_code == 410
        rc._agora = originais[2]
        assert not banco.exists() and not stores

        # Hashes, pacote, propostas e confirmação inválidos não inicializam o store.
        token = preparar()
        assert confirmar(token, r=editar(revisao, 'K2', 'Outro contato')).status_code == 409
        assert not stores
        for dados in ({}, {'hash_esperado': 'a'}, {**hashes(), 'hash_correcoes_esperado': 'g'*64}):
            invalida = adm.post(PREPARAR, files=arquivos(), data=dados)
            assert invalida.status_code == 422 and invalida.headers['cache-control'] == 'no-store'
        assert adm.post(PREPARAR, files=arquivos(), data={**hashes(), 'hash_correcoes_esperado': '0'*64}).status_code == 409
        pendente = editar(revisao, 'K2', 'x'*161)
        assert adm.post(PREPARAR, files=arquivos(r=pendente), data=hashes(r=pendente)).status_code == 422
        assert not stores and not banco.exists()
        limite = rc.pe.MAX_PLANILHA_BYTES
        rc.pe.MAX_PLANILHA_BYTES = 32
        try:
            assert adm.post(PREPARAR, files=arquivos(), data=hashes()).status_code == 413
            token_size = token
            assert confirmar(token_size).status_code in (410, 413)
        finally:
            rc.pe.MAX_PLANILHA_BYTES = limite

        # Limite de sessões de preparação e descarte são somente memória.
        tokens = [preparar() for _ in range(4)]
        assert confirmar(tokens[0]).status_code == 410
        assert not stores and not banco.exists()

        # Descarte ou revogação durante a leitura impede abertura do banco de escrita.
        original_leitura = rc._ler_confirmacao_base_fiscal
        def descartar_apos_leitura(*args):
            previa = original_leitura(*args)
            rc._descartar_preparacao_base_fiscal(tokens[-1])
            return previa
        rc._ler_confirmacao_base_fiscal = descartar_apos_leitura
        try:
            assert confirmar(tokens[-1]).status_code == 410
        finally:
            rc._ler_confirmacao_base_fiscal = original_leitura
        token_usuario = preparar(perfis[3])
        def revogar_apos_leitura(*args):
            previa = original_leitura(*args)
            rc.permissoes.definir(ids[3], PERMISSOES[:2], por='fixture')
            return previa
        rc._ler_confirmacao_base_fiscal = revogar_apos_leitura
        try:
            assert confirmar(token_usuario, perfis[3]).status_code == 403
        finally:
            rc._ler_confirmacao_base_fiscal = original_leitura
            rc.permissoes.definir(ids[3], PERMISSOES, por='fixture')
        assert not stores and not banco.exists()

        # Também há limite global, além do limite de três por usuário.
        limite_global = rc.PREPARACOES_BASE_FISCAL_MAXIMAS
        rc.PREPARACOES_BASE_FISCAL_MAXIMAS = 2
        try:
            preparar(); preparar(perfis[3]); preparar()
            assert len(rc._preparacoes_base_fiscal) == 2
        finally:
            rc.PREPARACOES_BASE_FISCAL_MAXIMAS = limite_global

        # Cadastro que muda entre preparação e confirmação exige nova revisão.
        token = preparar()
        store = originais[1](str(banco))
        try:
            criado = store.criar_cliente({'nome': 'Outra empresa sintética', 'documento': cnpj},
                                          usuario_id=1, usuario_login='fixture')
        finally:
            store.fechar()
        assert confirmar(token).status_code == 409
        assert confirmar(token).status_code == 410
        with closing(sqlite3.connect(banco)) as conn:
            assert conn.execute('SELECT count(*) FROM cliente').fetchone()[0] == 1
        assert adm.post(PREPARAR, files=arquivos(), data=hashes()).status_code == 422

        # Nova empresa: duas confirmações concorrentes devem retornar o mesmo recibo.
        f2, r2, doc2 = fonte_relatorio(2)
        token2 = preparar(f=f2, r=r2)
        def confirmar_copia(_):
            c = TestClient(app)
            c.cookies.update(adm.cookies)
            try:
                return confirmar(token2, c, f2, r2)
            finally:
                c.close()
        with ThreadPoolExecutor(max_workers=2) as pool:
            respostas = list(pool.map(confirmar_copia, range(2)))
        for r in respostas:
            assert r.status_code == 200, r.text
            assert r.headers['cache-control'] == 'no-store'
            assert r.json()['formato'] == 'base_fiscal_resultado_v1'
            assert r.json()['contagens'] == {'incluidas': 1, 'existentes': 0, 'ignoradas': 0}
        assert respostas[0].json()['importacao_id'] == respostas[1].json()['importacao_id']
        assert sorted(r.json()['repetida'] for r in respostas) == [False, True]
        with closing(sqlite3.connect(banco)) as conn:
            assert conn.execute('SELECT count(*) FROM cliente').fetchone()[0] == 2
            assert conn.execute('SELECT count(*) FROM segredo').fetchone()[0] == 0
            assert conn.execute('SELECT count(*) FROM certificado').fetchone()[0] == 0
            antes = list(conn.iterdump())
        repetida = confirmar(token2, f=f2, r=r2)
        assert repetida.status_code == 200 and repetida.json()['repetida'] is True
        with closing(sqlite3.connect(banco)) as conn:
            assert list(conn.iterdump()) == antes

        # A autoridade deve ser conferida já dentro da transação: abrir/migrar
        # o store e esperar o escritor SQLite também consomem tempo.
        f4, r4, _ = fonte_relatorio(4)
        token4 = preparar(perfis[3], f4, r4)
        def revogar_no_construtor(*args, **kwargs):
            store = abrir_store(*args, **kwargs)
            rc.permissoes.definir(ids[3], PERMISSOES[:2], por='fixture')
            return store
        rc.cs.ClientesStore = revogar_no_construtor
        try:
            assert confirmar(token4, perfis[3], f4, r4).status_code == 403
        finally:
            rc.cs.ClientesStore = abrir_store
            rc.permissoes.definir(ids[3], PERMISSOES, por='fixture')

        token4 = preparar(f=f4, r=r4)
        def descartar_no_construtor(*args, **kwargs):
            store = abrir_store(*args, **kwargs)
            rc._descartar_preparacao_base_fiscal(token4)
            return store
        rc.cs.ClientesStore = descartar_no_construtor
        try:
            assert confirmar(token4, f=f4, r=r4).status_code == 410
        finally:
            rc.cs.ClientesStore = abrir_store

        token4 = preparar(f=f4, r=r4)
        tentou_begin = threading.Event()
        liberou_escritor = threading.Event()
        threads = []
        def expirar_aguardando_transacao(*args, **kwargs):
            store = abrir_store(*args, **kwargs)
            escritor = sqlite3.connect(str(banco), isolation_level=None, check_same_thread=False)
            escritor.execute('BEGIN IMMEDIATE')
            def liberar():
                try:
                    if tentou_begin.wait(5):
                        rc._agora = lambda: originais[2]() + timedelta(minutes=31)
                    escritor.execute('COMMIT')
                finally:
                    escritor.close()
                    liberou_escritor.set()
            thread = threading.Thread(target=liberar)
            threads.append(thread)
            store._conn.set_trace_callback(
                lambda sql: tentou_begin.set() if sql == 'BEGIN IMMEDIATE' else None)
            thread.start()
            return store
        rc.cs.ClientesStore = expirar_aguardando_transacao
        try:
            assert confirmar(token4, f=f4, r=r4).status_code == 410
            assert tentou_begin.is_set() and liberou_escritor.is_set()
        finally:
            for thread in threads:
                thread.join(timeout=10)
            rc.cs.ClientesStore = abrir_store
            rc._agora = originais[2]
        with closing(sqlite3.connect(banco)) as conn:
            assert list(conn.iterdump()) == antes, 'Autorização expirada não grava cadastro, recibo ou trilha de domínio'

        # Encerrar a sessão elimina a autoridade do token, mesmo com arquivos iguais.
        f3, r3, _ = fonte_relatorio(3)
        token3 = preparar(perfis[3], f3, r3)
        assert perfis[3].post('/api/logout').status_code == 200
        assert confirmar(token3, perfis[3], f3, r3).status_code == 401

        caminho_original = rc.caminho_db_clientes
        quebrado = Path(_DADOS.name) / 'cadastro-invalido.db'
        quebrado.write_bytes(b'cadastro sintetico invalido')
        rc.caminho_db_clientes = lambda: str(quebrado)
        try:
            indisponivel = adm.post(PREPARAR, files=arquivos(f3, r3), data=hashes(f3, r3))
            assert indisponivel.status_code == 503 and indisponivel.headers['cache-control'] == 'no-store'
            assert confirmar(token2, f=f2, r=r2).status_code == 404, 'Token preso ao banco'
        finally:
            rc.caminho_db_clientes = caminho_original
        simular_original = rc._simular_base_fiscal
        def falha_interna(*args):
            raise RuntimeError('Detalhe privado sintético que não deve vazar')
        rc._simular_base_fiscal = falha_interna
        try:
            interna = adm.post(PREPARAR, files=arquivos(f3, r3), data=hashes(f3, r3))
            assert interna.status_code == 500 and interna.headers['cache-control'] == 'no-store'
            assert 'Detalhe privado' not in interna.text
        finally:
            rc._simular_base_fiscal = simular_original
        assert all(a.file.closed for a in uploads)
        assert not any(p.exists() or p.parent.exists() for p in caminhos)
        eventos = adm.get('/api/admin/historico', params={'limite': 500}).json()['itens']
        relevantes = [e for e in eventos if e['acao'] == 'clientes.importar']
        assert any('confirmacao base fiscal' in e['detalhe'] for e in relevantes)
        for privado in (cnpj, doc2, hashes()['hash_esperado'], 'Empresa sintética', 'revisao-sintetica.xlsx', token2):
            assert privado not in str(relevantes)
    finally:
        rc._receber, rc.cs.ClientesStore, rc._agora, socket.create_connection = originais
        for c in clientes:
            c.close()
        gc.collect()
        _DADOS.cleanup()
    print('OK - carga Fiscal HTTP: autorização, sessão, revalidação, transação e retry sintéticos')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
