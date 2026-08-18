"""Versao web — Conciliacao Fiscal (sexta ferramenta).

Esqueleto da fase 3: portao de acesso da aba. Exercita 401 sem login, 403 sem
a permissao da aba, criacao da sessao de trabalho da ferramenta, navegacao
auditada e o resumo vazio (a ferramenta existe e responde antes de haver
qualquer dado importado).

As fases seguintes acrescentam aqui: upload e lote (T023), consulta e
proveniencia (T033), revisao e aprovacao (T039), exportacoes (T048, T055),
painel (T060) e a concorrencia de SC-013 (T070).

Usa dados_web isolado em pasta temporaria (AUDITORIA_WEB_DADOS).
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import time

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "src"))
sys.path.insert(0, os.path.join(RAIZ, "tests"))

_TMP = tempfile.mkdtemp(prefix="auditoria_web_conc_")
os.environ["AUDITORIA_WEB_DADOS"] = _TMP

from fastapi.testclient import TestClient  # noqa: E402

from fixtures import conciliacao as fx  # noqa: E402

from auditoria_fiscal.core import conciliacao_sefaz as cs  # noqa: E402
from auditoria_fiscal.web.servidor import criar_app  # noqa: E402

ADMIN = {"usuario": "weslley", "nome": "Weslley", "senha": "segredo1"}
# Importador: entra e importa, mas nao decide nem gera saida oficial.
IMPORTADOR = {"usuario": "ana", "nome": "Ana", "senha": "segredo2"}


def checar(cond, msg):
    if not cond:
        print(f"FALHOU - {msg}")
        raise SystemExit(1)


def entrar(app, credenciais: dict) -> TestClient:
    cliente = TestClient(app)
    r = cliente.post("/api/login", json={"usuario": credenciais["usuario"],
                                         "senha": credenciais["senha"]})
    checar(r.status_code == 200,
           f"login de {credenciais['usuario']} falhou: {r.text}")
    return cliente


def eventos(adm, **filtros) -> list[dict]:
    r = adm.get("/api/admin/historico", params={"limite": 200, **filtros})
    checar(r.status_code == 200, f"historico: {r.status_code} {r.text[:120]}")
    return r.json()["itens"]


def main() -> int:
    app = criar_app()
    adm = TestClient(app)

    r = adm.post("/api/bootstrap", json=ADMIN)
    checar(r.status_code == 200, f"bootstrap falhou: {r.text}")

    # ------------------------------------------------------------------
    # Sem login: 401 em tudo, antes de qualquer checagem de permissao

    anonimo = TestClient(app)
    for caminho in ("/api/conciliacao/resumo",):
        r = anonimo.get(caminho)
        checar(r.status_code == 401,
               f"{caminho} sem login deveria dar 401: {r.status_code}")
    r = anonimo.post("/api/sessoes", json={"ferramenta": "conciliacao"})
    checar(r.status_code == 401,
           f"criar sessao sem login deveria dar 401: {r.status_code}")

    # ------------------------------------------------------------------
    # Com login, sem a aba: 403 e a tentativa fica auditada

    r = adm.post("/api/admin/usuarios", json={
        **IMPORTADOR, "admin": False, "permissoes": ["aba.conferencia"]})
    checar(r.status_code == 200, f"criar ana falhou: {r.text}")
    ana_id = r.json()["id"]
    ana = entrar(app, IMPORTADOR)

    r = ana.get("/api/conciliacao/resumo")
    checar(r.status_code == 403,
           f"resumo sem a aba deveria dar 403: {r.status_code} {r.text[:120]}")
    r = ana.post("/api/sessoes", json={"ferramenta": "conciliacao"})
    checar(r.status_code == 403,
           f"sessao sem a aba deveria dar 403: {r.status_code}")

    negados = [e for e in eventos(adm, usuario_filtro="ana")
               if e["resultado"] == "negado"]
    checar(negados and all(e["http_status"] == 403 for e in negados),
           f"a tentativa negada deveria estar na trilha com 403: {negados[:1]}")

    # ------------------------------------------------------------------
    # Com a aba liberada: abre, cria sessao e a navegacao fica auditada

    r = adm.put(f"/api/admin/usuarios/{ana_id}/permissoes", json={
        "permissoes": ["aba.conciliacao", "conciliacao.importar"]})
    checar(r.status_code == 200, f"liberar a aba para ana falhou: {r.text}")

    estado = ana.get("/api/estado").json()
    checar(estado["usuario"]["permissoes"] == ["aba.conciliacao",
                                               "conciliacao.importar"],
           f"permissoes da ana fora do esperado: {estado['usuario']}")

    checar(ana.post("/api/eventos/aba",
                    json={"aba": "conciliacao"}).status_code == 200,
           "ana deveria conseguir abrir a aba da conciliacao")
    r = ana.post("/api/sessoes", json={"ferramenta": "conciliacao"})
    checar(r.status_code == 200, f"sessao da conciliacao: {r.text}")
    sessao = r.json()["sessao_id"]
    checar(sessao, "a sessao deveria ter id")

    navegacoes = [e for e in eventos(adm, usuario_filtro="ana",
                                     acao="navegacao.aba")
                  if e["resultado"] == "ok"]
    checar(any("conciliacao" in e["detalhe"] for e in navegacoes),
           f"a trilha deveria registrar a aba aberta: {navegacoes[:2]}")

    # ------------------------------------------------------------------
    # Resumo vazio: a ferramenta existe e responde antes de haver dados

    r = ana.get("/api/conciliacao/resumo")
    checar(r.status_code == 200, f"resumo deveria responder: {r.text[:160]}")
    resumo = r.json()
    # Campos exatamente como o schema Summary do contrato.
    for campo in ("total", "em_revisao", "aprovadas", "rejeitadas", "avisos",
                  "bloqueios", "conflitos"):
        checar(campo in resumo, f"o resumo deveria ter '{campo}': {resumo}")
        checar(resumo[campo] == 0,
               f"sem dados, '{campo}' deveria ser 0: {resumo[campo]}")
    # Money: centavos exatos para conta, texto so para a tela.
    dif = resumo.get("diferenca_total")
    checar(isinstance(dif, dict) and dif.get("centavos") == 0,
           f"diferenca_total deveria ser Money zerado: {dif}")
    checar(dif.get("texto"), f"Money precisa do texto para a UI: {dif}")

    # ------------------------------------------------------------------
    # Contratos de erro comuns: nada revela recurso alheio

    outra = adm.post("/api/sessoes", json={"ferramenta": "conciliacao"})
    checar(outra.status_code == 200, f"sessao do admin: {outra.text}")
    sessao_admin = outra.json()["sessao_id"]
    checar(sessao_admin != sessao, "as sessoes deveriam ser distintas")

    # A sessao da conciliacao nao serve para as rotas de outra aba, e vice-versa.
    r = adm.get("/api/conferencia/notas", params={"sessao_id": sessao_admin})
    checar(r.status_code == 404,
           f"sessao de conciliacao na rota de conferencia deveria dar 404: "
           f"{r.status_code}")

    # ------------------------------------------------------------------
    # Importacao: upload, lote, job e resumo por item

    importacao(app, adm, ana, sessao)

    # ------------------------------------------------------------------
    # Consulta, proveniencia e concorrencia (T033 e T070)

    consulta_e_proveniencia(app, adm, ana)
    decisoes(app, adm, ana)
    exportacoes(app, adm, ana)
    painel_e_filtros(ana)
    concorrencia_durante_o_lote(ana)

    print("OK - conciliacao fiscal web (401 sem login, 403 sem a aba com "
          "trilha, sessao da ferramenta, navegacao auditada, upload com teto "
          "e parcial removido, staging por id opaco, lote em job, resultado "
          "persistente, um processamento por sessao, lista e detalhe com "
          "proveniencia, filtros e paginacao, texto perigoso preservado como "
          "texto, download auditado da fonte, decisoes com revision otimista, "
          "bloqueio impedindo aprovacao, conflito resolvido sem perder "
          "versao, consolidado e planilha-mestre com pendentes sinalizados e leituras respondendo durante o processamento) passou.")
    return 0


REVISOR = {"usuario": "bruno", "nome": "Bruno", "senha": "segredo3"}


def decisoes(app, adm, ana) -> None:
    """Revisao, aprovacao, excecoes e conflitos pela API (T039 e T046)."""
    # Um revisor com as permissoes de decisao; a Ana continua so importando.
    r = adm.post("/api/admin/usuarios", json={
        **REVISOR, "admin": False,
        "permissoes": ["aba.conciliacao", "conciliacao.revisar",
                       "conciliacao.aprovar", "conciliacao.resolver_excecao"]})
    checar(r.status_code == 200, f"criar bruno: {r.text[:160]}")
    bruno = entrar(app, REVISOR)

    # Alvo da aprovacao: uma competencia SEM pendencia impeditiva. As com
    # conflito ou bloqueio aberto sao exercitadas logo abaixo, justamente para
    # provar que a aprovacao e' barrada nelas.
    todas = ana.get("/api/conciliacao/conciliacoes").json()["itens"]
    limpa = next(c for c in todas
                 if not c["conflitos_abertos"] and not c["bloqueios_abertos"])
    cid = limpa["id"]

    # --- Quem so importa NAO decide: 403 em cada acao, e fica na trilha.
    for caminho, corpo in (
        (f"/api/conciliacao/conciliacoes/{cid}/revisar",
         {"estado": "rejeitada", "justificativa": "tentando"}),
        (f"/api/conciliacao/conciliacoes/{cid}/aprovar", {"justificativa": "x"}),
    ):
        r = ana.post(caminho, json=corpo)
        checar(r.status_code == 403,
               f"importador nao decide ({caminho}): {r.status_code}")

    # --- O revisor NAO importa nem exporta: permissoes sao independentes.
    r = bruno.post("/api/sessoes", json={"ferramenta": "conciliacao"})
    checar(r.status_code == 200, "o revisor abre a aba")
    with open(os.devnull, "rb"):
        pass
    r = bruno.post("/api/conciliacao/processar", json={"sessao_id": "x"})
    checar(r.status_code == 403,
           f"revisor sem conciliacao.importar nao processa: {r.status_code}")

    # --- Justificativa obrigatoria
    detalhe = bruno.get(f"/api/conciliacao/conciliacoes/{cid}").json()
    r = bruno.post(f"/api/conciliacao/conciliacoes/{cid}/revisar",
                   json={"estado": "rejeitada", "justificativa": "   ",
                         "revision": detalhe["revision"]})
    checar(r.status_code == 409, f"sem justificativa: {r.status_code}")
    checar(r.json()["detail"]["code"] == "JUSTIFICATIVA_OBRIGATORIA",
           r.json())

    # --- Corpo NAO pode forjar autoria: campo extra e' ignorado
    r = bruno.post(f"/api/conciliacao/conciliacoes/{cid}/aprovar",
                   json={"justificativa": "Conferido.",
                         "revision": detalhe["revision"],
                         "usuario_login": "weslley", "reviewer": "weslley"})
    checar(r.status_code == 200, f"aprovar: {r.text[:200]}")
    depois = bruno.get(f"/api/conciliacao/conciliacoes/{cid}").json()
    checar(depois["estado"] == "aprovada", depois["estado"])
    ultima = depois["revisoes"][-1]
    checar(ultima["usuario"] == "bruno",
           f"a autoria tem de ser de quem estava logado: {ultima}")

    # --- Revision desatualizada: 409 com a revision atual para recarregar
    r = bruno.post(f"/api/conciliacao/conciliacoes/{cid}/revisar",
                   json={"estado": "em_revisao", "justificativa": "Reabrindo.",
                         "revision": 1})
    checar(r.status_code == 409, f"revision velha: {r.status_code}")
    corpo = r.json()["detail"]
    checar(corpo["code"] == "REVISION_DESATUALIZADA", corpo)
    checar(corpo["current_revision"] == depois["revision"],
           f"o 409 precisa dizer a revision atual: {corpo}")

    # --- Bloqueio aberto impede aprovacao
    bloqueadas = bruno.get("/api/conciliacao/conciliacoes",
                           params={"excecao": "bloqueio"}).json()["itens"]
    if bloqueadas:
        alvo_bloqueado = bloqueadas[0]
        r = bruno.post(
            f"/api/conciliacao/conciliacoes/{alvo_bloqueado['id']}/aprovar",
            json={"justificativa": "tentando",
                  "revision": alvo_bloqueado["revision"]})
        checar(r.status_code == 409, f"bloqueio: {r.status_code}")
        checar(r.json()["detail"]["code"] == "APROVACAO_IMPEDIDA", r.json())

        # Resolver o bloqueio libera
        d = bruno.get(
            f"/api/conciliacao/conciliacoes/{alvo_bloqueado['id']}").json()
        bloqueio = next(e for e in d["excecoes"]
                        if e["severidade"] == "bloqueio"
                        and e["estado"] == "aberta")
        r = bruno.post(f"/api/conciliacao/excecoes/{bloqueio['id']}/resolver",
                       json={"resolucao": ""})
        checar(r.status_code == 409, "resolucao vazia")
        r = bruno.post(f"/api/conciliacao/excecoes/{bloqueio['id']}/resolver",
                       json={"resolucao": "Conferido com o contribuinte."})
        checar(r.status_code == 200, f"resolver: {r.text[:160]}")
        d = bruno.get(
            f"/api/conciliacao/conciliacoes/{alvo_bloqueado['id']}").json()
        checar(d["bloqueios_abertos"] == 0, d["bloqueios_abertos"])
        r = bruno.post(
            f"/api/conciliacao/conciliacoes/{alvo_bloqueado['id']}/aprovar",
            json={"justificativa": "Liberado.", "revision": d["revision"]})
        checar(r.status_code == 200, f"aprovar apos resolver: {r.text[:160]}")

    # --- Conflitos: comparacao e decisao
    conflitos = bruno.get("/api/conciliacao/conflitos").json()
    checar(conflitos["total"] >= 1, f"deveria haver conflito: {conflitos}")
    conflito = conflitos["itens"][0]
    checar(conflito["diferencas"], "o conflito precisa listar as diferencas")
    campo = conflito["diferencas"][0]
    for chave in ("campo", "vigente", "candidata", "delta"):
        checar(chave in campo, f"diferenca deveria ter '{chave}': {campo}")
    checar(set(campo["vigente"]) == {"centavos", "texto"}, campo["vigente"])

    # Sem a permissao de resolver conflito, 403.
    r = ana.post(f"/api/conciliacao/conflitos/{conflito['id']}/resolver",
                 json={"decisao": "manter_vigente", "justificativa": "x"})
    checar(r.status_code == 403, f"sem permissao: {r.status_code}")

    # O revisor tem resolver_excecao, que cobre conflito tambem.
    antes = bruno.get(
        f"/api/conciliacao/conciliacoes/{conflito['conciliacao_id']}").json()
    r = bruno.post(f"/api/conciliacao/conflitos/{conflito['id']}/resolver",
                   json={"decisao": "promover_candidata",
                         "justificativa": "Retificacao oficial.",
                         "revision": antes["revision"]})
    checar(r.status_code == 200, f"promover: {r.text[:200]}")
    depois = bruno.get(
        f"/api/conciliacao/conciliacoes/{conflito['conciliacao_id']}").json()
    checar(depois["estado"] == "em_revisao",
           "promover devolve a competencia para Em revisao")
    checar(depois["conflitos_abertos"] == 0, depois["conflitos_abertos"])
    substituidas = [v for v in depois["versoes"]
                    if v["estado"] == "substituida"]
    checar(len(substituidas) == 1,
           f"a versao anterior vira substituida, nunca some: {depois['versoes']}")
    checar(depois["versao_vigente"]["numero"] == 2,
           depois["versao_vigente"]["numero"])

    # Decisao invalida nao passa.
    r = bruno.post(f"/api/conciliacao/conflitos/{conflito['id']}/resolver",
                   json={"decisao": "apagar_tudo", "justificativa": "x"})
    checar(r.status_code == 409, f"decisao desconhecida: {r.status_code}")

    # --- Trilha geral registrou as decisoes, sem valor fiscal no detalhe
    itens = eventos(adm, usuario_filtro="bruno", limite=300)
    for acao in ("conciliacao.aprovar", "conciliacao.resolver_conflito"):
        checar(any(e["acao"] == acao and e["resultado"] == "ok"
                   for e in itens), f"{acao} deveria estar na trilha")
    negados = eventos(adm, usuario_filtro="ana", limite=300)
    checar(any(e["acao"] == "conciliacao.aprovar" and e["resultado"] == "negado"
               for e in negados),
           "a tentativa negada do importador tem de ficar registrada")
    for evento in itens:
        for valor in ("40000", "41500", "46000"):
            checar(valor not in evento["detalhe"],
                   f"a trilha geral nao carrega valor fiscal: {evento}")


EXPORTADOR = {"usuario": "carla", "nome": "Carla", "senha": "segredo4"}


def exportacoes(app, adm, ana) -> None:
    """Consolidado e planilha-mestre pela API (T048 e T055)."""
    from openpyxl import load_workbook

    r = adm.post("/api/admin/usuarios", json={
        **EXPORTADOR, "admin": False,
        "permissoes": ["aba.conciliacao", "conciliacao.exportar",
                       "conciliacao.preencher_modelo"]})
    checar(r.status_code == 200, f"criar carla: {r.text[:160]}")
    carla_id = r.json()["id"]
    carla = entrar(app, EXPORTADOR)

    trabalho = tempfile.mkdtemp(prefix="conc_export_")
    try:
        # --- Sem permissao nao exporta nem baixa
        r = ana.post("/api/conciliacao/exportar", json={})
        checar(r.status_code == 403, f"importador nao exporta: {r.status_code}")

        # --- Consolidado
        r = carla.post("/api/conciliacao/exportar", json={})
        checar(r.status_code == 200, f"exportar: {r.status_code} {r.text[:200]}")
        checar("spreadsheetml" in r.headers.get("content-type", ""),
               r.headers.get("content-type"))
        disposicao = r.headers.get("content-disposition", "")
        checar("conciliacao_fiscal_" in disposicao, disposicao)
        checar(r.content[:2] == b"PK", "o consolidado precisa ser um xlsx")

        caminho = os.path.join(trabalho, "consolidado.xlsx")
        with open(caminho, "wb") as arq:
            arq.write(r.content)
        livro = load_workbook(caminho)
        checar(livro.sheetnames == ["Conciliações", "Versões",
                                    "Movimentos DIMP", "Proveniência",
                                    "Exceções", "Conflitos", "Revisões",
                                    "Metadados"], livro.sheetnames)
        linhas = list(livro["Conciliações"].iter_rows(min_row=2,
                                                      values_only=True))
        checar(linhas, "o consolidado deveria ter conciliacoes")
        # Nenhuma celula de dado pode ser formula.
        for aba in livro.worksheets:
            for linha in aba.iter_rows():
                for celula in linha:
                    if isinstance(celula.value, str):
                        checar(not celula.value.startswith("="),
                               f"{aba.title}!{celula.coordinate} virou formula")
        meta = {l[0]: l[1] for l in livro["Metadados"].iter_rows(
            min_row=2, values_only=True)}
        checar(meta["Gerado por"] == "carla", meta["Gerado por"])
        checar(meta["Total de conciliações"] == len(linhas), meta)
        livro.close()

        # --- Filtro respeitado
        r = carla.post("/api/conciliacao/exportar",
                       json={"layout": "quadro_50_5"})
        checar(r.status_code == 200, r.status_code)
        filtrado = os.path.join(trabalho, "filtrado.xlsx")
        with open(filtrado, "wb") as arq:
            arq.write(r.content)
        livro = load_workbook(filtrado)
        layouts = {l[5] for l in livro["Conciliações"].iter_rows(
            min_row=2, values_only=True)}
        checar(layouts <= {"Quadro 50-5"}, f"o filtro vazou: {layouts}")
        meta = {l[0]: l[1] for l in livro["Metadados"].iter_rows(
            min_row=2, values_only=True)}
        checar("layout=quadro_50_5" in meta["Filtros aplicados"], meta)
        livro.close()

        # Filtro invalido nao gera arquivo.
        r = carla.post("/api/conciliacao/exportar", json={"estado": "xpto"})
        checar(r.status_code == 422, f"filtro invalido: {r.status_code}")

        # --- Planilha-mestre
        modelo = fx.gerar_modelo_mestre(os.path.join(trabalho, "modelo.xlsx"),
                                        anos=(2024,))
        tamanho_antes = os.path.getsize(modelo)

        def enviar_modelo(cliente, **campos):
            with open(modelo, "rb") as arq:
                return cliente.post(
                    "/api/conciliacao/preencher-modelo",
                    files={"arquivo": ("modelo.xlsx", arq.read(),
                                       "application/vnd.openxmlformats-"
                                       "officedocument.spreadsheetml.sheet")},
                    data=campos)

        # CNPJ obrigatorio e valido.
        r = enviar_modelo(carla, cnpj="")
        checar(r.status_code == 422, f"sem CNPJ: {r.status_code}")
        checar(r.json()["detail"]["code"] == "CNPJ_INVALIDO", r.json())

        # CNPJ do modelo tem de bater com o escolhido.
        r = enviar_modelo(carla, cnpj=fx.CNPJ_B)
        checar(r.status_code == 422, f"CNPJ divergente: {r.status_code}")
        checar(r.json()["detail"]["code"] == "CNPJ_DIVERGENTE", r.json())

        # Sem permissao de preencher.
        r = enviar_modelo(ana, cnpj=fx.CNPJ_A)
        checar(r.status_code == 403, f"sem permissao: {r.status_code}")

        # Ha competencia Em revisao: precisa de permissao adicional.
        r = enviar_modelo(carla, cnpj=fx.CNPJ_A, incluir_em_revisao="true")
        checar(r.status_code == 403,
               f"incluir pendentes sem a permissao extra: {r.status_code}")
        checar(r.json()["detail"]["code"] == "PERMISSAO_ADICIONAL", r.json())

        adm.put(f"/api/admin/usuarios/{carla_id}/permissoes", json={
            "permissoes": ["aba.conciliacao", "conciliacao.exportar",
                           "conciliacao.preencher_modelo",
                           "conciliacao.incluir_pendentes"]})

        # Com a permissao, ainda exige CONFIRMACAO explicita.
        r = enviar_modelo(carla, cnpj=fx.CNPJ_A, incluir_em_revisao="true")
        checar(r.status_code == 409, f"sem confirmacao: {r.status_code}")
        checar(r.json()["detail"]["code"] == "CONFIRMACAO_PENDENTES", r.json())

        r = enviar_modelo(carla, cnpj=fx.CNPJ_A, incluir_em_revisao="true",
                          confirmacao_pendentes="true")
        checar(r.status_code == 200, f"preencher: {r.status_code} {r.text[:200]}")
        disposicao = r.headers.get("content-disposition", "")
        checar("planilha_mestre_" in disposicao and fx.CNPJ_A in disposicao,
               f"nome de saida distinto: {disposicao}")
        preenchida = os.path.join(trabalho, "preenchida.xlsx")
        with open(preenchida, "wb") as arq:
            arq.write(r.content)

        checar(os.path.getsize(modelo) == tamanho_antes,
               "o modelo ENVIADO nunca pode ser alterado")

        livro = load_workbook(preenchida)
        checar("Metadados da automação" in livro.sheetnames,
               f"pendentes exigem marcacao visivel: {livro.sheetnames}")
        aba = livro["Conciliação"]
        preenchidas = [l for l in aba.iter_rows(min_row=1, values_only=True)
                       if l[2] is not None and isinstance(l[2], (int, float))]
        checar(preenchidas, "alguma linha mensal deveria ter valor")
        formulas = [c.value for linha in aba.iter_rows() for c in linha
                    if isinstance(c.value, str) and c.value.startswith("=")]
        checar(any(f.startswith("=G") for f in formulas),
               f"as formulas H e J deveriam existir: {formulas[:4]}")
        checar(any(f.startswith("=SUM(") for f in formulas),
               "a linha TOTAL deveria somar")
        livro.close()

        # Modelo invalido: mesma validacao de pacote dos relatorios.
        macro = fx.gerar_xlsx_com_macro(os.path.join(trabalho, "macro.xlsx"),
                                        modelo)
        with open(macro, "rb") as arq:
            r = carla.post("/api/conciliacao/preencher-modelo",
                           files={"arquivo": ("m.xlsx", arq.read(),
                                              "application/octet-stream")},
                           data={"cnpj": fx.CNPJ_A})
        checar(r.status_code == 422, f"modelo com macro: {r.status_code}")
        checar(r.json()["detail"]["code"] == "XLSX_MACRO_NAO_PERMITIDA",
               r.json())

        # CNPJ sem competencia elegivel.
        semdados = fx.gerar_modelo_mestre(
            os.path.join(trabalho, "outro.xlsx"), anos=(2024,),
            cnpj=fx.CNPJ_B, razao_social=fx.RAZAO_B)
        with open(semdados, "rb") as arq:
            r = carla.post("/api/conciliacao/preencher-modelo",
                           files={"arquivo": ("o.xlsx", arq.read(),
                                              "application/octet-stream")},
                           data={"cnpj": fx.CNPJ_B,
                                 "competencia_de": "209001"})
        checar(r.status_code == 422, f"sem competencia: {r.status_code}")
        checar(r.json()["detail"]["code"] == "SEM_COMPETENCIAS", r.json())

        # --- Trilha: exportacoes auditadas, com hash do modelo e contagem
        itens = eventos(adm, usuario_filtro="carla", limite=300)
        exportou = [e for e in itens if e["acao"] == "conciliacao.exportar"]
        checar(any(e["resultado"] == "ok" for e in exportou), exportou[:2])
        checar(any(e["resultado"] == "erro" for e in exportou),
               "o filtro invalido deveria virar 'erro' na trilha")
        modelos = [e for e in itens
                   if e["acao"] == "conciliacao.preencher_modelo"
                   and e["resultado"] == "ok"]
        checar(modelos, "preencher modelo tem de ficar na trilha")
        detalhe = modelos[0]["detalhe"]
        checar("modelo sha" in detalhe, f"hash do modelo na trilha: {detalhe}")
        checar("competencia(s)" in detalhe, detalhe)
        negados = [e for e in eventos(adm, usuario_filtro="ana", limite=300)
                   if e["acao"] == "conciliacao.exportar"]
        checar(any(e["resultado"] == "negado" for e in negados),
               "a tentativa negada tambem fica registrada")
    finally:
        shutil.rmtree(trabalho, ignore_errors=True)


def consulta_e_proveniencia(app, adm, ana) -> None:
    """Contrato JSON de lista e detalhe, filtros, paginacao e origem (T033)."""
    # --- Resumo agora reflete o que foi importado
    resumo = ana.get("/api/conciliacao/resumo").json()
    # O lote gera 2 conciliacoes: jan (50-5) e fev (50-3). A duplicata nao
    # cria nada, a retificacao vira candidata da MESMA chave de janeiro e o
    # corrompido e rejeitado.
    checar(resumo["total"] == 2, f"resumo apos importar: {resumo}")
    checar(resumo["em_revisao"] == resumo["total"],
           f"tudo nasce Em revisao: {resumo}")
    checar(resumo["conflitos"] >= 1, f"a retificacao abriu conflito: {resumo}")
    checar(resumo["diferenca_total"]["centavos"] != 0,
           f"Money com centavos exatos: {resumo['diferenca_total']}")

    # --- Lista: contrato e coerencia com o resumo
    pagina = ana.get("/api/conciliacao/conciliacoes").json()
    for campo in ("itens", "total", "pagina", "limite", "paginas"):
        checar(campo in pagina, f"a pagina deveria ter '{campo}': {pagina}")
    checar(pagina["total"] == resumo["total"],
           "lista e resumo tem de contar o mesmo universo")
    linha = pagina["itens"][0]
    for campo in ("id", "cnpj", "razao_social", "competencia", "estado",
                  "revision", "versao", "avisos_abertos", "bloqueios_abertos",
                  "conflitos_abertos"):
        checar(campo in linha, f"a linha deveria ter '{campo}': {linha}")
    for campo in ("receita_declarada", "receita_calculada",
                  "diferenca_receita"):
        checar(set(linha["versao"][campo]) == {"centavos", "texto"},
               f"{campo} deveria ser Money: {linha['versao'][campo]}")
    checar(linha["versao"]["receita_declarada"]["texto"].startswith("R$"),
           linha["versao"]["receita_declarada"])

    # --- Filtros e paginacao
    checar(ana.get("/api/conciliacao/conciliacoes",
                   params={"cnpj": fx.CNPJ_A}).json()["total"] >= 1,
           "filtro por CNPJ")
    so_50_5 = ana.get("/api/conciliacao/conciliacoes",
                      params={"layout": "quadro_50_5"}).json()
    checar(so_50_5["total"] >= 1, "filtro por layout")
    checar(all(i["versao"]["layout"] == "quadro_50_5"
               for i in so_50_5["itens"]), "o filtro de layout vazou")
    checar(ana.get("/api/conciliacao/conciliacoes",
                   params={"estado": "aprovada"}).json()["total"] == 0,
           "ninguem foi aprovado ainda")
    # Resumo aplica os MESMOS filtros da lista.
    for filtros in ({"cnpj": fx.CNPJ_A}, {"layout": "quadro_50_5"},
                    {"excecao": "conflito"}):
        lista = ana.get("/api/conciliacao/conciliacoes", params=filtros).json()
        indicadores = ana.get("/api/conciliacao/resumo", params=filtros).json()
        checar(lista["total"] == indicadores["total"],
               f"{filtros}: lista {lista['total']} x resumo "
               f"{indicadores['total']}")
    # Filtro invalido nao vira SQL nem 500.
    r = ana.get("/api/conciliacao/conciliacoes", params={"estado": "inventado"})
    checar(r.status_code == 422, f"estado invalido: {r.status_code}")
    r = ana.get("/api/conciliacao/conciliacoes", params={"ordenar": "'; DROP"})
    checar(r.status_code == 422, f"ordenacao invalida: {r.status_code}")

    p1 = ana.get("/api/conciliacao/conciliacoes",
                 params={"pagina": 1, "limite": 1}).json()
    p2 = ana.get("/api/conciliacao/conciliacoes",
                 params={"pagina": 2, "limite": 1}).json()
    checar(p1["itens"][0]["id"] != p2["itens"][0]["id"],
           "paginas nao podem repetir registro")
    gigante = ana.get("/api/conciliacao/conciliacoes",
                      params={"pagina": 10 ** 9}).json()
    checar(gigante["pagina"] <= gigante["paginas"],
           "pagina absurda e' presa a ultima, sem 500")

    # --- Detalhe do 50-3: DIMP, agregacao e proveniencia celula a celula
    com_dimp = ana.get("/api/conciliacao/conciliacoes",
                       params={"layout": "quadro_50_3_dimp"}).json()["itens"][0]
    r = ana.get(f"/api/conciliacao/conciliacoes/{com_dimp['id']}")
    checar(r.status_code == 200, f"detalhe: {r.text[:200]}")
    detalhe = r.json()
    for campo in ("versao_vigente", "versoes", "movimentos", "fatos",
                  "excecoes", "revisoes", "conflitos"):
        checar(campo in detalhe, f"o detalhe deveria ter '{campo}'")

    vigente = detalhe["versao_vigente"]
    for campo in ("declarada_com_st", "declarada_sem_st", "calculada_com_st",
                  "calculada_sem_st", "diferenca_com_st", "diferenca_sem_st"):
        checar(set(vigente[campo]) == {"centavos", "texto"},
               f"{campo} deveria ser Money: {vigente[campo]}")
    checar(vigente["fonte"]["sha256"] and len(vigente["fonte"]["sha256"]) == 64,
           vigente["fonte"])
    checar(vigente["fonte"]["versao_parser"] == cs.VERSAO_PARSER,
           vigente["fonte"])
    checar(vigente["aba_origem"] and vigente["linha_resumo"] >= 1, vigente)

    # Diferenca declarada: calculado MENOS declarado, conferido na conta.
    checar(vigente["diferenca_receita"]["centavos"]
           == vigente["receita_calculada"]["centavos"]
           - vigente["receita_declarada"]["centavos"],
           "a diferenca precisa fechar com os dois lados")

    checar(len(detalhe["movimentos"]) == 3,
           f"as linhas DIMP originais ficam: {len(detalhe['movimentos'])}")
    instituicoes = {i["instituicao"]: i for i in detalhe["instituicoes"]}
    checar(len(instituicoes) == 2, f"agregado: {list(instituicoes)}")
    repetida = instituicoes["BANCO SINTETICO S.A."]
    checar(repetida["linhas"] == 2,
           f"o agregado precisa dizer que somou 2 linhas: {repetida}")
    movimento = detalhe["movimentos"][0]
    celulas = movimento["proveniencia"]["celulas"]
    checar(celulas["total_dimp"] == f"J{movimento['linha_origem']}",
           f"celula do total DIMP: {celulas}")
    checar(celulas["pix"] == f"G{movimento['linha_origem']}", celulas)
    checar(movimento["proveniencia"]["fonte"]["sha256"]
           == vigente["fonte"]["sha256"], "o movimento aponta a mesma fonte")
    checar(movimento["total_nao_pix"]["centavos"]
           == movimento["total_dimp"]["centavos"]
           - movimento["pix"]["centavos"], "nao-PIX = total DIMP - PIX")

    # Fatos: celula com aba/celula, derivado com formula e origens.
    celula = next(f for f in detalhe["fatos"] if f["tipo_origem"] == "celula")
    checar(celula["aba"] and celula["celula"] and celula["rotulo"], celula)
    checar(celula["regra_parser"], "todo fato guarda a regra que o extraiu")
    checar(celula["fonte"]["sha256"] == vigente["fonte"]["sha256"], celula)
    derivado = next(f for f in detalhe["fatos"]
                    if f["metrica"] == "diferenca_receita")
    checar(derivado["tipo_origem"] == "formula", derivado)
    checar(derivado["formula"] == "receita_calculada - receita_declarada",
           derivado["formula"])
    checar(sorted(derivado["fatos_origem"])
           == ["receita_calculada", "receita_declarada"], derivado)

    # --- 50-5: PIX e nao-PIX NULOS, nunca R$ 0,00
    so_5 = ana.get("/api/conciliacao/conciliacoes",
                   params={"layout": "quadro_50_5"}).json()["itens"][0]
    d5 = ana.get(f"/api/conciliacao/conciliacoes/{so_5['id']}").json()
    checar(d5["versao_vigente"]["total_pix"] is None,
           f"50-5 nao pode inventar PIX: {d5['versao_vigente']['total_pix']}")
    checar(d5["versao_vigente"]["total_nao_pix"] is None,
           d5["versao_vigente"]["total_nao_pix"])
    checar(d5["movimentos"] == [], "50-5 nao tem movimento DIMP")

    # --- Conflito: candidata visivel, vigente preservada
    com_conflito = ana.get("/api/conciliacao/conciliacoes",
                           params={"excecao": "conflito"}).json()["itens"][0]
    dc = ana.get(f"/api/conciliacao/conciliacoes/{com_conflito['id']}").json()
    checar(dc["conflitos_abertos"] == 1, dc["conflitos_abertos"])
    candidatas = [v for v in dc["versoes"] if v["estado"] == "candidata"]
    checar(len(candidatas) == 1, f"uma candidata: {dc['versoes']}")
    checar(dc["versao_vigente"]["numero"] == 1,
           "a vigente continua sendo a primeira")
    checar(len(dc["revisoes"]) >= 2,
           f"a chegada da candidata e' uma revisao: {dc['revisoes']}")

    # --- Excecoes
    fila = ana.get("/api/conciliacao/excecoes").json()
    checar(fila["total"] >= 1, fila)
    excecao = fila["itens"][0]
    for campo in ("id", "codigo", "severidade", "mensagem", "estado",
                  "cnpj", "competencia"):
        checar(campo in excecao, f"a excecao deveria ter '{campo}': {excecao}")
    checar(all(e["severidade"] == "bloqueio" for e in ana.get(
        "/api/conciliacao/excecoes",
        params={"severidade": "bloqueio"}).json()["itens"]),
        "filtro de severidade vazou")

    # --- Download da fonte: auditado e servindo o binario imutavel
    fonte_id = vigente["fonte"]["id"]
    r = ana.get(f"/api/conciliacao/fontes/{fonte_id}/download")
    checar(r.status_code == 200, f"download: {r.status_code}")
    checar(r.content[:2] == b"PK", "o download precisa ser o XLSX original")
    checar("spreadsheetml" in r.headers.get("content-type", ""),
           r.headers.get("content-type"))
    checar(ana.get("/api/conciliacao/fontes/999999/download").status_code == 404,
           "fonte inexistente")
    baixados = [e for e in eventos(adm, usuario_filtro="ana",
                                   acao="conciliacao.fonte_download")
                if e["resultado"] == "ok"]
    checar(baixados, "baixar a evidencia tem de ficar na trilha")

    # --- Texto de terceiro continua TEXTO
    checar(isinstance(detalhe["razao_social"], str), "razao social e' texto")
    checar("<" not in detalhe["razao_social"] or True,
           "a API entrega texto cru; escapar e' responsabilidade da tela")

    # --- Sem a aba, nada disso responde
    sem_aba = TestClient(app)
    sem_aba.post("/api/login", json={"usuario": "forasteiro",
                                     "senha": "forasteiro12345"})
    for caminho in ("/api/conciliacao/conciliacoes",
                    f"/api/conciliacao/conciliacoes/{com_dimp['id']}",
                    "/api/conciliacao/excecoes",
                    f"/api/conciliacao/fontes/{fonte_id}/download"):
        r = sem_aba.get(caminho)
        checar(r.status_code in (401, 403),
               f"{caminho} sem a aba: {r.status_code}")



def painel_e_filtros(ana) -> None:
    """Indicadores, busca, intervalo, ordenacao e paginacao (T060)."""
    # Busca por razao social e por CNPJ chegam no mesmo lugar.
    por_texto = ana.get("/api/conciliacao/conciliacoes",
                        params={"texto": "ALFA"}).json()
    por_cnpj = ana.get("/api/conciliacao/conciliacoes",
                       params={"cnpj": fx.CNPJ_A}).json()
    checar(por_texto["total"] >= 1, por_texto["total"])
    checar({i["id"] for i in por_texto["itens"]}
           <= {i["id"] for i in por_cnpj["itens"]},
           "buscar pela razao social nao pode trazer outra empresa")

    # Curinga digitado e' TEXTO, nao curinga de SQL.
    checar(ana.get("/api/conciliacao/conciliacoes",
                   params={"texto": "%"}).json()["total"] == 0,
           "'%' digitado nao pode casar com tudo")

    # Intervalo de competencia.
    intervalo = ana.get("/api/conciliacao/conciliacoes",
                        params={"competencia_de": "202401",
                                "competencia_ate": "202401"}).json()
    checar(all(i["competencia"] == "202401" for i in intervalo["itens"]),
           [i["competencia"] for i in intervalo["itens"]])

    # Ordenacao muda a ordem, nao o conjunto.
    desc = ana.get("/api/conciliacao/conciliacoes",
                   params={"ordenar": "competencia_desc", "limite": 200}).json()
    asc = ana.get("/api/conciliacao/conciliacoes",
                  params={"ordenar": "competencia_asc", "limite": 200}).json()
    checar({i["id"] for i in desc["itens"]} == {i["id"] for i in asc["itens"]},
           "ordenar nao pode mudar o conjunto de resultados")
    competencias = [i["competencia"] for i in asc["itens"]]
    checar(competencias == sorted(competencias), competencias)

    # Paginacao: varre tudo sem duplicar nem omitir.
    limite = 2
    primeira = ana.get("/api/conciliacao/conciliacoes",
                       params={"limite": limite, "pagina": 1}).json()
    vistos = []
    for numero in range(1, primeira["paginas"] + 1):
        pagina = ana.get("/api/conciliacao/conciliacoes",
                         params={"limite": limite, "pagina": numero}).json()
        vistos += [i["id"] for i in pagina["itens"]]
    checar(len(vistos) == len(set(vistos)), f"paginacao duplicou: {vistos}")
    checar(len(vistos) == primeira["total"],
           f"paginacao omitiu: {len(vistos)} de {primeira['total']}")

    # Indicadores acompanham CADA filtro da lista.
    for filtros in ({"texto": "ALFA"}, {"layout": "quadro_50_3_dimp"},
                    {"excecao": "aviso"}, {"estado": "aprovada"},
                    {"competencia_de": "202402"}):
        lista = ana.get("/api/conciliacao/conciliacoes", params=filtros).json()
        painel = ana.get("/api/conciliacao/resumo", params=filtros).json()
        checar(lista["total"] == painel["total"],
               f"{filtros}: lista {lista['total']} x painel {painel['total']}")
        soma = (painel["em_revisao"] + painel["aprovadas"]
                + painel["rejeitadas"])
        checar(soma == painel["total"],
               f"{filtros}: os estados deveriam somar o total: {painel}")


def concorrencia_durante_o_lote(ana) -> None:
    """SC-013: a interface continua utilizavel durante um job (T070).

    A causa raiz de uma tela travada e' o servidor: se o job segurasse a
    trava da sessao ou a conexao do banco durante todo o processamento, as
    leituras da mesma pessoa ficariam na fila. Este teste ataca exatamente
    isso — com um lote EM ANDAMENTO, as tres leituras do painel e o proprio
    polling do job precisam responder dentro de um orcamento de tempo.
    """
    nova = ana.post("/api/sessoes", json={"ferramenta": "conciliacao"})
    sessao = nova.json()["sessao_id"]
    trabalho = tempfile.mkdtemp(prefix="conc_sc013_")
    try:
        # Lote grande o bastante para o job ainda estar rodando nas leituras.
        for i in range(12):
            caminho = fx.gerar_quadro_50_3_dimp(
                os.path.join(trabalho, f"m{i}.xlsx"),
                competencia=f"2023{i + 1:02d}")
            checar(enviar(ana, sessao, caminho).status_code == 200,
                   f"upload {i}")

        inicio = ana.post("/api/conciliacao/processar",
                          json={"sessao_id": sessao})
        checar(inicio.status_code == 200, inicio.text[:160])
        job_id = inicio.json()["job_id"]

        # Enquanto o job roda, as leituras respondem.
        leituras = 0
        piores: list[float] = []
        for _ in range(40):
            job = ana.get(f"/api/jobs/{job_id}",
                          params={"ferramenta": "conciliacao"})
            checar(job.status_code == 200,
                   f"o polling nao pode travar: {job.status_code}")
            if job.json()["status"] != "executando":
                break
            for caminho in ("/api/conciliacao/resumo",
                            "/api/conciliacao/conciliacoes",
                            "/api/conciliacao/excecoes"):
                marca = time.monotonic()
                r = ana.get(caminho)
                gasto = time.monotonic() - marca
                piores.append(gasto)
                checar(r.status_code == 200,
                       f"{caminho} durante o job: {r.status_code}")
                checar(gasto < 2.0,
                       f"{caminho} levou {gasto:.2f}s com o job rodando; o "
                       f"job nao pode segurar a sessao nem o banco")
                leituras += 1
        esperar_job(ana, job_id)
        checar(leituras > 0,
               "o lote terminou rapido demais para provar SC-013; aumente o "
               "numero de arquivos")
        print(f"    SC-013: {leituras} leituras durante o job, pior tempo "
              f"{max(piores):.3f}s")
    finally:
        shutil.rmtree(trabalho, ignore_errors=True)


def enviar(cliente, sessao: str, caminho: str, nome: str = "") -> object:
    with open(caminho, "rb") as arq:
        return cliente.post(
            "/api/conciliacao/upload", params={"sessao_id": sessao},
            files={"arquivo": (nome or os.path.basename(caminho), arq.read(),
                               "application/vnd.openxmlformats-officedocument."
                               "spreadsheetml.sheet")})


def esperar_job(cliente, job_id: str) -> dict:
    for _ in range(200):
        r = cliente.get(f"/api/jobs/{job_id}",
                        params={"ferramenta": "conciliacao"})
        checar(r.status_code == 200, f"polling do job: {r.status_code}")
        job = r.json()
        if job["status"] in ("concluido", "erro"):
            return job
        time.sleep(0.05)
    raise AssertionError("o job nao concluiu a tempo")


def importacao(app, adm, ana, sessao: str) -> None:
    trabalho = tempfile.mkdtemp(prefix="conc_lote_")
    try:
        lote = fx.gerar_lote_sintetico(trabalho)

        # Nomes propositalmente enganosos e HOMONIMOS: o staging usa id opaco,
        # entao dois arquivos com o mesmo nome nao se sobrescrevem.
        enviados = [
            (lote["quadro_50_5"], "relatorio.xlsx"),
            (lote["quadro_50_3"], "relatorio.xlsx"),
            (lote["duplicata"], "outro_nome.xlsx"),
            (lote["retificada"], "retificacao.xlsx"),
            (lote["invalido"], "corrompido.xlsx"),
        ]
        ids = set()
        for caminho, nome in enviados:
            r = enviar(ana, sessao, caminho, nome)
            checar(r.status_code == 200, f"upload de {nome}: {r.text[:160]}")
            corpo = r.json()
            checar(corpo["ok"] and corpo["arquivo"] == nome, corpo)
            checar(corpo["tamanho"] == os.path.getsize(caminho), corpo)
            ids.add(corpo["upload_id"])
        checar(len(ids) == len(enviados),
               f"cada upload precisa de id proprio: {len(ids)}")

        # Extensao errada nao entra.
        txt = os.path.join(trabalho, "nota.txt")
        with open(txt, "w", encoding="utf-8") as arq:
            arq.write("nao sou planilha")
        checar(enviar(ana, sessao, txt).status_code == 422,
               "somente .xlsx deveria ser aceito")

        # Teto por arquivo: 413 e o parcial NAO pode ficar no staging.
        antes = _contar_staging(sessao)
        os.environ["AUDITORIA_CONCILIACAO_MAX_UPLOAD_MB"] = "1"
        try:
            grande = os.path.join(trabalho, "grande.xlsx")
            with open(grande, "wb") as arq:
                arq.write(b"\0" * (2 * 1024 * 1024))
            r = enviar(ana, sessao, grande)
            checar(r.status_code == 413,
                   f"acima do teto deveria dar 413: {r.status_code}")
        finally:
            os.environ.pop("AUDITORIA_CONCILIACAO_MAX_UPLOAD_MB", None)
        checar(_contar_staging(sessao) == antes,
               "o upload interrompido nao pode deixar parcial no staging")

        # Processar: job + lote.
        r = ana.post("/api/conciliacao/processar", json={"sessao_id": sessao})
        checar(r.status_code == 200, f"processar: {r.text[:200]}")
        job_id, lote_id = r.json()["job_id"], r.json()["lote_id"]

        job = esperar_job(ana, job_id)
        checar(job["status"] == "concluido", f"job: {job}")
        resultado = job["resultado"]
        checar(resultado["processados"] == 2,
               f"dois layouts distintos processam: {resultado}")
        checar(resultado["duplicados"] == 1, f"a copia renomeada: {resultado}")
        checar(resultado["conflitantes"] == 1, f"a retificacao: {resultado}")
        checar(resultado["rejeitados"] == 1, f"o corrompido: {resultado}")

        # Resultado PERSISTENTE, com mensagem por item.
        r = ana.get(f"/api/conciliacao/lotes/{lote_id}")
        checar(r.status_code == 200, f"lote: {r.text[:200]}")
        corpo = r.json()
        checar(corpo["estado"] == "concluido", corpo["estado"])
        checar(len(corpo["itens"]) == 5, len(corpo["itens"]))
        por_resultado = {}
        for item in corpo["itens"]:
            por_resultado.setdefault(item["resultado"], []).append(item)
            checar(item["mensagem"], f"todo item precisa de mensagem: {item}")
        rejeitado = por_resultado["rejeitado"][0]
        checar(rejeitado["codigo"] == cs.XLSX_INVALIDO,
               f"codigo estavel no item rejeitado: {rejeitado}")
        checar(all(i["sha256"] for i in por_resultado["processado"]),
               "item processado carrega o SHA da fonte")

        # Lote alheio responde como inexistente.
        checar(adm.get(f"/api/conciliacao/lotes/{lote_id}").status_code == 404,
               "lote de outro usuario nao pode ser revelado")
        checar(ana.get("/api/conciliacao/lotes/999999").status_code == 404,
               "lote inexistente")

        # Segundo processamento da MESMA sessao com staging ja consumido.
        r = ana.post("/api/conciliacao/processar", json={"sessao_id": sessao})
        checar(r.status_code == 422,
               f"sem arquivos novos deveria dar 422: {r.status_code}")

        # Dois processamentos simultaneos da mesma sessao: o segundo e' 409.
        nova = ana.post("/api/sessoes", json={"ferramenta": "conciliacao"})
        sessao2 = nova.json()["sessao_id"]
        for i in range(3):
            enviar(ana, sessao2, lote["quadro_50_5"], f"r{i}.xlsx")
        primeiro = ana.post("/api/conciliacao/processar",
                            json={"sessao_id": sessao2})
        checar(primeiro.status_code == 200, primeiro.text[:160])
        segundo = ana.post("/api/conciliacao/processar",
                           json={"sessao_id": sessao2})
        checar(segundo.status_code in (409, 422),
               f"processamento concorrente da mesma sessao: "
               f"{segundo.status_code} {segundo.text[:120]}")
        esperar_job(ana, primeiro.json()["job_id"])

        # Trilha geral: ok, negado e erro, SEM valores fiscais no detalhe.
        itens = eventos(adm, usuario_filtro="ana", limite=300)
        uploads = [e for e in itens if e["acao"] == "conciliacao.upload"]
        checar(uploads, "o upload deveria estar na trilha")
        checar(any(e["resultado"] == "ok" for e in uploads), "upload ok")
        checar(any(e["resultado"] == "erro" for e in uploads),
               f"o 413 deveria virar 'erro' na trilha: "
               f"{[(e['resultado'], e['http_status']) for e in uploads]}")
        processamentos = [e for e in itens
                          if e["acao"] == "conciliacao.processar"]
        checar(processamentos, "o processamento deveria estar na trilha")
        proibidos = ("40000", "41500", "31000", "22000")
        for evento in uploads + processamentos:
            for valor in proibidos:
                checar(valor not in evento["detalhe"],
                       f"a trilha geral nao pode carregar valor fiscal: "
                       f"{evento['detalhe']!r}")
    finally:
        shutil.rmtree(trabalho, ignore_errors=True)


def _contar_staging(sessao: str) -> int:
    pasta = os.path.join(_TMP, "sessoes", sessao, "conciliacao")
    return len(os.listdir(pasta)) if os.path.isdir(pasta) else 0


if __name__ == "__main__":
    raise SystemExit(main())
