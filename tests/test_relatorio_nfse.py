"""Saidas da Emissao de NFS-e (spec 006, historia 8): planilha-modelo,
planilha-resumo do lote e zip das DPS. Tudo marcado como DPS NAO TRANSMITIDA,
texto neutralizado contra formula e nenhum arquivo parcial em caso de falha.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import os
import sys
import tempfile
import zipfile

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "src"))

from openpyxl import load_workbook  # noqa: E402

from auditoria_fiscal.core import nfse_dps as nd  # noqa: E402
from auditoria_fiscal.core import nfse_entrada as ne  # noqa: E402
from auditoria_fiscal.ferramentas import nfse_store as ns  # noqa: E402
from auditoria_fiscal.ferramentas import relatorio_nfse as rn  # noqa: E402
from fixtures import nfse as fx  # noqa: E402

HOJE = dt.date(2026, 11, 15)
AGORA = dt.datetime(2026, 11, 15, 10, 30, tzinfo=nd.FUSO_SALVADOR)
U = {"usuario_id": 1, "usuario_login": "fiscal"}

TEXTO = """Competência: 11/2026
Responsável: =HYPERLINK("http://x.invalid","clique")
CPF: 123.456.789-09
Aluno: Pedro Exemplo
Turma: Berçário
Valor: 1.250,00

Competência: 11/2026
Responsável: Bruno Teste
CPF: 111.444.777-35
Aluno: Caio Teste
Turma: Infantil 2
Valor: 980,50
"""


def checar(cond, msg):
    if not cond:
        print(f"FALHOU - {msg}")
        raise SystemExit(1)


def testar_planilha_modelo(pasta: str) -> None:
    destino = os.path.join(pasta, "modelo.xlsx")
    rn.gerar_planilha_modelo(destino)
    livro = load_workbook(destino)
    checar(livro.sheetnames == ["Notas", "Instruções"], livro.sheetnames)
    notas = livro["Notas"]
    cabecalho = [c.value for c in notas[1]]
    checar(cabecalho == list(ne.ROTULOS_MODELO), cabecalho)
    coluna_cpf = cabecalho.index("CPF do responsável") + 1
    checar(notas.cell(row=2, column=coluna_cpf).number_format == "@",
           "coluna de CPF formatada como texto (o Excel nao apaga o zero)")
    checar(all(c.value is None for linha in notas.iter_rows(min_row=2) for c in linha),
           "a planilha-modelo nao traz dado nenhum")
    instrucoes = " ".join(str(c.value or "") for l in livro["Instruções"].iter_rows()
                          for c in l)
    checar("Competência" in instrucoes and "uma linha por aluno" in instrucoes,
           instrucoes[:200])

    # Volta completa: o que o sistema entrega, o proprio importador le.
    notas.append(["11/2026", "Ana Exemplo", "012.345.678-90", None, None,
                  "Pedro Exemplo", "Berçário", "1.250,00"])
    preenchido = os.path.join(pasta, "preenchido.xlsx")
    livro.save(preenchido)
    r = ne.ler_planilha(preenchido)
    checar(len(r.linhas) == 1 and r.avisos == [], (len(r.linhas), r.avisos))
    checar(r.linhas[0].campos["responsavel_cpf"] == "012.345.678-90", r.linhas[0].campos)


def _lote_aprovado(pasta: str) -> tuple[ns.NfseStore, int]:
    store = ns.NfseStore(os.path.join(pasta, "rel.db"), os.path.join(pasta, "origens"))
    pid = store.criar_perfil(fx.perfil_completo(), **U)["perfil_id"]
    lote_id = store.importar_texto(pid, TEXTO, hoje=HOJE, **U)["lote_id"]
    store.aprovar(lote_id, revision=1, agora=AGORA, **U)
    return store, lote_id


def testar_resumo(pasta: str) -> None:
    store, lote_id = _lote_aprovado(pasta)
    destino = os.path.join(pasta, "resumo.xlsx")
    rn.gerar_resumo(store.lote(lote_id), store.dps_do_lote(lote_id), destino)
    livro = load_workbook(destino)
    checar(livro.sheetnames == ["Notas", "Lote"], livro.sheetnames)
    notas = livro["Notas"]
    cabecalho = [c.value for c in notas[1]]
    col = {nome: i for i, nome in enumerate(cabecalho)}
    linha1 = [c for c in notas[2]]
    checar(linha1[col["Responsável"]].value.startswith("'="),
           f"texto que comeca com = e' neutralizado: {linha1[col['Responsável']].value}")
    valor = linha1[col["Valor"]]
    checar(float(valor.value) == 1250.0 and "R$" in valor.number_format,
           (valor.value, valor.number_format))
    checar(linha1[col["CPF"]].value == "123.456.789-09", linha1[col["CPF"]].value)
    checar(linha1[col["Nº DPS"]].value == 1, linha1[col["Nº DPS"]].value)
    checar(notas.max_row == 3, f"duas notas: {notas.max_row}")
    lote = " ".join(str(c.value or "") for l in livro["Lote"].iter_rows() for c in l)
    checar("DPS NÃO TRANSMITIDA" in lote and "produção restrita" in lote, lote[:200])
    checar(store.lote(lote_id)["lote"]["hash_aprovacao"] in lote, "hash da aprovacao")
    store.fechar()


def testar_zip(pasta: str) -> None:
    store, lote_id = _lote_aprovado(os.path.join(pasta, "z"))
    dps = store.dps_do_lote(lote_id)
    destino = os.path.join(pasta, "dps.zip")
    rn.gerar_zip(lote_id, dps, destino)
    with zipfile.ZipFile(destino) as z:
        nomes = z.namelist()
        checar("LEIA-ME.txt" in nomes and len(nomes) == len(dps) + 1, nomes)
        checar("NÃO foram transmitidas" in z.read("LEIA-ME.txt").decode("utf-8"),
               "o zip explica que as DPS nao foram transmitidas")
        for d in dps:
            nome = rn.nome_xml(lote_id, d)
            checar(nome in nomes and "NAO-TRANSMITIDA" in nome, nome)
            checar(hashlib.sha256(z.read(nome)).hexdigest() == d["sha256"],
                   "XML no zip e' o mesmo gravado")
    store.fechar()


def testar_falha_sem_parcial(pasta: str) -> None:
    store, lote_id = _lote_aprovado(os.path.join(pasta, "f"))
    dps = store.dps_do_lote(lote_id)
    quebrado = [dict(dps[0]), dict(dps[1], xml=None)]
    destino = os.path.join(pasta, "quebrado.zip")
    try:
        rn.gerar_zip(lote_id, quebrado, destino)
        checar(False, "zip com XML ausente deveria falhar")
    except (TypeError, ValueError):
        pass
    checar(not os.path.exists(destino) and not os.path.exists(destino + ".parcial"),
           "falha nao deixa zip parcial")
    inexistente = os.path.join(pasta, "nao", "existe", "resumo.xlsx")
    try:
        rn.gerar_resumo(store.lote(lote_id), dps, inexistente)
        checar(False, "pasta inexistente deveria falhar")
    except OSError:
        pass
    checar(not os.path.exists(inexistente), "falha nao deixa planilha parcial")
    store.fechar()


def main() -> int:
    with tempfile.TemporaryDirectory() as pasta:
        os.makedirs(os.path.join(pasta, "z"))
        os.makedirs(os.path.join(pasta, "f"))
        testar_planilha_modelo(pasta)
        testar_resumo(pasta)
        testar_zip(pasta)
        testar_falha_sem_parcial(pasta)
    print("OK - saidas da NFS-e (planilha-modelo sem dado e lida pelo proprio "
          "importador, resumo com moeda numerica e texto neutralizado, zip com "
          "LEIA-ME e os mesmos XML gravados, tudo marcado como nao transmitido, "
          "falha sem arquivo parcial) passou.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
