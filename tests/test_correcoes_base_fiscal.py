"""Propostas sintéticas: origem íntegra, decisões explícitas e prévia sem escrita."""
from copy import deepcopy
from dataclasses import asdict
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import sys
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from openpyxl import Workbook, load_workbook
from auditoria_fiscal.core.clientes import ErroClientes
from auditoria_fiscal.core.documentos_br import dv_cnpj
from auditoria_fiscal.core.planilha_empresas import ler_planilha_empresas
from auditoria_fiscal.ferramentas.relatorio_base_fiscal import gerar_relatorio_base_fiscal
from auditoria_fiscal.core.correcoes_base_fiscal import preparar_previa, comparar_cadastro


def cnpj(base):
    return base + dv_cnpj(base)


D1, D2, D3 = (cnpj(b) for b in ('019876540001', '529876540001', '639876540001'))


def fonte(pasta):
    p = pasta / 'fonte.xlsx'
    w = Workbook()
    s = w.active
    s.title = 'Fiscal sintético'
    s.append([None, 'RESPONSAVEL', 'EMPRESAS', 'RESPONSAVEL', 'CNPJ', 'REGIME DA EMPRESA'])
    s.append([None, 'Pessoa sintética', 'Empresa A', 'Contato A', D1, 'Lucro Real'])
    s.append([None, None, 'Empresa B', None, 'pendente', None])
    s.append([None, 'Outra pessoa', 'Empresa C', 'Contato C', D1, 'Lucro Real'])
    z = w.create_sheet('ZERADAS')
    z.append([None, 'EMPRESA', 'CNPJ'])
    z.append([None, 'Empresa Z', D3])
    w.save(p)
    w.close()
    return ler_planilha_empresas(p)


def gravar(pasta, original, mudar=None):
    p = pasta / 'correcoes.xlsx'
    w = load_workbook(BytesIO(original))
    if mudar:
        mudar(w)
    w.save(p)
    w.close()
    return p


def resolvido(w):
    s = w['Correções']
    s['L2'] = 'MANTER'
    s['G3'] = D2
    s['I3'] = 'Simples Nacional'
    s['J3'] = 'Pessoa proposta'
    s['K3'] = '<b>Contato proposto</b>'
    s['L4'] = 'Ignorar'
    s['M4'] = 'Ocorrência repetida revisada'


def recusa(acao, codigo='PLANILHA_INVALIDA'):
    try:
        acao()
    except ErroClientes as e:
        assert e.codigo == codigo, (e.codigo, e.mensagem)
        return
    raise AssertionError('Arquivo inválido foi aceito')


def editar_xml(p, nome, acao):
    with ZipFile(p) as z:
        partes = [(i, z.read(i.filename)) for i in z.infolist()]
    with ZipFile(p, 'w', ZIP_DEFLATED) as z:
        for info, dados in partes:
            z.writestr(info, acao(dados) if info.filename == nome else dados)


def main():
    with tempfile.TemporaryDirectory(prefix='correcoes_fiscais_sinteticas_') as tmp:
        pasta = Path(tmp)
        leitura = fonte(pasta)
        original_leitura = deepcopy(asdict(leitura))
        relatorio = gerar_relatorio_base_fiscal(leitura, nome_arquivo='fonte.xlsx')
        p = gravar(pasta, relatorio)
        r = preparar_previa(leitura, p)
        assert r['somente_leitura'] is True
        assert r['formato'] == 'base_fiscal_importacao_v1'
        assert r['hash_arquivo'] == leitura.hash_arquivo
        assert r['hash_correcoes'] == sha256(p.read_bytes()).hexdigest()
        assert r['contagens']['linhas'] == 4
        assert [l['status'] for l in r['linhas']] == ['pendente', 'pendente', 'pendente', 'incluir']
        assert r['linhas'][3]['efetivo']['sem_movimento'] is True
        assert not r['linhas'][3]['pendencias'], 'Colunas inexistentes não geram pendência'
        assert not any('usuario_id' in str(l) for l in r['linhas'])

        p = gravar(pasta, relatorio, resolvido)
        r = preparar_previa(leitura, p)
        assert [l['status'] for l in r['linhas']] == ['incluir', 'incluir', 'ignorar', 'incluir']
        assert r['contagens']['incluir'] == 3 and r['contagens']['ignorar'] == 1
        assert r['linhas'][1]['original']['documento'] == 'pendente'
        assert r['linhas'][1]['efetivo']['documento'] == D2
        assert r['linhas'][1]['proposto']['documento'] == D2
        assert r['linhas'][1]['efetivo']['contato_nome'] == '<b>Contato proposto</b>'
        assert any(a['campo'] == 'documento' and a['celula'] == 'Correções!G3' for a in r['linhas'][1]['alteracoes'])
        assert asdict(leitura) == original_leitura
        snapshot = deepcopy(r)
        clientes = [{'id': 1, 'nome': ' Empresa   A ', 'documento': D1, 'contato_nome': 'Contato A', 'situacao': 'ativo', 'revision': 3},
                    {'id': 2, 'nome': 'Nome diferente', 'documento': D2, 'contato_nome': '', 'situacao': 'ativo', 'revision': 4},
                    {'id': 3, 'nome': 'Empresa Z', 'documento': D3, 'contato_nome': '', 'situacao': 'inativo', 'revision': 2}]
        cad_antes = deepcopy(clientes)
        confrontado = comparar_cadastro(r, clientes)
        assert [l['status'] for l in confrontado['linhas']] == ['existente', 'conflito', 'ignorar', 'conflito']
        assert confrontado['linhas'][0]['cadastro']['revision'] == 3
        assert r == snapshot and clientes == cad_antes
        assert confrontado['contagens']['existente'] == 1 and confrontado['contagens']['conflito'] == 2
        assert preparar_previa(leitura, p) == r

        # Ignorar é decisão sobre esta ocorrência, inclusive quando a fonte tem
        # campos inválidos; não exige inventar correções para excluir da proposta.
        def ignorar_tudo(w):
            for numero in range(2, 6):
                w['Correções'][f'L{numero}'] = 'ignorar'
                w['Correções'][f'M{numero}'] = 'Não incluir nesta proposta sintética'
        ignoradas = preparar_previa(leitura, gravar(pasta, relatorio, ignorar_tudo))
        assert ignoradas['contagens']['ignorar'] == 4
        assert ignoradas['linhas'][1]['pendencias'], 'Ignorar não apaga problemas da fonte'
        assert comparar_cadastro(ignoradas, clientes)['contagens']['ignorar'] == 4

        # Ausência estrutural não é proposta para apagar o contato cadastrado.
        r = preparar_previa(leitura, gravar(pasta, relatorio, resolvido))
        cadastro_z = [{'id': 9, 'nome': 'Empresa Z', 'documento': D3,
                       'contato_nome': 'Contato preservado', 'situacao': 'ativo', 'revision': 1}]
        z = comparar_cadastro(r, cadastro_z)['linhas'][3]
        assert z['status'] == 'existente'
        assert z['cadastro']['contato_nome'] == 'Contato preservado'
        assert any('preserva' in a.lower() for a in z['avisos'])

        # Cadastros ambíguos não são escolhidos pela ordem do snapshot.
        duplicados = [clientes[0], {**clientes[0], 'id': 99}]
        ambigua = comparar_cadastro(r, duplicados)['linhas'][0]
        assert ambigua['status'] == 'conflito' and ambigua['cadastro'] is None
        assert any('Mais de um cadastro' in p for p in ambigua['pendencias'])

        # Uma coincidência cadastral nunca torna apta uma linha ainda pendente.
        def regime_pendente(w):
            resolvido(w); w['Correções']['I2'] = 'Regime não identificado'
        pendente = preparar_previa(leitura, gravar(pasta, relatorio, regime_pendente))
        assert comparar_cadastro(pendente, clientes)['linhas'][0]['status'] == 'pendente'

        # Erro na fonte é exportado como texto literal em A:F; essa origem
        # intacta precisa poder ser corrigida em G:K, sem executar fórmulas.
        original_fonte = pasta / 'fonte.xlsx'
        w = load_workbook(original_fonte)
        w.active['E2'] = '#N/A'
        w.active['C2'] = '#REF!'
        w.active['C2'].data_type = 's'
        w.save(original_fonte); w.close()
        leitura_erro = ler_planilha_empresas(original_fonte)
        relatorio_erro = gerar_relatorio_base_fiscal(leitura_erro, nome_arquivo='fonte.xlsx')
        def corrigir_erro(w):
            resolvido(w); w['Correções']['G2'] = D1
        reparada = preparar_previa(leitura_erro, gravar(pasta, relatorio_erro, corrigir_erro))
        assert reparada['linhas'][0]['status'] == 'incluir'
        assert reparada['linhas'][0]['original']['documento'] == '#N/A'
        assert reparada['linhas'][0]['efetivo']['nome'] == '#REF!'

        # Ordenação não altera IDs; linhas/abas ocultas continuam na prévia.
        def reordenar(w):
            resolvido(w)
            s = w['Correções']; valores = list(s.values)
            s.delete_rows(2, 4)
            for row in reversed(valores[1:]): s.append(row)
            s.row_dimensions[2].hidden = True
            s.sheet_state = 'hidden'
        rr = preparar_previa(leitura, gravar(pasta, relatorio, reordenar))
        assert [l['id'] for l in rr['linhas']] == ['1:2', '1:3', '1:4', '2:2']
        assert rr['linhas'][1]['efetivo']['documento'] == D2

        # O conjunto completo e os valores originais são obrigatórios.
        for mudar in [lambda w: w['Correções'].__setitem__('A2', 'não existe'),
                      lambda w: w['Correções'].__setitem__('A3', '1:2'),
                      lambda w: w['Correções'].delete_rows(2),
                      lambda w: w['Correções'].__setitem__('D2', 'Outro original'),
                      lambda w: w['Correções'].__setitem__('C2', '2'),
                      lambda w: w['Correções'].__setitem__('G1', 'Outro cabeçalho'),
                      lambda w: w['Correções'].__setitem__('N2', 'coluna extra'),
                      lambda w: w['Correções'].__setitem__('G2', '=1+1'),
                      lambda w: w['Correções'].__setitem__('H2', '#REF!'),
                      lambda w: w['Correções'].merge_cells('G2:G3'),
                      lambda w: w['Resumo'].__setitem__('B4', '=1+1')]:
            recusa(lambda: preparar_previa(leitura, gravar(pasta, relatorio, mudar)))
        recusa(lambda: preparar_previa(leitura, gravar(pasta, relatorio, lambda w: w['Resumo'].__setitem__('B4', 'a'*64))), 'BASE_FISCAL_DIVERGENTE')

        # Erros de propostas não usam fallback nem truncam para parecer válidos.
        for coluna, valor, campo in [('G2', 'inválido', 'documento'), ('G2', D1[1:], 'documento'),
                                     ('H2', 'x'*161, 'nome'), ('I2', 'Regime inventado', 'regime'),
                                     ('J2', 'x'*161, 'responsavel_fiscal'), ('K2', 'x'*161, 'contato_nome')]:
            def mudar(w):
                resolvido(w); w['Correções'][coluna] = valor
            ruim = preparar_previa(leitura, gravar(pasta, relatorio, mudar))['linhas'][0]
            assert ruim['status'] == 'pendente', coluna
            assert ruim['proposto'][campo] == valor
            assert ruim['pendencias'], coluna
            assert ruim['efetivo'][campo] != ruim['original'][campo], coluna
        for coluna, valor in [('L2', 'automático'), ('M4', ''), ('M4', 'x'*2001)]:
            def mudar(w):
                resolvido(w); w['Correções'][coluna] = valor
            ruim = preparar_previa(leitura, gravar(pasta, relatorio, mudar))
            assert ruim['linhas'][0]['status'] == 'pendente' or ruim['linhas'][2]['status'] == 'pendente'
        # Correção pode criar colisão com uma linha antes independente.
        def colidir(w):
            resolvido(w); w['Correções']['G3'] = D3
        ruim = preparar_previa(leitura, gravar(pasta, relatorio, colidir))
        assert ruim['linhas'][1]['status'] == ruim['linhas'][3]['status'] == 'pendente'
        # Ou resolver o grupo ao corrigir o documento, sem fundir dados.
        def separar(w):
            w['Correções']['G4'] = D2
        solto = preparar_previa(leitura, gravar(pasta, relatorio, separar))
        assert solto['linhas'][0]['status'] == solto['linhas'][2]['status'] == 'incluir'

        # Um número inteiro hostil é valor inválido, não erro 500 por overflow.
        def numero_hostil(w):
            resolvido(w); w['Correções']['G2'] = 123
        p = gravar(pasta, relatorio, numero_hostil)
        editar_xml(p, 'xl/worksheets/sheet4.xml',
                   lambda b: b.replace(b'<v>123</v>', b'<v>' + b'9' * 400 + b'</v>'))
        hostil = preparar_previa(leitura, p)['linhas'][0]
        assert hostil['status'] == 'pendente' and len(hostil['proposto']['documento']) == 400

        # Fórmula na fonte continua pendente até ser proposta explicitamente.
        formula = deepcopy(leitura)
        formula.linhas[0].origem['contato_nome']['formula'] = True
        fp = gravar(pasta, gerar_relatorio_base_fiscal(formula, nome_arquivo='fonte.xlsx'), resolvido)
        assert preparar_previa(formula, fp)['linhas'][0]['status'] == 'pendente'
        def substituir(w):
            resolvido(w); w['Correções']['K2'] = 'Contato revisado'
        fp = gravar(pasta, gerar_relatorio_base_fiscal(formula, nome_arquivo='fonte.xlsx'), substituir)
        assert preparar_previa(formula, fp)['linhas'][0]['status'] == 'incluir'

        # As demais abas não influenciam propostas, mas continuam sob validação estrutural.
        def detalhe(w):
            resolvido(w); w['Conferência']['D2'] = 'não é autoridade'; w['Origens']['A15000'] = 'origem sintética'
        r = preparar_previa(leitura, gravar(pasta, relatorio, detalhe))
        assert r['linhas'][0]['efetivo']['nome'] == 'Empresa A'
        def exceder(w): w['Origens']['A50002'] = 'excesso'
        recusa(lambda: preparar_previa(leitura, gravar(pasta, relatorio, exceder)))
        # Dimensão falsa não apaga linhas reais.
        p = gravar(pasta, relatorio, resolvido)
        editar_xml(p, 'xl/worksheets/sheet4.xml', lambda b: __import__('re').sub(rb'<dimension[^>]*/>', b'<dimension ref="A1:A1"/>', b))
        assert preparar_previa(leitura, p)['contagens']['linhas'] == 4
        editar_xml(p, 'xl/worksheets/sheet4.xml', lambda b: b.replace(b'<row r="2">', b'<row r="2"><c r="A2" t="inlineStr"><is><t>duplicada</t></is></c>', 1))
        recusa(lambda: preparar_previa(leitura, p))
        # DTD/macros/vínculos não chegam ao leitor de alto nível.
        for nome, conteudo in [('xl/vbaProject.bin', b'invalido'), ('xl/externalLinks/externalLink1.xml', b'<x/>')]:
            p = gravar(pasta, relatorio)
            with ZipFile(p, 'a') as z: z.writestr(nome, conteudo)
            recusa(lambda: preparar_previa(leitura, p))
        p = gravar(pasta, relatorio)
        editar_xml(p, 'xl/workbook.xml', lambda b: b'<!DOCTYPE workbook [<!ENTITY a "x">]>' + b)
        recusa(lambda: preparar_previa(leitura, p))
        # Arquivos ficam liberados depois de sucesso/erro.
        p.unlink()
    print('OK - correções Fiscal: origem/hash, propostas, duplicidades, cadastro, limites e pacotes seguros')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
