"""Relatório da fonte Fiscal, sem cadastro nem decisões de saneamento (019).

O parser já validou a fonte; este módulo conserva seus valores e identificadores.
As propostas humanas ficam em colunas vazias, separadas dos valores lidos.
Strings são células de texto explícitas, sem alterar o conteúdo com apóstrofos.
"""
from __future__ import annotations

from io import BytesIO
import re
import sys
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter

from ..core.clientes import ErroClientes
from ..core.planilha_empresas import LeituraEmpresas, LinhaEmpresa


_AUSENTE = "Coluna ausente no layout"
_XML_INVALIDO = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_FORMULA = "[fórmula não lida]"
_CABECALHOS = {
    "Resumo": ("Item", "Valor"),
    "Conferência": (
        "ID de origem", "Aba de origem", "Linha de origem", "Empresa",
        "CNPJ informado", "CNPJ normalizado", "Situação do CNPJ",
        "Regime informado", "Regime reconhecido", "Responsável Fiscal informado",
        "Contato da empresa", "Sem movimento", "CNPJ repetido", "Pendências de leitura",
        "Colunas ausentes no layout", "Linha oculta na fonte", "Aba oculta na fonte"),
    "Duplicidades": (
        "ID de origem", "CNPJ normalizado", "Aba de origem", "Linha de origem",
        "Empresa", "Regime informado", "Regime reconhecido", "Responsável Fiscal informado",
        "Contato da empresa", "Sem movimento", "Pendências de leitura"),
    "Correções": (
        "ID de origem", "Aba de origem", "Linha de origem", "Empresa original",
        "CNPJ informado original", "CNPJ normalizado original", "CNPJ proposto",
        "Empresa proposta", "Regime proposto", "Responsável Fiscal proposto",
        "Contato proposto", "Decisão sobre duplicidade", "Observações da revisão"),
    "Origens": (
        "ID de origem", "Aba de origem", "Linha de origem", "Campo", "Célula",
        "Célula de origem", "Valor informado", "Formato da célula", "Intervalo mesclado",
        "Propagado", "Fórmula não lida", "Linha oculta na fonte", "Aba oculta na fonte"),
}
_ROTULOS_CONTAGENS = (
    ("linhas", "Registros no arquivo"),
    ("documentos_validos", "Registros com CNPJ válido"),
    ("cnpjs_distintos", "CNPJs válidos distintos"),
    ("documentos_pendentes", "Documentos pendentes"),
    ("grupos_duplicados", "Grupos de CNPJ repetido"),
    ("linhas_duplicadas", "Ocorrências de CNPJ repetido"),
    ("sem_movimento", "Registros sem movimento"),
    ("sem_responsavel", "Responsável ausente em coluna existente"),
    ("sem_regime", "Regime pendente em coluna existente"),
    ("com_pendencias", "Registros com pendências de leitura"),
)


def _campo(linha: LinhaEmpresa, campo: str, valor: str) -> str:
    return _AUSENTE if campo in linha.campos_ausentes_no_layout else valor


def _sim_nao(valor: bool) -> str:
    return "Sim" if valor else "Não"


def _movimento(linha: LinhaEmpresa) -> str:
    if linha.sem_movimento is None:
        return "Não informado"
    return "Sim — permanece na carteira" if linha.sem_movimento else "Não"


def _resumo(leitura: LeituraEmpresas, nome_arquivo: str):
    yield "Relatório", "Conferência da Base Fiscal — preparação para importação"
    yield "Arquivo de origem", nome_arquivo
    yield "SHA-256 da fonte", leitura.hash_arquivo
    yield "Uso do relatório", ("Material de revisão: este relatório ainda não é aceito como arquivo de "
                                "importação. Nenhum cadastro ou vínculo foi gravado.")
    yield "Como revisar", ("Confira as abas Conferência, Duplicidades e Origens. Anote propostas somente "
                            "nas colunas de revisão da aba Correções, preservando o ID de origem.")
    yield "Duplicidades", ("Cada ocorrência continua separada. Compare todas as linhas do mesmo CNPJ; "
                            "o relatório não escolhe um cadastro nem um responsável.")
    yield "Sem movimento", "Empresas sem movimento permanecem na carteira; não significa baixa, inatividade ou dispensa de obrigações."
    yield "Ausência no layout", "Coluna ausente no layout não é uma célula vazia nem uma pendência de leitura."
    yield "Fórmulas", "Fórmulas da fonte não são executadas nem copiadas; são identificadas pelo marcador [fórmula não lida]."
    yield "Abrangência", "Arquivo inteiro, inclusive linhas e abas ocultas; sem filtros da tela e sem exclusão de duplicidades."
    # Evita materializar novamente todos os dicionários de origem do resumo JSON.
    linhas = leitura.linhas
    contagens = {
        "linhas": len(linhas),
        "documentos_validos": sum(bool(l.documento) for l in linhas),
        "cnpjs_distintos": len({l.documento for l in linhas if l.documento}),
        "documentos_pendentes": sum(not l.documento for l in linhas),
        "grupos_duplicados": len({l.documento for l in linhas if l.duplicado}),
        "linhas_duplicadas": sum(l.duplicado for l in linhas),
        "sem_movimento": sum(l.sem_movimento is True for l in linhas),
        "sem_responsavel": sum("responsavel_fiscal" not in l.campos_ausentes_no_layout and not l.responsavel_fiscal_informado.strip() for l in linhas),
        "sem_regime": sum("regime" not in l.campos_ausentes_no_layout and not l.regime for l in linhas),
        "com_pendencias": sum(bool(l.pendencias) for l in linhas),
    }
    for chave, rotulo in _ROTULOS_CONTAGENS:
        yield rotulo, contagens[chave]
    for indice, aba in enumerate(leitura.abas, 1):
        yield f"Aba reconhecida {indice}", aba["nome"]
        yield f"Layout da aba {indice}", aba["tipo"]
        yield f"Cabeçalho da aba {indice}", aba["linha_cabecalho"]
        yield f"Registros da aba {indice}", aba["total"]
        yield f"Aba {indice} oculta na fonte", _sim_nao(aba["oculta"])
    for indice, aviso in enumerate(leitura.avisos, 1):
        yield f"Aviso {indice}", aviso


def _conferencia(leitura: LeituraEmpresas):
    for linha in leitura.linhas:
        yield (linha.id, linha.aba, linha.linha, linha.nome,
               linha.documento_informado, linha.documento, linha.documento_status,
               _campo(linha, "regime", linha.regime_informado), _campo(linha, "regime", linha.regime),
               _campo(linha, "responsavel_fiscal", linha.responsavel_fiscal_informado),
               _campo(linha, "contato_nome", linha.contato_nome), _movimento(linha),
               _sim_nao(linha.duplicado), "; ".join(linha.pendencias),
               ", ".join(linha.campos_ausentes_no_layout), _sim_nao(linha.oculta), _sim_nao(linha.aba_oculta))


def _duplicidades(leitura: LeituraEmpresas):
    # Uma linha por ocorrência, jamais uma lista quadrática de pares na célula.
    for linha in leitura.linhas:
        if linha.duplicado:
            yield (linha.id, linha.documento, linha.aba, linha.linha, linha.nome,
                   _campo(linha, "regime", linha.regime_informado), _campo(linha, "regime", linha.regime),
                   _campo(linha, "responsavel_fiscal", linha.responsavel_fiscal_informado),
                   _campo(linha, "contato_nome", linha.contato_nome), _movimento(linha),
                   "; ".join(linha.pendencias))


def _correcoes(leitura: LeituraEmpresas):
    for linha in leitura.linhas:
        yield (linha.id, linha.aba, linha.linha, linha.nome, linha.documento_informado,
               linha.documento, None, None, None, None, None, None, None)


def _origens(leitura: LeituraEmpresas):
    for linha in leitura.linhas:
        for campo, origem in linha.origem.items():
            yield (linha.id, linha.aba, linha.linha, campo, origem["celula"],
                   origem["celula_origem"], _FORMULA if origem["formula"] else origem["valor_informado"],
                   origem["formato"], origem["intervalo_mesclado"],
                   _sim_nao(origem["propagado"]), _sim_nao(origem["formula"]),
                   _sim_nao(linha.oculta), _sim_nao(linha.aba_oculta))


def _validar(aba: str, titulos: tuple, linhas) -> None:
    for numero, linha in enumerate(linhas, 2):
        for titulo, valor in zip(titulos, linha):
            if not isinstance(valor, str):
                continue
            motivo = "caractere não representável em XLSX" if _XML_INVALIDO.search(valor) else ""
            if not motivo and len(valor.encode("utf-16-le")) // 2 > 32767:
                motivo = "texto acima do limite de 32.767 caracteres da célula"
            if motivo:
                referencia = f"registro {linha[0]}" if aba != "Resumo" else f"linha {numero}"
                raise ErroClientes("PLANILHA_INVALIDA", f"Não foi possível exportar {aba}, {referencia}, campo {titulo}: {motivo}.")


def gerar_relatorio_base_fiscal(leitura: LeituraEmpresas, *, nome_arquivo: str) -> bytes:
    """Gera as cinco abas em memória sem modificar a fonte nem abrir bancos."""
    geradores = {
        "Resumo": lambda: _resumo(leitura, nome_arquivo),
        "Conferência": lambda: _conferencia(leitura),
        "Duplicidades": lambda: _duplicidades(leitura),
        "Correções": lambda: _correcoes(leitura),
        "Origens": lambda: _origens(leitura),
    }
    # openpyxl trunca strings longas ao criar células. Validar antes da escrita
    # evita esse comportamento e também impede uma saída parcialmente válida.
    for nome, gerar in geradores.items():
        _validar(nome, _CABECALHOS[nome], gerar())

    livro = Workbook(write_only=True)
    fonte = Font(name="Calibri", size=11, color="172C43")
    fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fundo_cabecalho = PatternFill("solid", fgColor="173C55")
    fundo_proposta = PatternFill("solid", fgColor="FFF3CD")
    alinhamento = Alignment(vertical="top", wrap_text=True)
    try:
        for nome, gerar in geradores.items():
            aba = livro.create_sheet(nome)
            titulos = _CABECALHOS[nome]
            aba.freeze_panes = "A2"
            aba.sheet_view.showGridLines = False
            aba.sheet_properties.pageSetUpPr.fitToPage = True
            aba.page_setup.orientation = "landscape"
            aba.page_setup.paperSize = "9"  # A4, também no modo write_only.
            aba.page_setup.fitToWidth = 1
            aba.page_setup.fitToHeight = 0
            aba.print_title_rows = "1:1"
            for indice, titulo in enumerate(titulos, 1):
                largura = 16 if "ID" in titulo or "Linha" in titulo else 28
                if titulo in {"Empresa", "Empresa original", "Empresa proposta", "Valor informado", "Pendências de leitura", "Observações da revisão"}:
                    largura = 48
                if nome == "Resumo":
                    largura = 43 if indice == 1 else 110
                aba.column_dimensions[get_column_letter(indice)].width = largura

            def escrever(valores, *, cabecalho=False):
                celulas = []
                for indice, valor in enumerate(valores, 1):
                    celula = WriteOnlyCell(aba, value=valor)
                    if isinstance(valor, str):
                        celula.data_type = "s"
                        celula.number_format = "@"
                    celula.font = fonte_cabecalho if cabecalho else fonte
                    celula.alignment = alinhamento
                    if cabecalho:
                        celula.fill = fundo_cabecalho
                    elif nome == "Correções" and indice >= 7:
                        celula.fill = fundo_proposta
                        celula.number_format = "@"
                    celulas.append(celula)
                aba.append(celulas)

            aba.row_dimensions[1].height = 34
            escrever(titulos, cabecalho=True)
            total = 1
            for valores in gerar():
                total += 1
                escrever(valores)
            aba.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{total}"
        saida = BytesIO()
        # Workbook.save não fecha o ZIP quando write_data falha. O contexto
        # próprio garante o fechamento também em erros de serialização.
        with ZipFile(saida, "w", ZIP_DEFLATED, allowZip64=True) as pacote:
            ExcelWriter(livro, pacote).write_data()
        return saida.getvalue()
    finally:
        # O modo write_only contém o uso de memória por células. O openpyxl usa
        # XML temporário internamente; normalmente write_data já o remove, mas falhas
        # de escrita também precisam liberar esses arquivos e seus handles.
        erro_original = sys.exc_info()[0] is not None
        falha_limpeza = None
        def tentar_fechar(acao):
            nonlocal falha_limpeza
            try:
                acao()
            except Exception as exc:
                if falha_limpeza is None:
                    falha_limpeza = exc

        for aba in livro.worksheets:
            writer = aba._writer
            if writer is None:
                continue
            # Não chamar aba.close(): ele escreve o rodapé de novo e pode
            # repetir justamente a falha que trouxe a execução até aqui.
            # Mesmo que um fechamento falhe, tentar os demais e a remoção.
            if not aba.closed:
                if aba._rows is not None:
                    tentar_fechar(aba._rows.close)
                tentar_fechar(writer.close)
            try:
                writer.cleanup()
            except FileNotFoundError:
                pass  # write_data já removeu o XML desta aba.
            except Exception as exc:
                if falha_limpeza is None:
                    falha_limpeza = exc
        tentar_fechar(livro.close)
        if falha_limpeza is not None and not erro_original:
            raise falha_limpeza
