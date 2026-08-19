"""Saidas do Patrimonio: relacao em Excel e termo de responsabilidade em PDF.

Modulo PURO: recebe dados ja consultados e um caminho de destino. Nao conhece
FastAPI, sessao nem banco.

A escrita segura do Excel vem de `core/planilha_segura`, a mesma usada pela
conciliacao — descricao de bem e nome de colaborador sao texto digitado por
pessoas, e um valor comecando por `=` viraria formula ao abrir a planilha.
"""

from __future__ import annotations

import os
from datetime import datetime

from openpyxl import Workbook

from ..core.patrimonio import CONSERVACOES, MOVIMENTOS, SITUACOES
from ..core.planilha_segura import (escrever_cabecalho, escrever_moeda,
                                    escrever_texto, salvar_atomico,
                                    sanitizar_texto)


def gerar_relacao(destino: str, *, bens: list[dict], gerado_por: str,
                  filtros: dict | None = None) -> str:
    """Relacao de bens em Excel, com uma aba de metadados."""
    livro = Workbook()
    aba = livro.active
    aba.title = "Bens"
    escrever_cabecalho(aba, [
        "Etiqueta", "Tipo", "Descrição", "Marca", "Modelo", "Nº de série",
        "Conservação", "Situação", "Responsável", "Local",
        "Valor de aquisição", "Data de aquisição", "Nota fiscal"])

    linha = 2
    for bem in bens:
        escrever_texto(aba, linha, 1, bem.get("etiqueta"))
        escrever_texto(aba, linha, 2, bem.get("tipo"))
        escrever_texto(aba, linha, 3, bem.get("descricao"))
        escrever_texto(aba, linha, 4, bem.get("marca"))
        escrever_texto(aba, linha, 5, bem.get("modelo"))
        escrever_texto(aba, linha, 6, bem.get("numero_serie"))
        escrever_texto(aba, linha, 7,
                       CONSERVACOES.get(bem.get("conservacao"), ""))
        escrever_texto(aba, linha, 8, SITUACOES.get(bem.get("situacao"), ""))
        escrever_texto(aba, linha, 9, bem.get("responsavel_nome"))
        escrever_texto(aba, linha, 10, bem.get("local_nome"))
        escrever_moeda(aba, linha, 11, bem.get("valor_aquisicao"))
        escrever_texto(aba, linha, 12, bem.get("data_aquisicao"))
        escrever_texto(aba, linha, 13, bem.get("nota_fiscal"))
        linha += 1

    meta = livro.create_sheet("Metadados")
    escrever_cabecalho(meta, ["Item", "Valor"])
    total_centavos = sum(b.get("valor_aquisicao") or 0 for b in bens)
    linhas_meta = [
        ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Gerado por", gerado_por),
        ("Filtros aplicados",
         ", ".join(f"{k}={v}" for k, v in sorted((filtros or {}).items()))
         or "nenhum"),
        ("Total de bens", len(bens)),
    ]
    numero = 2
    for item, valor in linhas_meta:
        escrever_texto(meta, numero, 1, item)
        if isinstance(valor, int):
            meta.cell(row=numero, column=2).value = valor
        else:
            escrever_texto(meta, numero, 2, valor)
        numero += 1
    escrever_texto(meta, numero, 1, "Valor total de aquisição")
    escrever_moeda(meta, numero, 2, total_centavos)
    meta.column_dimensions["B"].width = 60

    salvar_atomico(livro, destino)
    return destino


# ----------------------------------------------------------------------
# Termo de responsabilidade

TEXTO_COMPROMISSO = (
    "Declaro ter recebido o bem acima descrito, em perfeitas condições de "
    "uso, comprometendo-me a utilizá-lo exclusivamente no exercício das "
    "minhas atividades profissionais, a zelar por sua conservação e a "
    "devolvê-lo à empresa quando solicitado ou ao término do vínculo, nas "
    "mesmas condições, ressalvado o desgaste natural pelo uso regular."
)


def gerar_termo(destino: str, *, bem: dict, responsavel: dict,
                emitido_por: str, empresa: str = "JB Fraga Contabilidade") -> str:
    """Termo de responsabilidade em PDF, pronto para assinatura.

    Usa `fpdf2`, que ja e' dependencia do projeto (vem com o
    brazilfiscalreport, usado no DANFE) — nenhuma biblioteca nova entra por
    causa deste arquivo.
    """
    from fpdf import FPDF

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, _pdf(empresa), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "TERMO DE RESPONSABILIDADE", align="C",
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("Helvetica", "", 10)
    linhas = [
        ("Etiqueta", bem.get("etiqueta", "")),
        ("Tipo", bem.get("tipo", "")),
        ("Descrição", bem.get("descricao", "")),
        ("Marca / modelo",
         " / ".join(x for x in (bem.get("marca"), bem.get("modelo")) if x)),
        ("Nº de série", bem.get("numero_serie") or "—"),
        ("Conservação", CONSERVACOES.get(bem.get("conservacao"), "")),
        ("Responsável", responsavel.get("nome", "")),
        ("Setor", responsavel.get("setor") or "—"),
        ("Sob responsabilidade desde",
         _data(responsavel.get("desde") or responsavel.get("iniciada_em"))),
    ]
    for rotulo, valor in linhas:
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(52, 7, _pdf(rotulo + ":"))
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 7, _pdf(str(valor)), new_x="LMARGIN", new_y="NEXT")

    pdf.ln(4)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, _pdf(TEXTO_COMPROMISSO), align="J",
                   new_x="LMARGIN", new_y="NEXT")

    pdf.ln(16)
    pdf.cell(0, 6, _pdf("_" * 56), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, _pdf(responsavel.get("nome", "")), align="C",
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(0, 6, _pdf("Assinatura do responsável"), align="C",
             new_x="LMARGIN", new_y="NEXT")

    pdf.ln(8)
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(0, 5, _pdf(
        f"Emitido por {emitido_por} em "
        f"{datetime.now().strftime('%d/%m/%Y %H:%M')}."),
        align="C", new_x="LMARGIN", new_y="NEXT")

    # Salvamento atomico pelo mesmo motivo do Excel: um PDF truncado abre com
    # erro na mao de quem confiou nele.
    parcial = f"{destino}.parcial"
    try:
        pdf.output(parcial)
        os.replace(parcial, destino)
    except Exception:
        for caminho in (parcial, destino):
            if os.path.exists(caminho):
                os.remove(caminho)
        raise
    return destino


def _pdf(texto) -> str:
    """Texto seguro para as fontes core do PDF (latin-1).

    O `sanitizar_texto` cuida da injecao de formula; aqui o problema e' outro:
    a fonte Helvetica embutida nao tem glifo para tudo, e um caractere fora da
    tabela quebraria a geracao inteira.
    """
    limpo = sanitizar_texto(texto)
    return limpo.encode("latin-1", "replace").decode("latin-1")


def _data(iso: str) -> str:
    if not iso:
        return "—"
    try:
        return datetime.fromisoformat(str(iso)).strftime("%d/%m/%Y")
    except ValueError:
        return str(iso)
