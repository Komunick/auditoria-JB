"""Saida do cadastro de clientes em Excel (spec 012, FR-015).

Modulo PURO: recebe dados ja consultados e um caminho. Por construcao NUNCA
recebe senha — o que chega e' so a lista de quais senhas existem — e por isso
nao tem como escreve-la. Nome de cliente e observacao passam pela escrita
segura de `core/planilha_segura` (texto comecando por `=` nao vira formula).
"""

from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook

from ..core import clientes as cl
from ..core.planilha_segura import (escrever_cabecalho, escrever_texto,
                                    salvar_atomico)
from .clientes_store import CAMPOS_SEGREDO, SITUACOES_CERTIFICADO


def _campo(valor: str, situacao: str) -> str:
    return valor or cl.SITUACOES_CAMPO.get(situacao, "")


def gerar_relacao(destino: str, *, clientes: list[dict],
                  senhas_presentes: dict[int, list[str]],
                  certificados: list[dict], gerado_por: str,
                  filtros: dict | None = None) -> str:
    livro = Workbook()
    aba = livro.active
    aba.title = "Clientes"
    escrever_cabecalho(aba, [
        "Cliente", "CNPJ/CPF", "Situação", "Município", "Inscrição estadual",
        "Inscrição municipal", "IPTU", "Senhas cadastradas",
        "Certificado vence em", "Pendências"])
    for numero, cliente in enumerate(clientes, start=2):
        presentes = senhas_presentes.get(cliente["id"], [])
        valores = [
            cliente["nome"], cl.formatar_documento(cliente.get("documento") or ""),
            cl.SITUACOES_CLIENTE.get(cliente["situacao"], ""),
            cliente.get("municipio", ""),
            _campo(cliente["ie"], cliente["ie_situacao"]),
            _campo(cliente["im"], cliente["im_situacao"]),
            _campo(cliente["iptu"], cliente["iptu_situacao"]),
            ", ".join(CAMPOS_SEGREDO.get(c, c) for c in presentes),
            cliente.get("certificado_fim") or "",
            "; ".join(cliente.get("pendencias") or []),
        ]
        for coluna, valor in enumerate(valores, start=1):
            escrever_texto(aba, numero, coluna, valor)

    certs = livro.create_sheet("Certificados")
    escrever_cabecalho(certs, [
        "Cliente", "Tipo", "Titular", "Emissor", "Início", "Vence em", "Dias",
        "Faixa", "Situação", "Renovação em andamento", "Onde está"])
    for numero, cert in enumerate(certificados, start=2):
        valores = [
            cert.get("cliente_nome"), cert.get("tipo"), cert.get("titular_nome"),
            cert.get("emissor"), cert.get("inicio"), cert.get("fim"), None,
            cl.rotulo_faixa(cert.get("faixa", "")),
            SITUACOES_CERTIFICADO.get(cert.get("situacao"), ""),
            "Sim" if cert.get("renovacao_em_andamento") else "",
            cert.get("local_guarda"),
        ]
        for coluna, valor in enumerate(valores, start=1):
            if coluna == 7:
                certs.cell(row=numero, column=7).value = cert.get("dias")
            else:
                escrever_texto(certs, numero, coluna, valor)

    meta = livro.create_sheet("Metadados")
    escrever_cabecalho(meta, ["Item", "Valor"])
    linhas = [
        ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Gerado por", gerado_por),
        ("Filtros aplicados",
         ", ".join(f"{k}={v}" for k, v in sorted((filtros or {}).items()))
         or "nenhum"),
        ("Total de clientes", str(len(clientes))),
        ("Senhas", "Este arquivo não contém senhas: elas só são vistas no "
                   "sistema, uma por vez, com registro de quem consultou."),
    ]
    for numero, (item, valor) in enumerate(linhas, start=2):
        escrever_texto(meta, numero, 1, item)
        escrever_texto(meta, numero, 2, valor)
    meta.column_dimensions["B"].width = 80

    salvar_atomico(livro, destino)
    return destino
