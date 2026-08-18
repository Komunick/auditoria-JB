"""Saidas Excel da Conciliacao Fiscal: consolidado e planilha-mestre.

Modulo PURO: recebe dados ja consultados e um caminho de destino, devolve
arquivo. Nao conhece FastAPI, sessao nem banco (constituicao, principio II).

Duas regras atravessam o arquivo inteiro:

1. **Texto de terceiro nunca vira formula.** Razao social, nome de arquivo e
   justificativa sao digitados por pessoas ou vem da SEFAZ. Um valor que
   comece por `=`, `+`, `-` ou `@` e' formula para o Excel — abrir a planilha
   executaria conteudo de origem externa. Todo texto passa por
   `sanitizar_texto` antes de ser escrito.
2. **Dinheiro e' celula numerica**, com formato monetario brasileiro, nunca
   string. String impede somar, ordenar e filtrar no proprio Excel, que e'
   exatamente o que a pessoa vai fazer com o consolidado.

O arquivo so aparece para quem pediu depois de salvo por inteiro: falha no
meio nao pode devolver bytes parciais com cara de planilha valida.
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..core.planilha_segura import (FORMATO_MOEDA, cnpj_formatado,
                                    escrever_cabecalho, escrever_moeda,
                                    escrever_texto, salvar_atomico,
                                    sanitizar_texto)

# Os helpers de escrita segura moram em `core/planilha_segura` porque o
# Patrimonio precisa das MESMAS garantias. Duplicar um controle de
# seguranca e' como uma das copias envelhece: alguem corrige um lado,
# esquece o outro, e a vulnerabilidade volta pela porta que ninguem
# estava olhando. Os aliases abaixo mantem os nomes ja usados no modulo.
_moeda = escrever_moeda
_texto = escrever_texto
_cabecalho = escrever_cabecalho
_cnpj_texto = cnpj_formatado
_salvar_atomico = salvar_atomico


ROTULO_ESTADO = {"em_revisao": "Em revisão", "aprovada": "Aprovada",
                 "rejeitada": "Rejeitada"}
ROTULO_LAYOUT = {"quadro_50_5": "Quadro 50-5",
                 "quadro_50_3_dimp": "Quadro 50-3/DIMP"}


def gerar_consolidado(destino: str, *, conciliacoes: list[dict],
                      gerado_por: str, filtros: dict,
                      gerado_em: str = "", versao_parser: str = "") -> str:
    """Consolidado auditavel: dados, historico e proveniencia num arquivo.

    `conciliacoes` sao detalhes completos (o que `store.detalhar` devolve),
    ja filtrados e autorizados por quem chamou. Este modulo nao decide quem
    pode ver o que.
    """
    livro = Workbook()

    # --- Conciliacoes: a visao corrente, uma linha por versao vigente
    aba = livro.active
    aba.title = "Conciliações"
    _cabecalho(aba, [
        "Competência", "CNPJ", "Razão social", "Estado da revisão",
        "Versão", "Layout", "Receita declarada", "Declarada com ST",
        "Declarada sem ST", "Receita calculada", "Calculada com ST",
        "Calculada sem ST", "Diferença de receita", "Diferença com ST",
        "Diferença sem ST", "PIX", "Cartão/outros (não PIX)",
        "Avisos abertos", "Bloqueios abertos", "Conflitos abertos",
        "Arquivo-fonte", "SHA-256"])
    linha = 2
    for c in conciliacoes:
        v = c.get("versao_vigente") or {}
        fonte = c.get("fonte") or {}
        _texto(aba, linha, 1, c["competencia"])
        _texto(aba, linha, 2, _cnpj_texto(c["cnpj"]))
        _texto(aba, linha, 3, c.get("razao_social"))
        _texto(aba, linha, 4, ROTULO_ESTADO.get(c["estado"], c["estado"]))
        aba.cell(row=linha, column=5).value = v.get("numero")
        _texto(aba, linha, 6, ROTULO_LAYOUT.get(v.get("layout"),
                                                v.get("layout")))
        for deslocamento, campo in enumerate((
                "receita_declarada", "declarada_com_st", "declarada_sem_st",
                "receita_calculada", "calculada_com_st", "calculada_sem_st",
                "diferenca_receita", "diferenca_com_st", "diferenca_sem_st",
                "total_pix", "total_nao_pix")):
            _moeda(aba, linha, 7 + deslocamento, v.get(campo))
        aba.cell(row=linha, column=18).value = c.get("avisos_abertos", 0)
        aba.cell(row=linha, column=19).value = c.get("bloqueios_abertos", 0)
        aba.cell(row=linha, column=20).value = c.get("conflitos_abertos", 0)
        _texto(aba, linha, 21, fonte.get("nome_original"))
        _texto(aba, linha, 22, fonte.get("sha256"))
        linha += 1

    # --- Versoes: TODAS, inclusive substituidas e descartadas
    versoes = livro.create_sheet("Versões")
    _cabecalho(versoes, [
        "Competência", "CNPJ", "Versão", "Situação da versão", "Layout",
        "Receita declarada", "Declarada com ST", "Declarada sem ST",
        "Receita calculada", "Calculada com ST", "Calculada sem ST",
        "Diferença de receita", "PIX", "Cartão/outros (não PIX)",
        "Aba", "Linha", "Criada em", "versao_id"])
    linha = 2
    for c in conciliacoes:
        for v in c.get("versoes", []):
            _texto(versoes, linha, 1, c["competencia"])
            _texto(versoes, linha, 2, _cnpj_texto(c["cnpj"]))
            versoes.cell(row=linha, column=3).value = v.get("numero")
            _texto(versoes, linha, 4, v.get("estado"))
            _texto(versoes, linha, 5, ROTULO_LAYOUT.get(v.get("layout"),
                                                        v.get("layout")))
            for deslocamento, campo in enumerate((
                    "receita_declarada", "declarada_com_st",
                    "declarada_sem_st", "receita_calculada",
                    "calculada_com_st", "calculada_sem_st",
                    "diferenca_receita", "total_pix", "total_nao_pix")):
                _moeda(versoes, linha, 6 + deslocamento, v.get(campo))
            _texto(versoes, linha, 15, v.get("aba_origem"))
            versoes.cell(row=linha, column=16).value = v.get("linha_resumo")
            _texto(versoes, linha, 17, v.get("criada_em"))
            versoes.cell(row=linha, column=18).value = v.get("id")
            linha += 1

    # --- Movimentos DIMP: uma linha por movimento ORIGINAL
    movimentos = livro.create_sheet("Movimentos DIMP")
    _cabecalho(movimentos, [
        "Competência", "CNPJ", "versao_id", "Instituição", "Linha de origem",
        "Débito", "Crédito", "Transferência", "PIX", "Voucher", "Outras",
        "Total DIMP", "Total não PIX", "Arquivo", "SHA-256", "Aba",
        "Célula do total"])
    linha = 2
    for c in conciliacoes:
        fonte = c.get("fonte") or {}
        for m in c.get("movimentos", []):
            _texto(movimentos, linha, 1, c["competencia"])
            _texto(movimentos, linha, 2, _cnpj_texto(c["cnpj"]))
            movimentos.cell(row=linha, column=3).value = m.get("versao_id")
            _texto(movimentos, linha, 4, m.get("instituicao"))
            movimentos.cell(row=linha, column=5).value = m.get("linha_origem")
            for deslocamento, campo in enumerate((
                    "debito", "credito", "transferencia", "pix", "voucher",
                    "outras", "total_dimp", "total_nao_pix")):
                _moeda(movimentos, linha, 6 + deslocamento, m.get(campo))
            _texto(movimentos, linha, 14, fonte.get("nome_original"))
            _texto(movimentos, linha, 15, fonte.get("sha256"))
            _texto(movimentos, linha, 16, m.get("aba"))
            _texto(movimentos, linha, 17, m.get("celula_total"))
            linha += 1

    # --- Proveniencia: uma linha por fato
    proveniencia = livro.create_sheet("Proveniência")
    _cabecalho(proveniencia, [
        "Competência", "CNPJ", "versao_id", "Métrica", "Valor",
        "Tipo de origem", "Arquivo", "SHA-256", "Aba", "Célula", "Rótulo",
        "Regra do parser", "Versão do parser", "Fórmula", "Fatos de origem"])
    linha = 2
    for c in conciliacoes:
        fonte = c.get("fonte") or {}
        for f in c.get("fatos", []):
            _texto(proveniencia, linha, 1, c["competencia"])
            _texto(proveniencia, linha, 2, _cnpj_texto(c["cnpj"]))
            proveniencia.cell(row=linha, column=3).value = f.get("versao_id")
            _texto(proveniencia, linha, 4, f.get("metrica"))
            _moeda(proveniencia, linha, 5, f.get("valor_numerico"))
            _texto(proveniencia, linha, 6, f.get("tipo_origem"))
            _texto(proveniencia, linha, 7, fonte.get("nome_original"))
            _texto(proveniencia, linha, 8, fonte.get("sha256"))
            _texto(proveniencia, linha, 9, f.get("aba"))
            _texto(proveniencia, linha, 10, f.get("celula"))
            _texto(proveniencia, linha, 11, f.get("rotulo"))
            _texto(proveniencia, linha, 12, f.get("regra_parser"))
            _texto(proveniencia, linha, 13, f.get("versao_parser"))
            _texto(proveniencia, linha, 14, f.get("formula"))
            _texto(proveniencia, linha, 15,
                   ", ".join(f.get("fatos_origem") or []))
            linha += 1

    # --- Excecoes
    excecoes = livro.create_sheet("Exceções")
    _cabecalho(excecoes, [
        "Competência", "CNPJ", "versao_id", "Código", "Severidade",
        "Mensagem", "Situação", "Resolução", "Resolvida por", "Criada em",
        "Resolvida em"])
    linha = 2
    for c in conciliacoes:
        for e in c.get("excecoes", []):
            _texto(excecoes, linha, 1, c["competencia"])
            _texto(excecoes, linha, 2, _cnpj_texto(c["cnpj"]))
            excecoes.cell(row=linha, column=3).value = e.get("versao_id")
            _texto(excecoes, linha, 4, e.get("codigo"))
            _texto(excecoes, linha, 5, e.get("severidade"))
            _texto(excecoes, linha, 6, e.get("mensagem"))
            _texto(excecoes, linha, 7, e.get("estado"))
            _texto(excecoes, linha, 8, e.get("resolucao"))
            _texto(excecoes, linha, 9, e.get("resolvida_por_login"))
            _texto(excecoes, linha, 10, e.get("criada_em"))
            _texto(excecoes, linha, 11, e.get("resolvida_em"))
            linha += 1

    # --- Conflitos
    conflitos = livro.create_sheet("Conflitos")
    _cabecalho(conflitos, [
        "Competência", "CNPJ", "Situação", "Versão vigente (id)",
        "Versão candidata (id)", "Justificativa", "Resolvido por",
        "Criado em", "Resolvido em"])
    linha = 2
    for c in conciliacoes:
        for k in c.get("conflitos", []):
            _texto(conflitos, linha, 1, c["competencia"])
            _texto(conflitos, linha, 2, _cnpj_texto(c["cnpj"]))
            _texto(conflitos, linha, 3, k.get("estado"))
            conflitos.cell(row=linha, column=4).value = k.get("versao_vigente_id")
            conflitos.cell(row=linha, column=5).value = k.get("versao_candidata_id")
            _texto(conflitos, linha, 6, k.get("justificativa"))
            _texto(conflitos, linha, 7, k.get("resolvido_por_login"))
            _texto(conflitos, linha, 8, k.get("criado_em"))
            _texto(conflitos, linha, 9, k.get("resolvido_em"))
            linha += 1

    # --- Revisoes
    revisoes = livro.create_sheet("Revisões")
    _cabecalho(revisoes, [
        "Competência", "CNPJ", "versao_id", "Revisão anterior",
        "Estado anterior", "Estado novo", "Justificativa", "Usuário",
        "Quando"])
    linha = 2
    for c in conciliacoes:
        for r in c.get("revisoes", []):
            _texto(revisoes, linha, 1, c["competencia"])
            _texto(revisoes, linha, 2, _cnpj_texto(c["cnpj"]))
            revisoes.cell(row=linha, column=3).value = r.get("versao_id")
            revisoes.cell(row=linha, column=4).value = r.get("revision_anterior")
            _texto(revisoes, linha, 5, r.get("estado_anterior"))
            _texto(revisoes, linha, 6, r.get("estado_novo"))
            _texto(revisoes, linha, 7, r.get("justificativa"))
            _texto(revisoes, linha, 8, r.get("usuario_login"))
            _texto(revisoes, linha, 9, r.get("criada_em"))
            linha += 1

    # --- Metadados
    meta = livro.create_sheet("Metadados")
    _cabecalho(meta, ["Item", "Valor"])
    estados = {c["estado"] for c in conciliacoes}
    linhas_meta = [
        ("Gerado em", gerado_em or datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Gerado por", gerado_por),
        ("Filtros aplicados",
         ", ".join(f"{k}={v}" for k, v in sorted(filtros.items())) or "nenhum"),
        ("Total de conciliações", len(conciliacoes)),
        ("Total de versões", sum(len(c.get("versoes", []))
                                 for c in conciliacoes)),
        ("Total de movimentos DIMP", sum(len(c.get("movimentos", []))
                                         for c in conciliacoes)),
        ("Total de fatos", sum(len(c.get("fatos", [])) for c in conciliacoes)),
        ("Total de exceções", sum(len(c.get("excecoes", []))
                                  for c in conciliacoes)),
        ("Total de conflitos", sum(len(c.get("conflitos", []))
                                   for c in conciliacoes)),
        ("Total de revisões", sum(len(c.get("revisoes", []))
                                  for c in conciliacoes)),
        ("Versão do parser", versao_parser),
        ("Semântica de 'Cartão/outros (não PIX)'",
         "Total DIMP menos PIX. Inclui voucher, transferência e outras "
         "operações; não é somente cartão."),
        ("Aviso",
         "Conforme os filtros escolhidos, este consolidado PODE conter "
         "competências não aprovadas. Confira a coluna 'Estado da revisão'."),
        ("Estados presentes",
         ", ".join(sorted(ROTULO_ESTADO.get(e, e) for e in estados)) or "—"),
    ]
    for numero, (item, valor) in enumerate(linhas_meta, start=2):
        _texto(meta, numero, 1, item)
        if isinstance(valor, int):
            meta.cell(row=numero, column=2).value = valor
        else:
            _texto(meta, numero, 2, valor)
    meta.column_dimensions["B"].width = 70
    for linha_meta in meta.iter_rows(min_row=2, min_col=2, max_col=2):
        linha_meta[0].alignment = Alignment(wrap_text=True, vertical="top")

    _salvar_atomico(livro, destino)
    return destino


# ----------------------------------------------------------------------
# Planilha-mestre

class ModeloIncompativel(Exception):
    """O modelo enviado nao pode ser mapeado sem ambiguidade."""

    codigo = "MODELO_INCOMPATIVEL"


MESES = ("JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO")

# Mapa do contrato (contracts/xlsx-outputs.md). Trocar coluna aqui exige
# adaptador novo e testado.
COL_DECLARADA = 3       # C
COL_DECL_COM_ST = 4     # D
COL_DECL_SEM_ST = 5     # E
COL_CALCULADA = 7       # G
COL_DIFERENCA = 8       # H  (fórmula =G-C)
COL_CALC_COM_ST = 9     # I
COL_CALC_SEM_ST = 10    # J  (fórmula =G-I, conforme legado)
COL_PIX = 12            # L
COL_NAO_PIX = 13        # M

_RE_ANO = re.compile(r"relat[oó]rio\s+(\d{4})", re.IGNORECASE)


def _normalizar(texto) -> str:
    return " ".join(str(texto or "").strip().upper().split())


def mapear_modelo(caminho: str) -> dict:
    """Descobre as secoes anuais e a linha de cada mes.

    Os meses sao reconhecidos pelo TEXTO, nunca por numero fixo de linha: o
    modelo do escritorio muda de layout entre anos e uma posicao fixa
    escreveria receita na linha errada em silencio.
    """
    livro = load_workbook(caminho)
    try:
        mapa: dict[int, dict] = {}
        aba_alvo = None
        for aba in livro.worksheets:
            anos_na_aba: dict[int, dict] = {}
            ano_atual = None
            for linha in aba.iter_rows():
                primeira = _normalizar(linha[0].value if linha else "")
                achado = _RE_ANO.search(str(linha[0].value or ""))
                if achado:
                    ano_atual = int(achado.group(1))
                    if ano_atual in anos_na_aba:
                        raise ModeloIncompativel(
                            f"O modelo tem mais de uma seção para {ano_atual}.")
                    anos_na_aba[ano_atual] = {"meses": {}, "total": None}
                    continue
                if ano_atual is None:
                    continue
                if primeira == "TOTAL":
                    anos_na_aba[ano_atual]["total"] = linha[0].row
                    continue
                if primeira in MESES:
                    mes = MESES.index(primeira) + 1
                    if mes in anos_na_aba[ano_atual]["meses"]:
                        raise ModeloIncompativel(
                            f"O mês {primeira} aparece duas vezes na seção "
                            f"{ano_atual}.")
                    anos_na_aba[ano_atual]["meses"][mes] = linha[0].row
            if anos_na_aba:
                if aba_alvo is not None:
                    raise ModeloIncompativel(
                        "Mais de uma aba do modelo tem seções anuais; não dá "
                        "para escolher sem ambiguidade.")
                aba_alvo, mapa = aba.title, anos_na_aba
        if not mapa:
            raise ModeloIncompativel(
                "Nenhuma seção 'RELATÓRIO AAAA' com meses foi encontrada no "
                "modelo.")
        return {"aba": aba_alvo, "anos": mapa}
    finally:
        livro.close()


def cnpj_do_modelo(caminho: str) -> str:
    livro = load_workbook(caminho)
    try:
        padrao = re.compile(r"\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}")
        for aba in livro.worksheets:
            for linha in aba.iter_rows(max_row=30):
                for celula in linha:
                    achado = padrao.search(str(celula.value or ""))
                    if achado:
                        return "".join(c for c in achado.group(0)
                                       if c.isdigit())
        return ""
    finally:
        livro.close()


def preencher_modelo(modelo: str, destino: str, *, competencias: dict,
                     detalhes_por_competencia: dict | None = None,
                     pendentes: list | None = None) -> dict:
    """Preenche uma COPIA do modelo. O original nunca e' tocado.

    `competencias` mapeia 'AAAAMM' -> dict de centavos da versao vigente.
    Celulas fora do mapa, formulas, estilos, larguras, mesclagens e demais
    abas ficam exatamente como estavam.
    """
    mapa = mapear_modelo(modelo)
    livro = load_workbook(modelo)
    try:
        aba = livro[mapa["aba"]]
        escritas = []
        detalhes = detalhes_por_competencia or {}

        for competencia, valores in sorted(competencias.items()):
            ano, mes = int(competencia[:4]), int(competencia[4:])
            secao = mapa["anos"].get(ano)
            if secao is None or mes not in secao["meses"]:
                raise ModeloIncompativel(
                    f"O modelo não tem linha para a competência {competencia}.")
            linha = secao["meses"][mes]

            _moeda(aba, linha, COL_DECLARADA, valores.get("receita_declarada"))
            _moeda(aba, linha, COL_DECL_COM_ST, valores.get("declarada_com_st"))
            _moeda(aba, linha, COL_DECL_SEM_ST, valores.get("declarada_sem_st"))
            _moeda(aba, linha, COL_CALCULADA, valores.get("receita_calculada"))
            _moeda(aba, linha, COL_CALC_COM_ST, valores.get("calculada_com_st"))
            # H e J sao FORMULAS no legado: quem abre a planilha ve a conta,
            # nao um numero solto que ninguem sabe de onde veio.
            aba.cell(row=linha, column=COL_DIFERENCA).value = (
                f"=G{linha}-C{linha}")
            aba.cell(row=linha, column=COL_DIFERENCA).number_format = FORMATO_MOEDA
            aba.cell(row=linha, column=COL_CALC_SEM_ST).value = (
                f"=G{linha}-I{linha}")
            aba.cell(row=linha, column=COL_CALC_SEM_ST).number_format = FORMATO_MOEDA

            detalhe = detalhes.get(competencia) or {}
            _texto(aba, linha, COL_PIX, detalhe.get("pix_texto", ""))
            _texto(aba, linha, COL_NAO_PIX, detalhe.get("nao_pix_texto", ""))
            for coluna in (COL_PIX, COL_NAO_PIX):
                aba.cell(row=linha, column=coluna).alignment = Alignment(
                    wrap_text=True, vertical="top")
            escritas.append(competencia)

        # Totais anuais: SUM só quando a seção tem os 12 meses e uma linha
        # TOTAL. Somar seção incompleta produziria um total que não é total.
        for ano, secao in mapa["anos"].items():
            if secao["total"] is None or len(secao["meses"]) != 12:
                continue
            linhas = sorted(secao["meses"].values())
            primeira, ultima = linhas[0], linhas[-1]
            for coluna in (COL_DECLARADA, COL_DECL_COM_ST, COL_DECL_SEM_ST,
                           COL_CALCULADA, COL_DIFERENCA, COL_CALC_COM_ST,
                           COL_CALC_SEM_ST):
                letra = get_column_letter(coluna)
                celula = aba.cell(row=secao["total"], column=coluna)
                celula.value = f"=SUM({letra}{primeira}:{letra}{ultima})"
                celula.number_format = FORMATO_MOEDA

        # Excel recalcula tudo ao abrir: as fórmulas escritas aqui não têm
        # valor em cache, e sem isto apareceriam vazias.
        livro.calculation.fullCalcOnLoad = True

        if pendentes:
            _marcar_pendentes(livro, pendentes)

        _salvar_atomico(livro, destino)
    finally:
        livro.close()

    return {"destino": destino, "competencias": escritas,
            "pendentes": list(pendentes or [])}


def _marcar_pendentes(livro: Workbook, pendentes: list) -> None:
    """Aviso VISIVEL de que o arquivo tem competencia nao aprovada.

    Comentario oculto nao serve: quem recebe a planilha por e-mail precisa
    ver a ressalva sem procurar (contrato de saida).
    """
    aviso = livro.create_sheet("Metadados da automação", 0)
    _cabecalho(aviso, ["Atenção", "Detalhe"])
    _texto(aviso, 2, 1, "Competências NÃO aprovadas incluídas")
    _texto(aviso, 2, 2, ", ".join(sorted(str(p) for p in pendentes)))
    _texto(aviso, 3, 1, "O que isso significa")
    _texto(aviso, 3, 2,
           "Estas competências ainda estavam Em revisão quando a planilha foi "
           "gerada. Os valores podem mudar após a revisão fiscal.")
    aviso.column_dimensions["B"].width = 80
    aviso["B2"].alignment = Alignment(wrap_text=True, vertical="top")
    aviso["B3"].alignment = Alignment(wrap_text=True, vertical="top")
    aviso.sheet_properties.tabColor = "C00000"
