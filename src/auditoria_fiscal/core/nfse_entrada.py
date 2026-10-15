"""Entrada do lote da Emissao de NFS-e (spec 006): planilha-modelo e texto
colado da conversa do WhatsApp.

Este modulo so LE. Ele devolve os valores brutos, como chegaram, e de onde
cada um veio (celula e rotulo da coluna, ou linha do texto e rotulo); quem
confere e normaliza e' `core/nfse_conferencia.py`. Separar as duas coisas e'
o que permite responder, depois de uma correcao, "o que a escola mandou?".

Duas garantias:

- **Nada some calado.** Coluna desconhecida, texto antes do primeiro bloco,
  linha sem rotulo e rotulo repetido viram aviso ou bloqueio, nunca descarte.
- **Nada e' adivinhado.** Coluna e' achada pelo rotulo, nao pela posicao;
  turma so e' separada do nome do aluno quando e' turma do perfil; um bloco
  sem competencia nao herda a de outro.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from decimal import Decimal

from . import nfse as nf
from . import planilha_segura

ARQUIVO_RECUSADO = "ARQUIVO_RECUSADO"
PLANILHA_SEM_CABECALHO = "PLANILHA_SEM_CABECALHO"
COLUNA_OBRIGATORIA_AUSENTE = "COLUNA_OBRIGATORIA_AUSENTE"
COLUNA_REPETIDA = "COLUNA_REPETIDA"
PLANILHA_GRANDE_DEMAIS = "PLANILHA_GRANDE_DEMAIS"
TEXTO_GRANDE_DEMAIS = "TEXTO_GRANDE_DEMAIS"
TEXTO_SEM_BLOCOS = "TEXTO_SEM_BLOCOS"

MAX_PLANILHA_BYTES = 10 * 1024 * 1024          # [proposta]
MAX_LINHAS_PLANILHA = 5000                     # [proposta]
MAX_TEXTO_BYTES = 256 * 1024                   # [proposta]
LINHAS_PROCURA_CABECALHO = 10

# Campos brutos de uma linha (data-model.md, tabela `linha`).
CAMPOS = ("competencia", "responsavel_nome", "responsavel_cpf",
          "responsavel_celular", "responsavel_email", "endereco_texto", "cep",
          "logradouro", "numero", "complemento", "bairro", "municipio", "uf",
          "municipio_ibge", "aluno_nome", "turma", "valor")

# Cabecalho da planilha-modelo que a JB entrega a escola, nesta ordem
# (data-model.md, "Planilha-modelo"). O leitor acha as colunas pelo rotulo.
ROTULOS_MODELO = ("Competência (MM/AAAA)", "Nome do responsável",
                  "CPF do responsável", "Celular", "E-mail", "Nome do aluno",
                  "Turma", "Valor (R$)", "CEP", "Logradouro", "Número",
                  "Complemento", "Bairro", "Município", "UF",
                  "Código IBGE do município")

OBRIGATORIAS = {
    "competencia": "Competência (MM/AAAA)",
    "responsavel_nome": "Nome do responsável",
    "responsavel_cpf": "CPF do responsável",
    "aluno_nome": "Nome do aluno",
    "turma": "Turma",
    "valor": "Valor (R$)",
}

# Rotulo normalizado (sem caixa, acento e parenteses) -> campo.
ROTULOS_PLANILHA = {
    "competencia": "competencia",
    "nome do responsavel": "responsavel_nome", "responsavel": "responsavel_nome",
    "cpf do responsavel": "responsavel_cpf", "cpf": "responsavel_cpf",
    "celular": "responsavel_celular", "telefone": "responsavel_celular",
    "fone": "responsavel_celular",
    "e-mail": "responsavel_email", "email": "responsavel_email",
    "nome do aluno": "aluno_nome", "aluno": "aluno_nome", "aluna": "aluno_nome",
    "turma": "turma",
    "valor": "valor",
    "cep": "cep",
    "endereco": "endereco_texto",
    "logradouro": "logradouro", "rua": "logradouro",
    "numero": "numero", "no": "numero", "n": "numero",
    "complemento": "complemento",
    "bairro": "bairro",
    "municipio": "municipio", "cidade": "municipio",
    "uf": "uf", "estado": "uf",
    "codigo ibge do municipio": "municipio_ibge", "codigo ibge": "municipio_ibge",
    "ibge": "municipio_ibge",
}

ROTULOS_TEXTO = {
    "competencia": "competencia",
    "responsavel": "responsavel_nome", "nome do responsavel": "responsavel_nome",
    "cpf": "responsavel_cpf", "cpf do responsavel": "responsavel_cpf",
    "celular": "responsavel_celular", "telefone": "responsavel_celular",
    "fone": "responsavel_celular", "whatsapp": "responsavel_celular",
    "e-mail": "responsavel_email", "email": "responsavel_email",
    "endereco": "endereco_texto",
    "aluno": "aluno_nome", "aluna": "aluno_nome", "nome do aluno": "aluno_nome",
    "turma": "turma",
    "valor": "valor",
}

# Rotulos que abrem bloco novo no texto colado.
INICIO_DE_BLOCO = ("competencia", "responsavel_nome")
# Com estes ja preenchidos, uma linha em branco seguida de inicio de bloco
# fecha o bloco atual (data-model.md, "Texto colado").
DADOS_DE_BLOCO = ("aluno_nome", "valor", "responsavel_cpf")


@dataclass
class LinhaBruta:
    ordem: int
    campos: dict
    tipos: dict = field(default_factory=dict)
    proveniencia: dict = field(default_factory=dict)
    aba: str | None = None
    linha_planilha: int | None = None
    bloco: int | None = None
    texto_linha_ini: int | None = None
    texto_linha_fim: int | None = None
    nao_reconhecido: tuple[str, ...] = ()
    rotulos_repetidos: tuple[str, ...] = ()


@dataclass
class ResultadoEntrada:
    tipo: str                       # planilha | texto
    linhas: list[LinhaBruta]
    avisos: list[str] = field(default_factory=list)
    ignoradas: int = 0


def _chave_rotulo(texto) -> str:
    chave = nf.chave_comparacao(texto)
    chave = re.sub(r"\(.*?\)", " ", chave)
    return " ".join(chave.strip(" .:-").split())


# ----------------------------------------------------------------------
# Planilha


def _valor_celula(valor) -> tuple[str | None, str]:
    """Celula -> (texto bruto, tipo). Numero vira texto exato; data, ISO."""
    if valor is None:
        return None, "texto"
    if isinstance(valor, bool):
        return str(valor), "texto"
    if isinstance(valor, int):
        return str(valor), "numero"
    if isinstance(valor, float):
        if valor.is_integer():
            return str(int(valor)), "numero"
        return format(Decimal(repr(valor)), "f"), "numero"
    if isinstance(valor, dt.datetime):
        return valor.date().isoformat(), "data"
    if isinstance(valor, dt.date):
        return valor.isoformat(), "data"
    texto = str(valor)
    return (texto if texto.strip() else None), "texto"


def _achar_cabecalho(linhas: list[tuple]) -> tuple[int, dict, list[str]] | None:
    """(indice da linha, {coluna: (campo, rotulo)}, rotulos desconhecidos)."""
    for indice, linha in enumerate(linhas[:LINHAS_PROCURA_CABECALHO]):
        mapa, desconhecidos = {}, []
        for coluna, celula in enumerate(linha):
            texto = str(celula).strip() if celula is not None else ""
            if not texto:
                continue
            campo = ROTULOS_PLANILHA.get(_chave_rotulo(texto))
            if campo:
                mapa[coluna] = (campo, texto)
            else:
                desconhecidos.append(texto)
        if len(mapa) >= 3:
            return indice, mapa, desconhecidos
    return None


def _letra(coluna: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(coluna + 1)


def ler_planilha(caminho: str, *, max_bytes: int | None = MAX_PLANILHA_BYTES,
                 max_linhas: int = MAX_LINHAS_PLANILHA) -> ResultadoEntrada:
    """Le a planilha-modelo (aba `Notas` ou a primeira com cabecalho).

    O pacote e' validado antes do openpyxl. Formula sem valor calculado e'
    ausencia (marcada na proveniencia); linha so com vazios e' ignorada e
    contada; linha parcial entra.
    """
    try:
        planilha_segura.validar_pacote_xlsx(caminho, max_bytes=max_bytes)
    except planilha_segura.PacoteRecusado as exc:
        raise nf.ErroNfse(ARQUIVO_RECUSADO, exc.mensagem) from exc

    from openpyxl import load_workbook

    valores = load_workbook(caminho, read_only=True, data_only=True)
    formulas = load_workbook(caminho, read_only=True, data_only=False)
    try:
        abas = sorted(valores.worksheets,
                      key=lambda a: 0 if nf.chave_comparacao(a.title) == "notas" else 1)
        for aba in abas:
            # A dimensao declarada no arquivo pode mentir (A1:XFD1048576);
            # recalcular faz o leitor percorrer so as linhas que existem.
            aba.reset_dimensions()
            linhas = list(aba.iter_rows(values_only=True))
            achado = _achar_cabecalho(linhas)
            if achado is None:
                continue
            aba_formulas = formulas[aba.title]
            aba_formulas.reset_dimensions()
            brutas = list(aba_formulas.iter_rows(values_only=True))
            return _ler_aba(aba.title, linhas, brutas, achado, max_linhas)
    finally:
        valores.close()
        formulas.close()
    raise nf.ErroNfse(PLANILHA_SEM_CABECALHO,
                      "Não achei o cabeçalho da planilha-modelo nas 10 primeiras "
                      "linhas (Competência, Nome do responsável, CPF do "
                      "responsável, Nome do aluno, Turma, Valor).")


def _ler_aba(titulo: str, linhas: list[tuple], brutas: list[tuple],
             achado: tuple, max_linhas: int) -> ResultadoEntrada:
    indice_cab, mapa, desconhecidos = achado
    vistos: dict[str, str] = {}
    for campo, rotulo in mapa.values():
        if campo in vistos:
            raise nf.ErroNfse(COLUNA_REPETIDA,
                              f"As colunas {vistos[campo]!r} e {rotulo!r} trazem o "
                              "mesmo dado; deixe só uma.")
        vistos[campo] = rotulo
    faltam = [nome for campo, nome in OBRIGATORIAS.items() if campo not in vistos]
    if faltam:
        raise nf.ErroNfse(COLUNA_OBRIGATORIA_AUSENTE,
                          "Falta a coluna obrigatória: " + ", ".join(
                              repr(f) for f in faltam) + ".")

    dados = linhas[indice_cab + 1:]
    # Linhas vazias do fim (formatacao sem dado) nao contam como ignoradas.
    while dados and all(_valor_celula(c)[0] is None for c in dados[-1]) and not any(
            isinstance(c, str) and c.startswith("=")
            for c in (brutas[indice_cab + len(dados)]
                      if indice_cab + len(dados) < len(brutas) else ())):
        dados = dados[:-1]
    if len(dados) > max_linhas:
        raise nf.ErroNfse(PLANILHA_GRANDE_DEMAIS,
                          f"A planilha tem {len(dados)} linhas; o limite é "
                          f"{max_linhas}.")

    resultado = ResultadoEntrada(tipo="planilha", linhas=[])
    if desconhecidos:
        resultado.avisos.append("Colunas ignoradas (não fazem parte da "
                                "planilha-modelo): " + ", ".join(desconhecidos) + ".")
    for deslocamento, linha in enumerate(dados):
        numero_planilha = indice_cab + 2 + deslocamento
        bruta = brutas[indice_cab + 1 + deslocamento] if (
            indice_cab + 1 + deslocamento) < len(brutas) else ()
        campos = {c: None for c in CAMPOS}
        tipos, prov = {}, {}
        tem_dado = False
        for coluna, (campo, rotulo) in mapa.items():
            celula = linha[coluna] if coluna < len(linha) else None
            texto, tipo = _valor_celula(celula)
            info = {"celula": f"{_letra(coluna)}{numero_planilha}",
                    "rotulo": rotulo, "regra": "planilha.coluna_por_rotulo",
                    "tipo_celula": tipo}
            formula = bruta[coluna] if coluna < len(bruta) else None
            if texto is None and isinstance(formula, str) and formula.startswith("="):
                info["formula_sem_valor"] = True
                tem_dado = True
            if texto is not None:
                tem_dado = True
            campos[campo] = texto
            tipos[campo] = tipo
            prov[campo] = info
        if not tem_dado:
            resultado.ignoradas += 1
            continue
        resultado.linhas.append(LinhaBruta(
            ordem=len(resultado.linhas) + 1, campos=campos, tipos=tipos,
            proveniencia=prov, aba=titulo, linha_planilha=numero_planilha))
    return resultado


# ----------------------------------------------------------------------
# Texto colado

_RE_PREFIXO_DATA = re.compile(
    r"^\s*\[?\d{1,2}/\d{1,2}/\d{2,4}[,\s]+\d{1,2}:\d{2}(?::\d{2})?\]?\s*(?:-\s*)?")
_RE_REMETENTE = re.compile(r"^([^:]{1,40}):\s*(.*)$")
_SEPARADORES_ROTULO = (re.compile(r":"), re.compile(r"\s[-–]\s"), re.compile(r"[-–]"))
_RE_CEP = re.compile(r"\b(\d{5})-?(\d{3})\b")


def _rotulo(linha: str):
    """`(campo, rotulo, valor)` se a linha comeca por um rotulo conhecido."""
    for separador in _SEPARADORES_ROTULO:
        achado = separador.search(linha)
        if not achado:
            continue
        rotulo = linha[:achado.start()].strip()
        campo = ROTULOS_TEXTO.get(_chave_rotulo(rotulo))
        if campo:
            return campo, rotulo, linha[achado.end():].strip()
    return None


def _sem_prefixo(linha: str) -> str:
    """Tira "[13/10/2026 10:15] Remetente: " de mensagem copiada, se houver."""
    achado = _RE_PREFIXO_DATA.match(linha)
    if not achado:
        return linha
    resto = linha[achado.end():]
    if _rotulo(resto):
        return resto
    remetente = _RE_REMETENTE.match(resto)
    if remetente and _rotulo(remetente.group(2)):
        return remetente.group(2)
    return resto


class _Bloco:
    def __init__(self, numero: int, inicio: int):
        self.numero = numero
        self.inicio = inicio
        self.fim = inicio
        self.campos = {c: None for c in CAMPOS}
        self.prov: dict = {}
        self.vistos: set[str] = set()
        self.turma_do_aluno = False
        self.nao_reconhecido: list[str] = []
        self.repetidos: list[str] = []

    def guardar(self, campo, valor, linha, rotulo, regra="texto.rotulo"):
        self.campos[campo] = valor or None
        self.prov[campo] = {"linha_texto": linha, "rotulo": rotulo, "regra": regra}
        self.vistos.add(campo)


def ler_texto(texto: str, turmas) -> ResultadoEntrada:
    """Texto colado -> uma linha por bloco (aluno).

    Bloco novo comeca em Competencia ou Responsavel quando o bloco atual ja
    tem esse rotulo, ou depois de linha em branco se o bloco atual ja tem
    aluno, valor ou CPF. Turma escrita junto do aluno so e' separada quando e'
    turma do perfil. Endereco fica em texto livre; o CEP e' extraido.
    """
    if len((texto or "").encode("utf-8")) > MAX_TEXTO_BYTES:
        raise nf.ErroNfse(TEXTO_GRANDE_DEMAIS,
                          f"O texto colado passa de {MAX_TEXTO_BYTES // 1024} KB.")
    blocos: list[_Bloco] = []
    preambulo: list[str] = []
    atual: _Bloco | None = None
    quebra = False
    for numero, original in enumerate((texto or "").splitlines(), start=1):
        linha = _sem_prefixo(original).strip()
        if not linha:
            quebra = True
            continue
        achado = _rotulo(linha)
        if achado is None:
            if atual is None:
                preambulo.append(linha)
            else:
                atual.nao_reconhecido.append(linha)
                atual.fim = numero
            quebra = False
            continue

        campo, rotulo, valor = achado
        abre = atual is None or (campo in INICIO_DE_BLOCO and (
            campo in atual.vistos
            or (quebra and any(c in atual.vistos for c in DADOS_DE_BLOCO))))
        if abre:
            atual = _Bloco(len(blocos) + 1, numero)
            blocos.append(atual)
        quebra = False
        atual.fim = numero

        if campo == "turma" and atual.turma_do_aluno:
            if nf.chave_comparacao(valor) != nf.chave_comparacao(atual.campos["turma"]):
                atual.repetidos.append(rotulo)
            continue
        if campo in atual.vistos:
            atual.repetidos.append(rotulo)
            continue
        if campo == "aluno_nome":
            aluno, turma = nf.separar_turma_do_aluno(valor, turmas)
            if turma is not None:
                atual.guardar("aluno_nome", aluno, numero, rotulo,
                              "texto.aluno_sem_turma")
                if "turma" not in atual.vistos:
                    atual.guardar("turma", turma, numero, rotulo,
                                  "texto.turma_no_aluno")
                    atual.turma_do_aluno = True
                continue
        atual.guardar(campo, valor, numero, rotulo)
        if campo == "endereco_texto":
            cep = _RE_CEP.search(valor or "")
            if cep:
                atual.guardar("cep", cep.group(1) + cep.group(2), numero, rotulo,
                              "texto.cep_no_endereco")

    if not blocos:
        raise nf.ErroNfse(TEXTO_SEM_BLOCOS,
                          "Nenhum bloco reconhecido. Cada aluno precisa de linhas "
                          "como 'Responsável: ...', 'CPF: ...', 'Aluno: ...', "
                          "'Valor: ...'.")
    resultado = ResultadoEntrada(tipo="texto", linhas=[])
    if preambulo:
        resultado.avisos.append("Texto antes do primeiro bloco, sem rótulo: "
                                + " / ".join(preambulo))
    for bloco in blocos:
        resultado.linhas.append(LinhaBruta(
            ordem=bloco.numero, campos=bloco.campos,
            tipos={c: "texto" for c in bloco.vistos},
            proveniencia=bloco.prov, bloco=bloco.numero,
            texto_linha_ini=bloco.inicio, texto_linha_fim=bloco.fim,
            nao_reconhecido=tuple(bloco.nao_reconhecido),
            rotulos_repetidos=tuple(bloco.repetidos)))
    return resultado
