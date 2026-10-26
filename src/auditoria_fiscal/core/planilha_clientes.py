"""Leitura da planilha de senhas que a JB usa hoje (spec 012, US3).

O que a planilha tem (estrutura lida em 22/10/2026, sem exibir valores):
cabecalho na linha 2, nome do cliente na coluna A sem rotulo, CNPJ gravado
como numero, situacao do cliente na COR da linha, uma legenda que divide as
linhas com os clientes, colunas sem cabecalho que a JB confirmou nao terem
nada (23/10/2026) e campos que misturam valor com "NAO TEM"/"PENDENTE".

Este modulo so LE e classifica. Nada e' gravado aqui; a prévia devolvida por
`resumo()` nunca leva senha — quem precisa dela pega `LinhaPlanilha.segredos`
e cifra antes de qualquer outra coisa.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import clientes as cl
from .planilha_segura import PacoteRecusado, validar_pacote_xlsx

PLANILHA_INVALIDA = "PLANILHA_INVALIDA"
CABECALHO_NAO_ENCONTRADO = "CABECALHO_NAO_ENCONTRADO"

MAX_PLANILHA_BYTES = 20 * 1024 * 1024
LINHAS_PROCURA_CABECALHO = 10

ACAO_INCLUIR = "incluir"
ACAO_IGNORAR = "ignorar"

# Rotulo do cabecalho (sem acento, caixa alta) -> campo.
_COLUNAS = {
    "NOME": "nome", "CLIENTE": "nome", "RAZAO SOCIAL": "nome",
    "EMPRESA": "nome",
    "CNPJ": "documento", "CPF": "documento", "CNPJ/CPF": "documento",
    "CPF/CNPJ": "documento",
    "INSCRICAO ESTADUAL": "ie", "IE": "ie",
    "INSCRICAO MUNICIPAL": "im", "IM": "im",
    "SENHA SEFAZ ESTADUAL": "sefaz", "SENHA SEFAZ": "sefaz",
    "SENHA PREFEITURA": "prefeitura",
    "CODIGO ACESSO": "codigo_acesso", "CODIGO DE ACESSO": "codigo_acesso",
    "IPTU": "iptu",
    "TVL": "tvl", "VENC TVL": "venc_tvl",
    "VENC ALVARA": "venc_alvara",
    "PUBLICIDADE": "publicidade",
    "ALVARA SAUDE": "alvara_saude",
}
# Colunas conhecidas que NAO sao dado de cliente.
_IGNORADAS = {"LEGENDA"}

CAMPOS_SEGREDO = ("sefaz", "prefeitura", "codigo_acesso")

# Legenda da planilha (cor de preenchimento da coluna do nome).
COR_BAIXADO = "FFFF0000"
COR_INATIVO = "FFFF9900"
COR_NAO_CLIENTE = "FF999999"
COR_LAURO = "FFF9CB9C"


@dataclass
class LinhaPlanilha:
    linha: int
    nome: str = ""
    documento: str = ""
    documento_informado: str = ""
    ie: str = ""
    ie_situacao: str = ""
    im: str = ""
    im_situacao: str = ""
    iptu: str = ""
    iptu_situacao: str = ""
    situacao: str = cl.SITUACAO_ATIVO
    municipio: str = ""
    cor: str = ""
    oculta: bool = False
    licencas: dict = field(default_factory=dict)
    segredos: dict = field(default_factory=dict)   # campo -> (valor, situacao)
    pendencias: list = field(default_factory=list)
    acao: str = ACAO_INCLUIR
    motivo: str = ""

    def resumo(self) -> dict:
        """O que a tela mostra. Senha vira 'presente' ou a situacao."""
        return {
            "linha": self.linha, "nome": self.nome,
            "documento": cl.formatar_documento(self.documento),
            "documento_informado": self.documento_informado,
            "ie": self.ie, "ie_situacao": self.ie_situacao,
            "im": self.im, "im_situacao": self.im_situacao,
            "situacao": self.situacao,
            "situacao_rotulo": cl.SITUACOES_CLIENTE.get(self.situacao, ""),
            "municipio": self.municipio, "oculta": self.oculta,
            "segredos": {campo: ("presente" if valor else situacao or "")
                         for campo, (valor, situacao) in self.segredos.items()},
            "licencas": self.licencas,
            "pendencias": list(self.pendencias),
            "acao": self.acao, "motivo": self.motivo,
        }


@dataclass
class LeituraPlanilha:
    hash_arquivo: str
    aba: str
    linha_cabecalho: int
    colunas: dict                      # letra -> campo
    colunas_sem_cabecalho: list
    colunas_ignoradas: list
    colunas_desconhecidas: list
    linhas: list

    def resumo(self) -> dict:
        incluir = [l for l in self.linhas if l.acao == ACAO_INCLUIR]
        return {
            "hash_arquivo": self.hash_arquivo, "aba": self.aba,
            "linha_cabecalho": self.linha_cabecalho,
            "colunas_sem_cabecalho": self.colunas_sem_cabecalho,
            "colunas_ignoradas": self.colunas_ignoradas,
            "colunas_desconhecidas": self.colunas_desconhecidas,
            "contagens": {
                "total": len(self.linhas), "incluir": len(incluir),
                "ignorar": len(self.linhas) - len(incluir),
                "com_pendencia": sum(1 for l in incluir if l.pendencias),
            },
            "linhas": [l.resumo() for l in self.linhas],
        }


def _hash(caminho: str) -> str:
    h = hashlib.sha256()
    with open(caminho, "rb") as arquivo:
        for bloco in iter(lambda: arquivo.read(1024 * 1024), b""):
            h.update(bloco)
    return h.hexdigest()


def _cor(celula) -> str:
    preenchimento = celula.fill
    if preenchimento is None or preenchimento.fill_type is None:
        return ""
    cor = preenchimento.fgColor
    if cor is not None and cor.type == "rgb" and isinstance(cor.rgb, str):
        return cor.rgb.upper()
    return "tema"


def _quase_branca(cor: str) -> bool:
    if not cor:
        return True
    if len(cor) != 8:
        return False
    try:
        canais = [int(cor[i:i + 2], 16) for i in (2, 4, 6)]
    except ValueError:
        return False
    return all(c >= 0xF5 for c in canais)


def _aplicar_cor(linha: LinhaPlanilha) -> None:
    cor = linha.cor
    if cor == COR_BAIXADO:
        linha.situacao = cl.SITUACAO_BAIXADO
    elif cor == COR_INATIVO:
        linha.situacao = cl.SITUACAO_INATIVO
    elif cor == COR_NAO_CLIENTE:
        linha.situacao = cl.SITUACAO_NAO_CLIENTE
    elif cor == COR_LAURO:
        linha.municipio = "Lauro de Freitas"
    elif not _quase_branca(cor):
        linha.pendencias.append(
            f"Cor da linha fora da legenda ({cor}): confirme a situação.")


def _licenca(numero, data_ou_situacao) -> dict:
    """Uma licenca: numero (quando a planilha tem) + vencimento ou situacao."""
    registro = {"numero": cl.classificar_campo(numero)[0] if numero is not None
                else "", "vencimento": "", "situacao": "", "texto_original": ""}
    if numero is not None and not registro["numero"]:
        registro["situacao"] = cl.classificar_campo(numero)[1]
    if data_ou_situacao is None or cl.texto_celula(data_ou_situacao) == "":
        return registro
    data = cl.data_plausivel(data_ou_situacao)
    if data is not None:
        registro["vencimento"] = data.isoformat()
        return registro
    valor, situacao = cl.classificar_campo(data_ou_situacao)
    registro["texto_original"] = cl.texto_celula(data_ou_situacao)
    if situacao:
        registro["situacao"] = situacao
    else:
        registro["situacao"] = cl.CAMPO_DATA_INVALIDA
    return registro


def ler_planilha_clientes(caminho: str, *,
                          max_bytes: int = MAX_PLANILHA_BYTES) -> LeituraPlanilha:
    try:
        validar_pacote_xlsx(caminho, max_bytes=max_bytes)
    except PacoteRecusado as exc:
        raise cl.ErroClientes(PLANILHA_INVALIDA, exc.mensagem) from exc
    try:
        livro = load_workbook(caminho, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl levanta de tudo
        raise cl.ErroClientes(PLANILHA_INVALIDA,
                              "Não foi possível abrir a planilha.") from exc
    try:
        return _ler(livro, _hash(caminho))
    finally:
        livro.close()


def _achar_cabecalho(aba) -> int:
    for linha in range(1, min(aba.max_row, LINHAS_PROCURA_CABECALHO) + 1):
        for coluna in range(1, aba.max_column + 1):
            if _COLUNAS.get(cl.chave_texto(aba.cell(linha, coluna).value)) \
                    == "documento":
                return linha
    raise cl.ErroClientes(
        CABECALHO_NAO_ENCONTRADO,
        "Não encontrei a coluna CNPJ nas primeiras linhas da planilha.")


def _ler(livro, hash_arquivo: str) -> LeituraPlanilha:
    aba = livro.worksheets[0]
    cabecalho = _achar_cabecalho(aba)

    colunas: dict[int, str] = {}
    ignoradas: list[int] = []
    desconhecidas: list[str] = []
    sem_rotulo: list[int] = []
    for coluna in range(1, aba.max_column + 1):
        rotulo = cl.chave_texto(aba.cell(cabecalho, coluna).value)
        if not rotulo:
            sem_rotulo.append(coluna)
        elif rotulo in _IGNORADAS:
            ignoradas.append(coluna)
        elif rotulo in _COLUNAS:
            colunas.setdefault(coluna, _COLUNAS[rotulo])
        else:
            desconhecidas.append(f"{get_column_letter(coluna)} ({rotulo})")
            ignoradas.append(coluna)

    # O nome mora na coluna sem rotulo a esquerda do CNPJ (coluna A na
    # planilha da JB) quando nenhum cabecalho diz "NOME".
    if "nome" not in colunas.values():
        coluna_doc = min(c for c, campo in colunas.items() if campo == "documento")
        candidatas = [c for c in sem_rotulo if c < coluna_doc]
        if candidatas:
            colunas[candidatas[-1]] = "nome"
            sem_rotulo.remove(candidatas[-1])

    primeira, ultima = cabecalho + 1, aba.max_row
    sem_cabecalho_com_dado = [
        c for c in sem_rotulo
        if any(cl.texto_celula(aba.cell(l, c).value)
               for l in range(primeira, ultima + 1))]

    linhas: list[LinhaPlanilha] = []
    vistos: dict[str, int] = {}
    for numero in range(primeira, ultima + 1):
        valores = {campo: aba.cell(numero, coluna).value
                   for coluna, campo in colunas.items()}
        if not any(cl.texto_celula(v) for v in valores.values()):
            continue
        linha = _linha(numero, valores, aba)
        if linha.documento:
            if linha.documento in vistos and linha.acao == ACAO_INCLUIR:
                linha.acao = ACAO_IGNORAR
                linha.motivo = (f"Documento repetido: já aparece na linha "
                                f"{vistos[linha.documento]}.")
            vistos.setdefault(linha.documento, numero)
        linhas.append(linha)

    letra = get_column_letter
    return LeituraPlanilha(
        hash_arquivo=hash_arquivo, aba=aba.title, linha_cabecalho=cabecalho,
        colunas={letra(c): campo for c, campo in sorted(colunas.items())},
        colunas_sem_cabecalho=[letra(c) for c in sem_cabecalho_com_dado],
        colunas_ignoradas=[letra(c) for c in ignoradas],
        colunas_desconhecidas=desconhecidas, linhas=linhas)


def _linha(numero: int, valores: dict, aba) -> LinhaPlanilha:
    linha = LinhaPlanilha(linha=numero)
    linha.nome = " ".join(cl.texto_celula(valores.get("nome")).split())[:160]
    if not linha.nome:
        linha.acao = ACAO_IGNORAR
        linha.motivo = "Linha sem nome do cliente."

    bruto = valores.get("documento")
    linha.documento_informado = cl.texto_celula(bruto)
    if linha.documento_informado:
        try:
            linha.documento = cl.normalizar_documento(bruto)
        except cl.ErroClientes as exc:
            linha.pendencias.append(f"CNPJ/CPF inválido: {exc.mensagem}")
    else:
        linha.pendencias.append("Sem CNPJ/CPF.")

    for campo in ("ie", "im", "iptu"):
        valor, situacao = cl.classificar_campo(valores.get(campo))
        setattr(linha, campo, valor[:40])
        setattr(linha, f"{campo}_situacao", situacao)

    for campo in CAMPOS_SEGREDO:
        if campo in valores:
            valor, situacao = cl.classificar_campo(valores.get(campo),
                                                   preservar=True)
            if valor or situacao:
                linha.segredos[campo] = (valor, situacao)

    licencas = {
        "tvl": _licenca(valores.get("tvl"), valores.get("venc_tvl")),
        "alvara": _licenca(None, valores.get("venc_alvara")),
        "publicidade": _licenca(None, valores.get("publicidade")),
        "alvara_saude": _licenca(None, valores.get("alvara_saude")),
    }
    linha.licencas = {tipo: reg for tipo, reg in licencas.items()
                      if any(reg.values())}
    for tipo, reg in linha.licencas.items():
        if reg["situacao"] != cl.CAMPO_DATA_INVALIDA:
            continue
        if cl.erro_do_excel(reg["texto_original"]):
            linha.pendencias.append(
                f"{ROTULOS_LICENCA[tipo]}: a célula tem erro do Excel "
                f"({reg['texto_original'].strip()}), não uma data.")
        else:
            linha.pendencias.append(
                f"Data inválida em {ROTULOS_LICENCA[tipo]}: "
                f"'{reg['texto_original']}'.")

    linha.cor = _cor(aba.cell(numero, 1))
    _aplicar_cor(linha)
    dimensao = aba.row_dimensions.get(numero)
    linha.oculta = bool(dimensao is not None and dimensao.hidden)
    if linha.oculta:
        linha.pendencias.append("Linha oculta na planilha: confira se ainda vale.")
    return linha


ROTULOS_LICENCA = {
    "tvl": "TVL", "alvara": "Alvará", "publicidade": "Publicidade",
    "alvara_saude": "Alvará de saúde",
}
