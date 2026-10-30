"""Correções humanas da Base Fiscal e simulação de carga, sem persistência.

A fonte relida é a autoridade das colunas originais. O relatório identifica
essa fonte e transporta somente propostas: nenhuma decisão cadastral é aplicada.
"""
from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
from hashlib import file_digest
import math
import os
import re

from openpyxl import load_workbook

from .clientes import ErroClientes, chave_texto, texto_celula
from .documentos_br import DocumentoInvalido, digitos_cnpj
from . import planilha_empresas as pe
from .planilha_segura import PacoteRecusado


_CABECALHOS = (
    "ID de origem", "Aba de origem", "Linha de origem", "Empresa original",
    "CNPJ informado original", "CNPJ normalizado original", "CNPJ proposto",
    "Empresa proposta", "Regime proposto", "Responsável Fiscal proposto",
    "Contato proposto", "Decisão sobre duplicidade", "Observações da revisão",
)
_LIMITES = {"Resumo": 10001, "Conferência": 10001, "Duplicidades": 10001,
            "Correções": 10001, "Origens": 50001}
_COLUNAS = {"documento": "G", "nome": "H", "regime": "I",
            "responsavel_fiscal": "J", "contato_nome": "K"}
_ROTULOS = {"documento": "CNPJ", "nome": "Nome da empresa", "regime": "Regime",
            "responsavel_fiscal": "Responsável Fiscal", "contato_nome": "Contato"}
_ERROS_EXCEL = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!",
                "#NULL!", "#SPILL!", "#CALC!", "#GETTING_DATA"}
_CAMPOS_CADASTRO = ("id", "nome", "documento", "contato_nome", "situacao", "revision")


def _erro(mensagem: str) -> ErroClientes:
    return ErroClientes("PLANILHA_INVALIDA", mensagem)


def _texto(valor) -> str:
    return texto_celula(valor, preservar=True)


def _espacos(texto: str) -> str:
    return " ".join(texto.split())


def _incluir_unico(lista: list, mensagem: str) -> None:
    if mensagem not in lista:
        lista.append(mensagem)


def _cnpj(valor) -> str:
    """CNPJ completo; uma proposta não recupera zeros perdidos pelo Excel."""
    if isinstance(valor, bool):
        return ""
    if isinstance(valor, (int, float)):
        if (isinstance(valor, float) and not math.isfinite(valor)) or valor <= 0 or int(valor) != valor:
            return ""
        texto = str(int(valor))
    elif isinstance(valor, str):
        texto = valor.strip()
        if not re.fullmatch(r"[0-9. /-]+", texto):
            return ""
        texto = re.sub(r"[. /-]", "", texto)
    else:
        return ""
    if not re.fullmatch(r"[0-9]{14}", texto):
        return ""
    try:
        return digitos_cnpj(texto)
    except DocumentoInvalido:
        return ""


def _celula_segura(celula, aba: str, referencia: str, *, original: bool = False) -> None:
    if celula.data_type in {"f", "e"} or (
            not original and isinstance(celula.value, str) and celula.value.strip() in _ERROS_EXCEL):
        raise _erro(f"{aba}!{referencia}: fórmulas e erros de célula não são aceitos nas correções.")


def _original_confere(valor, esperado) -> bool:
    # CNPJ e identificadores do relatório são texto, enquanto linha é número.
    # O Excel não pode trocar esses tipos silenciosamente e perder zeros.
    if isinstance(esperado, int):
        return type(valor) in (int, float) and valor == esperado
    return (valor is None and esperado == "") or (isinstance(valor, str) and valor == esperado)


def _ler_propostas(leitura: pe.LeituraEmpresas, caminho: str) -> tuple[str, dict]:
    livro = None
    try:
        abas = pe._preflight(caminho, pe.MAX_PLANILHA_BYTES, limites_linhas=_LIMITES)
        por_nome = {a.nome: a for a in abas}
        if set(por_nome) != set(_LIMITES):
            raise _erro("Envie o relatório completo, com as cinco abas geradas pelo sistema.")
        for nome in ("Resumo", "Correções"):
            if por_nome[nome].mesclagens:
                raise _erro(f"A aba {nome} não pode conter células mescladas.")
        with open(caminho, "rb") as fonte:
            hash_correcoes = file_digest(fonte, "sha256").hexdigest()
        livro = load_workbook(caminho, read_only=True, data_only=False, keep_links=False)
        if [a.title for a in livro.worksheets] != [a.nome for a in abas]:
            raise _erro("A estrutura de abas diverge do pacote conferido.")
        resumo = livro["Resumo"]
        resumo.reset_dimensions()
        metadados = list(resumo.iter_rows(min_row=4, max_row=4, min_col=1, max_col=2))[0]
        for i, celula in enumerate(metadados):
            _celula_segura(celula, "Resumo", f"{'AB'[i]}4")
        hash_relatorio = metadados[1].value
        if metadados[0].value != "SHA-256 da fonte" or not isinstance(hash_relatorio, str) or not re.fullmatch(r"[0-9a-f]{64}", hash_relatorio):
            raise _erro("Resumo!A4:B4 deve preservar a identificação SHA-256 da fonte.")
        if hash_relatorio != leitura.hash_arquivo:
            raise ErroClientes("BASE_FISCAL_DIVERGENTE",
                               "O relatório pertence a outra fonte. Confira o arquivo original e o relatório correspondente.")
        ws = livro["Correções"]
        ws.reset_dimensions()
        fisica = por_nome["Correções"]
        originais = {l.id: l for l in leitura.linhas}
        if len(originais) != len(leitura.linhas) or len(originais) > pe.MAX_LINHAS:
            raise _erro("A fonte possui identificadores inválidos ou registros acima do limite.")
        propostas = {}
        for numero, row in enumerate(ws.iter_rows(max_row=max(fisica.max_linha, 1),
                                                   max_col=max(fisica.max_coluna, 13)), 1):
            for i, celula in enumerate(row):
                _celula_segura(celula, "Correções", f"{pe.get_column_letter(i + 1)}{numero}", original=i < 6)
            if any(c.value not in (None, "") for c in row[13:]):
                raise _erro(f"Correções, linha {numero}: coluna adicional preenchida.")
            valores = [c.value for c in row[:13]]
            if numero == 1:
                if tuple(valores) != _CABECALHOS:
                    raise _erro("Os 13 cabeçalhos da aba Correções devem ser preservados.")
                continue
            if all(v in (None, "") for v in valores):
                continue
            identificador = valores[0]
            if not isinstance(identificador, str) or identificador not in originais or identificador in propostas:
                raise _erro(f"Correções!A{numero}: ID de origem desconhecido ou repetido.")
            linha = originais[identificador]
            esperado = (linha.id, linha.aba, linha.linha, linha.nome,
                        linha.documento_informado, linha.documento)
            if not all(_original_confere(v, e) for v, e in zip(valores[:6], esperado)):
                raise _erro(f"Correções, linha {numero}: as colunas originais A:F foram alteradas.")
            propostas[identificador] = (numero, valores[6:])
        if set(propostas) != set(originais):
            raise _erro("O relatório deve conservar uma ocorrência de cada ID da fonte; há registros ausentes.")
        return hash_correcoes, propostas
    except ErroClientes:
        raise
    except PacoteRecusado as exc:
        raise _erro(exc.mensagem) from exc
    except Exception as exc:
        raise _erro("Não foi possível ler o relatório XLSX; confira sua estrutura e seu conteúdo.") from exc
    finally:
        if livro is not None:
            livro.close()


def _preparar_linha(linha: pe.LinhaEmpresa, numero: int, valores: list) -> dict:
    original = {"nome": linha.nome, "documento": linha.documento_informado,
                "regime": linha.regime_informado,
                "responsavel_fiscal": linha.responsavel_fiscal_informado,
                "contato_nome": linha.contato_nome}
    efetivo = {**original, "documento": linha.documento or linha.documento_informado,
               "regime": linha.regime or linha.regime_informado,
               "sem_movimento": linha.sem_movimento}
    proposto = {campo: _texto(v) for campo, v in zip(_COLUNAS, valores[:5])}
    pendencias, alteracoes = [], []
    for campo, bruto in zip(_COLUNAS, valores[:5]):
        proposta = proposto[campo]
        preenchida = bool(proposta.strip())
        if preenchida:
            antes = efetivo[campo]
            valor = _espacos(proposta)
            if campo == "documento":
                valor = _cnpj(bruto) or proposta.strip()
            elif campo == "regime":
                valor = pe._REGIMES.get(chave_texto(proposta), proposta.strip())
            elif not isinstance(bruto, str):
                pendencias.append(f"{_ROTULOS[campo]}: preencha a proposta como texto.")
            efetivo[campo] = valor
            if valor != antes:
                alteracoes.append({"campo": campo, "antes": original[campo],
                                   "depois": valor, "celula": f"Correções!{_COLUNAS[campo]}{numero}"})
        origem = linha.origem.get(campo, {})
        erro_fonte = any(p.startswith(f"Erro de célula em {campo} (") for p in linha.pendencias)
        if not preenchida and (origem.get("formula") or erro_fonte):
            pendencias.append(f"{_ROTULOS[campo]}: a fórmula ou erro da fonte exige uma proposta explícita.")
    if not _espacos(efetivo["nome"]):
        pendencias.append("Nome da empresa ausente.")
    if not _cnpj(efetivo["documento"]):
        pendencias.append("CNPJ ausente ou inválido; informe os 14 dígitos completos com verificador válido.")
    for campo in ("nome", "contato_nome", "responsavel_fiscal"):
        if len(efetivo[campo]) > 160:
            pendencias.append(f"{_ROTULOS[campo]} excede 160 caracteres.")
    if "regime" not in linha.campos_ausentes_no_layout or proposto["regime"].strip():
        if efetivo["regime"] not in pe._REGIMES.values():
            pendencias.append("Regime ausente ou não reconhecido.")
    if "responsavel_fiscal" not in linha.campos_ausentes_no_layout and not efetivo["responsavel_fiscal"].strip():
        pendencias.append("Responsável Fiscal ausente.")

    decisao_bruta, observacao_bruta = valores[5:7]
    decisao = chave_texto(_texto(decisao_bruta)).lower()
    observacao = _texto(observacao_bruta).strip()
    decisao_valida = decisao in {"", "manter", "ignorar"} and (decisao_bruta is None or isinstance(decisao_bruta, str))
    observacao_valida = (observacao_bruta is None or isinstance(observacao_bruta, str)) and len(observacao) <= 2000
    if not decisao_valida:
        pendencias.append("Decisão inválida: use manter, ignorar ou deixe vazio.")
    if not observacao_valida:
        pendencias.append("A observação deve ser texto de até 2.000 caracteres.")
    if decisao == "ignorar" and not observacao:
        pendencias.append("Ignorar exige justificativa na observação da revisão.")
    ignorar = decisao == "ignorar" and decisao_valida and observacao_valida and bool(observacao)
    return {"id": linha.id, "aba": linha.aba, "linha": linha.linha,
            "original": original, "proposto": proposto, "efetivo": efetivo,
            "alteracoes": alteracoes, "decisao": decisao, "observacao": observacao,
            "status": "ignorar" if ignorar else "pendente" if pendencias else "incluir",
            "pendencias": pendencias,
            "avisos": ["Responsável Fiscal informado sem vínculo com usuário. Regime e sem movimento não foram aplicados ao cadastro."],
            "cadastro": None, "campos_ausentes_no_layout": list(linha.campos_ausentes_no_layout)}


def _contagens(linhas: list[dict]) -> dict:
    return {"linhas": len(linhas),
            **{estado: sum(l["status"] == estado for l in linhas)
               for estado in ("incluir", "existente", "conflito", "pendente", "ignorar")},
            "alteradas": sum(bool(l["alteracoes"]) for l in linhas)}


def preparar_previa(leitura: pe.LeituraEmpresas, caminho_correcoes) -> dict:
    """Lê propostas e classifica sem mudar a fonte nem consultar um banco."""
    hash_correcoes, propostas = _ler_propostas(leitura, os.fspath(caminho_correcoes))
    linhas = [_preparar_linha(l, *propostas[l.id]) for l in leitura.linhas]
    grupos = defaultdict(list)
    for linha in linhas:
        documento = _cnpj(linha["efetivo"]["documento"])
        if documento:
            grupos[documento].append(linha)
    for grupo in grupos.values():
        if len(grupo) < 2 or all(l["status"] == "ignorar" for l in grupo):
            continue
        ativas = [l for l in grupo if l["status"] != "ignorar"]
        resolvido = len(ativas) == 1 and ativas[0]["decisao"] == "manter"
        if not resolvido:
            for linha in ativas:
                linha["status"] = "pendente"
                _incluir_unico(linha["pendencias"], "CNPJ repetido após as correções: marque uma ocorrência como manter e as demais como ignorar com justificativa, ou corrija os documentos.")
    return {"formato": "base_fiscal_importacao_v1", "somente_leitura": True,
            "hash_arquivo": leitura.hash_arquivo, "hash_correcoes": hash_correcoes,
            "contagens": _contagens(linhas), "linhas": linhas,
            "avisos": ["Esta prévia não grava dados e deve ser revalidada antes de qualquer futura importação.",
                       "Regime, responsável Fiscal e sem movimento são informações propostas; não geram vínculos, características ou tarefas.",
                       "Ignorar retira apenas a ocorrência desta proposta de carga; não exclui nem altera cadastros existentes."]}


def comparar_cadastro(previa: dict, clientes) -> dict:
    """Confronta CNPJs com um snapshot mínimo sem transformar pendência em aptidão."""
    resultado = deepcopy(previa)
    por_documento = defaultdict(list)
    for cliente in clientes:
        documento = cliente.get("documento")
        if documento:
            por_documento[documento].append(cliente)
    for linha in resultado["linhas"]:
        if linha["status"] == "ignorar":
            continue
        candidatos = por_documento.get(_cnpj(linha["efetivo"]["documento"]), [])
        if not candidatos:
            continue
        if len(candidatos) != 1:
            _incluir_unico(linha["pendencias"], "Mais de um cadastro possui este CNPJ; é necessária revisão do cadastro.")
            if linha["status"] != "pendente":
                linha["status"] = "conflito"
            continue
        cliente = candidatos[0]
        linha["cadastro"] = {c: cliente.get(c) for c in _CAMPOS_CADASTRO}
        diferencas = []
        for campo in ("nome", "contato_nome"):
            if campo in linha.get("campos_ausentes_no_layout", []) and not linha["proposto"].get(campo, "").strip():
                _incluir_unico(linha["avisos"], f"{_ROTULOS[campo]} ausente no layout: a proposta preserva o valor do cadastro existente.")
                continue  # Informação inexistente não propõe apagar o contato cadastrado.
            if _espacos(linha["efetivo"][campo]) != _espacos(str(cliente.get(campo) or "")):
                diferencas.append(f"{_ROTULOS[campo]} diverge do cadastro existente.")
        if cliente.get("situacao") != "ativo":
            diferencas.append("O cliente existente não está ativo; não será reativado por esta prévia.")
        for mensagem in diferencas:
            _incluir_unico(linha["avisos"], mensagem)
        if linha["status"] != "pendente":
            linha["status"] = "conflito" if diferencas else "existente"
    resultado["contagens"] = _contagens(resultado["linhas"])
    return resultado
