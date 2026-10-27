"""Leitura de carteiras XLS/XLSX: somente previa, sugestoes e proveniencia.

Checklist nao vira tarefa e sugestao nunca atribui cliente/caracteristica sozinha.
Formulas, inclusive resultados em cache do XLS, nao entram nos campos importados.
O chamador apaga o upload em finally, como na importacao do cadastro da aba 9.
"""
from __future__ import annotations

import hashlib
import io
import os
import re
import struct
from dataclasses import asdict, dataclass, field
from difflib import SequenceMatcher

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .clientes import ErroClientes, chave_texto, normalizar_documento, texto_celula
from .planilha_segura import PacoteRecusado, validar_pacote_xlsx
from .tarefas import ErroTarefas, validar_email

MAX_PLANILHA_BYTES = 20 * 1024 * 1024
MAX_LINHAS = 10000
MAX_COLUNAS = 100
MAX_ABAS = 50
MAX_TEXTO = 10000
_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_CAMPOS = {
    "EMPRESAS": "nome", "EMPRESA": "nome", "CLIENTE": "nome",
    "CLIENTES": "nome", "NOME": "nome", "RAZAO SOCIAL": "nome",
    "CNPJ": "documento", "CPF": "documento", "CNPJ/CPF": "documento", "CPF/CNPJ": "documento",
    "CONTATO": "contato_nome", "NOME DO CONTATO": "contato_nome",
    "LOCAL": "municipio", "MUNICIPIO": "municipio", "CIDADE": "municipio",
    "ATIVIDADE": "atividade", "E-MAIL": "contato_email", "EMAIL": "contato_email", "E MAIL": "contato_email",
    "OBSERVACAO": "observacao", "OBSERVACOES": "observacao", "OBS": "observacao", "OBS.": "observacao",
}
_SUGESTOES = (
    ("ADIANTAMENTO QUINZENAL", "Adiantamento quinzenal"),
    ("SEM MOVIMENTO", "Sem movimento"), ("LUCRO REAL", "Lucro real"),
    ("LUCRO PRESUMIDO", "Lucro presumido"), ("SIMPLES NACIONAL", "Simples Nacional"),
    ("MEI", "MEI"), ("CONTRIBUICAO NEGOCIAL", "Contribuição negocial"),
)


@dataclass
class LinhaCarteira:
    linha: int
    nome: str = ""
    documento: str = ""
    documento_informado: str = ""
    contato_nome: str = ""
    contato_email: str = ""
    contato_email_informado: str = ""
    municipio: str = ""
    atividade: str = ""
    observacao: str = ""
    sugestoes_caracteristicas: list[str] = field(default_factory=list)
    pendencias: list[str] = field(default_factory=list)
    origem: dict = field(default_factory=dict)

    def resumo(self) -> dict:
        return asdict(self)


@dataclass
class LeituraCarteira:
    hash_arquivo: str
    aba: str
    linha_cabecalho: int
    colunas: dict
    colunas_ignoradas: list
    linhas: list[LinhaCarteira]

    def resumo(self) -> dict:
        return {
            "hash_arquivo": self.hash_arquivo, "aba": self.aba,
            "linha_cabecalho": self.linha_cabecalho, "colunas": dict(self.colunas),
            "colunas_ignoradas": list(self.colunas_ignoradas),
            "contagens": {"total": len(self.linhas),
                          "com_pendencia": sum(bool(l.pendencias) for l in self.linhas)},
            "linhas": [l.resumo() for l in self.linhas],
        }


def _erro(mensagem: str) -> ErroTarefas:
    return ErroTarefas("PLANILHA_INVALIDA", mensagem)


def _limites(linhas: int, colunas: int, abas: int) -> None:
    if linhas > MAX_LINHAS or colunas > MAX_COLUNAS or abas > MAX_ABAS:
        raise _erro(f"A planilha excede o limite de {MAX_LINHAS} linhas, {MAX_COLUNAS} colunas ou {MAX_ABAS} abas.")


def _formulas_xls(conteudo: bytes) -> list[set[tuple[int, int]]]:
    """Preflight BIFF: marcas de formula por aba, antes de ler seus caches.

    OLE e' apenas desempacotado pelo leitor de container do xlrd. O tamanho e
    o cabecalho do container ja foram limitados; macros e criptografia recusados.
    """
    if conteudo.startswith(_OLE):
        if (len(conteudo) < 512 or conteudo[28:30] != b"\xfe\xff"
                or struct.unpack_from("<H", conteudo, 30)[0] not in (9, 12)
                or struct.unpack_from("<H", conteudo, 32)[0] != 6):
            raise _erro("Cabeçalho XLS inválido.")
        from xlrd.compdoc import CompDoc
        documento = CompDoc(conteudo, logfile=io.StringIO())
        if any(n.name.upper() in {"VBA", "_VBA_PROJECT", "_VBA_PROJECT_CUR", "MACROS"}
               for n in documento.dirlist):
            raise _erro("Planilha XLS com macro não é permitida.")
        stream = documento.get_named_stream("Workbook") or documento.get_named_stream("Book")
        if not stream or len(stream) > MAX_PLANILHA_BYTES:
            raise _erro("Não foi encontrado um livro XLS válido.")
    else:
        if len(conteudo) < 8 or struct.unpack_from("<H", conteudo)[0] not in (0x0009, 0x0209, 0x0409, 0x0809):
            raise _erro("O arquivo não tem assinatura de planilha XLS.")
        stream = conteudo
    formulas = {}
    offsets = []
    aba_atual = None
    pos = 0
    while pos + 4 <= len(stream):
        inicio = pos
        tipo, tamanho = struct.unpack_from("<HH", stream, pos)
        pos += 4
        if pos + tamanho > len(stream):
            raise _erro("Registro XLS truncado.")
        dados = stream[pos:pos + tamanho]
        pos += tamanho
        if tipo == 0x002F:
            raise _erro("Planilha XLS protegida por senha não é permitida.")
        if tipo == 0x0085 and len(dados) >= 4:
            offsets.append(struct.unpack_from("<I", dados)[0])
        if tipo in (0x0009, 0x0209, 0x0409, 0x0809) and len(dados) >= 4:
            subtipo = struct.unpack_from("<H", dados, 2)[0]
            if subtipo == 0x0040:
                raise _erro("Planilha XLS com macro não é permitida.")
            if subtipo == 0x0010:
                aba_atual = inicio
                formulas.setdefault(inicio, set())
        elif tipo == 0x000A:
            aba_atual = None
        elif tipo in (0x0006, 0x0206, 0x0406) and aba_atual is not None and len(dados) >= 4:
            formulas[aba_atual].add(struct.unpack_from("<HH", dados))
    ordem = offsets or list(formulas)
    return [formulas[o] for o in ordem if o in formulas]


def _ler_xls(conteudo: bytes) -> tuple[str, list[list[tuple[object, bool]]]]:
    import xlrd
    formulas = _formulas_xls(conteudo)
    livro = xlrd.open_workbook(file_contents=conteudo, on_demand=True, logfile=io.StringIO())
    try:
        if not livro.nsheets:
            raise _erro("A planilha não contém abas.")
        aba = livro.sheet_by_index(0)
        _limites(aba.nrows, aba.ncols, livro.nsheets)
        proibidas = formulas[0] if formulas else set()
        linhas = []
        for r in range(aba.nrows):
            linha = []
            for c in range(aba.ncols):
                celula = aba.cell(r, c)
                valor = celula.value
                if celula.ctype == xlrd.XL_CELL_DATE:
                    valor = xlrd.xldate_as_datetime(valor, livro.datemode)
                linha.append((valor, (r, c) in proibidas or celula.ctype == xlrd.XL_CELL_ERROR))
            linhas.append(linha)
        return aba.name, linhas
    finally:
        livro.release_resources()


def _ler_xlsx(caminho: str, max_bytes: int) -> tuple[str, list[list[tuple[object, bool]]]]:
    validar_pacote_xlsx(caminho, max_bytes=max_bytes)
    livro = load_workbook(caminho, read_only=True, data_only=False, keep_links=False)
    try:
        aba = livro.worksheets[0]
        _limites(aba.max_row or 0, aba.max_column or 0, len(livro.worksheets))
        # Dimensoes declaradas menores que os dados nao podem descartar registros.
        aba.reset_dimensions()
        linhas = []
        # Iterador tambem limitado: dimensoes declaradas por um arquivo hostil podem mentir.
        for r, linha in enumerate(aba.iter_rows(), 1):
            _limites(r, len(linha), len(livro.worksheets))
            linhas.append([(c.value, c.data_type in ("f", "e")) for c in linha])
        return aba.title, linhas
    finally:
        livro.close()


def ler_planilha_carteira(caminho: str, *, max_bytes: int = MAX_PLANILHA_BYTES) -> LeituraCarteira:
    caminho = os.fspath(caminho)
    if not caminho.lower().endswith((".xls", ".xlsx")):
        raise _erro("Envie uma planilha .xls ou .xlsx.")
    try:
        tamanho = os.path.getsize(caminho)
        if not tamanho or tamanho > min(max_bytes, MAX_PLANILHA_BYTES):
            raise _erro("Planilha vazia ou acima do limite de tamanho.")
        with open(caminho, "rb") as arquivo:
            conteudo = arquivo.read(MAX_PLANILHA_BYTES + 1)
        if len(conteudo) > MAX_PLANILHA_BYTES:
            raise _erro("A planilha excede o limite de tamanho.")
        sha = hashlib.sha256(conteudo).hexdigest()
        if caminho.lower().endswith(".xlsx"):
            aba, linhas = _ler_xlsx(caminho, max_bytes)
        else:
            aba, linhas = _ler_xls(conteudo)
        return _interpretar(sha, aba, linhas)
    except ErroTarefas:
        raise
    except PacoteRecusado as exc:
        raise _erro(exc.mensagem) from exc
    except Exception as exc:
        # Nao retorna o conteudo de uma celula, caminho do upload ou detalhes internos do parser.
        raise _erro("Não foi possível ler a planilha; confira se o arquivo está íntegro.") from exc


def _interpretar(sha: str, aba: str, linhas: list) -> LeituraCarteira:
    cabecalho = None
    colunas = {}
    ignoradas = []
    for numero, linha in enumerate(linhas[:10]):
        campos = [_CAMPOS.get(chave_texto(v)) if not formula else None for v, formula in linha]
        # Uma linha de titulo isolado 'Empresas' nao vence o cabecalho estruturado seguinte.
        if "nome" in campos and sum(c is not None for c in campos) >= 2:
            cabecalho = numero
            for c, ((valor, formula), campo) in enumerate(zip(linha, campos)):
                if campo:
                    if campo in colunas.values():
                        raise _erro(f"O cabeçalho repete o campo {campo}; mantenha uma coluna por campo.")
                    colunas[c] = campo
                elif texto_celula(valor) and not formula:
                    ignoradas.append(texto_celula(valor))
            break
    if cabecalho is None:
        raise ErroTarefas("CABECALHO_NAO_ENCONTRADO", "Não encontrei Empresas e outra coluna reconhecida nas primeiras dez linhas.")
    resultado = LeituraCarteira(sha, aba, cabecalho + 1,
                               {get_column_letter(c + 1): nome for c, nome in colunas.items()}, ignoradas, [])
    for numero, linha in enumerate(linhas[cabecalho + 1:], cabecalho + 2):
        valores = {}
        pendencias = []
        celulas = {}
        tem_conteudo = False
        for c, campo in colunas.items():
            valor, formula = linha[c] if c < len(linha) else (None, False)
            texto = texto_celula(valor)
            tem_conteudo |= bool(texto) or formula
            celula = f"{get_column_letter(c + 1)}{numero}"
            celulas[campo] = celula
            if formula:
                pendencias.append(f"Fórmula ou erro não importado em {celula}; substitua por valor.")
                valores[campo] = ""
            elif len(texto) > MAX_TEXTO:
                raise _erro(f"Texto acima do limite na célula {celula}.")
            else:
                valores[campo] = valor if campo == "documento" else texto
        if not tem_conteudo:
            continue
        registro = LinhaCarteira(numero, **{k: v for k, v in valores.items() if k != "documento"})
        registro.pendencias.extend(pendencias)
        registro.origem = {"hash_arquivo": sha, "aba": aba, "linha": numero, "celulas": celulas}
        documento = valores.get("documento")
        registro.documento_informado = texto_celula(documento)
        if registro.documento_informado:
            try:
                registro.documento = normalizar_documento(documento)
            except ErroClientes:
                registro.pendencias.append(f"CNPJ/CPF inválido em {celulas['documento']}; confirme o cliente manualmente.")
        registro.contato_email_informado = registro.contato_email
        try:
            registro.contato_email = validar_email(registro.contato_email)
        except ErroTarefas:
            registro.contato_email = ""
            registro.pendencias.append(f"E-mail inválido em {celulas['contato_email']}; corrija antes de importar o contato.")
        observacao = chave_texto(registro.observacao)
        registro.sugestoes_caracteristicas = [nome for chave, nome in _SUGESTOES
                                              if re.search(r"(?<!\w)" + re.escape(chave) + r"(?!\w)", observacao)]
        resultado.linhas.append(registro)
    return resultado


def sugerir_clientes(linha: LinhaCarteira | dict, clientes: list[dict]) -> list[dict]:
    linha = linha.resumo() if isinstance(linha, LinhaCarteira) else linha
    documento = linha.get("documento") or ""
    nome = chave_texto(linha.get("nome"))
    if documento:
        exatos = [{"cliente_id": c["id"], "nome": c["nome"], "tipo": "documento", "similaridade": 1.0}
                  for c in clientes if c.get("documento") == documento]
        if exatos:
            return exatos[:3]
    if not nome:
        return []
    sugestoes = []
    for cliente in clientes:
        chave = chave_texto(cliente.get("nome"))
        if not chave:
            continue
        similaridade = SequenceMatcher(None, nome, chave).ratio()
        if similaridade >= 0.35:
            sugestoes.append({"cliente_id": cliente["id"], "nome": cliente["nome"],
                              "tipo": "nome" if chave == nome else "semelhante",
                              "similaridade": round(similaridade, 4)})
    return sorted(sugestoes, key=lambda c: (-c["similaridade"], c["cliente_id"]))[:3]
