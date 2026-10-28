"""Previa da base Fiscal: dados de origem, sem cadastro ou associacao automatica.

O formato e reconhecido por cabecalhos explicitos. Nenhuma linha e mesclada com
outra, nem sem movimento e convertido em inatividade. O chamador e responsavel
pelo ciclo de vida do upload; este modulo apenas le e fecha o arquivo.
"""
from __future__ import annotations

import hashlib
import math
import os
import posixpath
import re
import zipfile
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from itertools import chain, islice

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .clientes import ErroClientes, chave_texto, texto_celula
from .documentos_br import DocumentoInvalido, digitos_cnpj
from .planilha_segura import PacoteRecusado, validar_pacote_xlsx

MAX_PLANILHA_BYTES = 20 * 1024 * 1024
MAX_LINHAS = 10000
MAX_LINHAS_FISICAS = 10020
MAX_COLUNAS = 32
MAX_ABAS = 20
MAX_REFERENCIAS_DUPLICADAS = 100000
_MARCADOR_FORMULA = "[fórmula não lida]"
_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CAMPOS = {"principal": {"responsavel_fiscal": 2, "nome": 3, "contato_nome": 4, "documento": 5, "regime": 6},
           "zeradas": {"nome": 2, "documento": 3}}
_REGIMES = {"SIMPLES NACIONAL": "Simples Nacional", "LUCRO REAL": "Lucro Real",
            "LUCRO PRESUMIDO": "Lucro Presumido", "PRESUMIDO": "Lucro Presumido"}


@dataclass
class LinhaEmpresa:
    id: str
    aba: str
    linha: int
    nome: str = ""
    documento: str = ""
    documento_informado: str = ""
    documento_status: str = "ausente"
    regime: str = ""
    regime_informado: str = ""
    responsavel_fiscal_informado: str = ""
    contato_nome: str = ""
    sem_movimento: bool | None = None
    duplicado: bool = False
    duplicado_em: list[str] = field(default_factory=list)
    pendencias: list[str] = field(default_factory=list)
    oculta: bool = False
    aba_oculta: bool = False
    origem: dict = field(default_factory=dict)


@dataclass
class LeituraEmpresas:
    hash_arquivo: str
    abas: list[dict]
    linhas: list[LinhaEmpresa]
    avisos: list[str]

    def resumo(self) -> dict:
        documentos = {l.documento for l in self.linhas if l.documento}
        repetidos = {l.documento for l in self.linhas if l.duplicado}
        return {
            "formato": "base_fiscal_v1", "hash_arquivo": self.hash_arquivo,
            "avisos": list(self.avisos), "abas": [dict(a) for a in self.abas],
            "contagens": {
                "linhas": len(self.linhas), "documentos_validos": sum(bool(l.documento) for l in self.linhas),
                "cnpjs_distintos": len(documentos), "documentos_pendentes": sum(not l.documento for l in self.linhas),
                "grupos_duplicados": len(repetidos), "linhas_duplicadas": sum(l.duplicado for l in self.linhas),
                "sem_movimento": sum(l.sem_movimento is True for l in self.linhas),
                "sem_responsavel": sum(not l.responsavel_fiscal_informado.strip() for l in self.linhas),
                "sem_regime": sum(not l.regime for l in self.linhas),
                "com_pendencias": sum(bool(l.pendencias) for l in self.linhas)},
            "linhas": [asdict(l) for l in self.linhas],
        }


@dataclass
class _AbaFisica:
    nome: str
    caminho: str
    oculta: bool
    max_linha: int = 0
    max_coluna: int = 0
    ocultas: set[int] = field(default_factory=set)
    mesclagens: list[tuple[str, int, int, int, int]] = field(default_factory=list)


def _erro(mensagem: str) -> ErroClientes:
    return ErroClientes("PLANILHA_INVALIDA", mensagem)


def _inteiro(valor, *, limite: int) -> int:
    if not re.fullmatch(r"[0-9]{1,8}", str(valor or "")):
        raise _erro("Coordenada inválida na estrutura da planilha.")
    numero = int(valor)
    if not 1 <= numero <= limite:
        raise _erro(f"A planilha excede o limite de {MAX_LINHAS_FISICAS} linhas físicas ou {MAX_COLUNAS} colunas por aba.")
    return numero


def _coordenada(valor: str) -> tuple[int, int]:
    achado = re.fullmatch(r"\$?([A-Z]{1,3})\$?([0-9]{1,8})", valor or "")
    if not achado:
        raise _erro("Referência de célula inválida na planilha.")
    coluna = 0
    for letra in achado[1]:
        coluna = coluna * 26 + ord(letra) - ord("A") + 1
    return _inteiro(achado[2], limite=MAX_LINHAS_FISICAS), _inteiro(str(coluna), limite=MAX_COLUNAS)


def _intervalo(valor: str) -> tuple[int, int, int, int]:
    partes = valor.split(":")
    if not 1 <= len(partes) <= 2:
        raise _erro("Intervalo de células inválido.")
    r1, c1 = _coordenada(partes[0])
    r2, c2 = _coordenada(partes[-1])
    if r2 < r1 or c2 < c1:
        raise _erro("Intervalo de células invertido.")
    return r1, c1, r2, c2


def _xml_seguro(dados: bytes):
    parser = etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True,
                             huge_tree=False, recover=False)
    raiz = etree.fromstring(dados, parser)
    if raiz.getroottree().docinfo.doctype:
        raise _erro("XML com DTD ou entidades não é permitido na planilha.")
    return raiz


def _medir_aba(raiz, aba: _AbaFisica) -> None:
    if raiz.tag != f"{{{_NS}}}worksheet":
        raise _erro("A base Fiscal deve conter somente abas de planilha.")
    if len(list(raiz.iter(f"{{{_NS}}}sheetData"))) > 1:
        raise _erro("A aba contém mais de uma área de dados no XML.")
    # O parser de alto nivel percorre eventos XML. Conferir todas as ocorrencias
    # evita que tags fora da estrutura escapem da medicao por um XPath restrito.
    for tag, pai in (("row", "sheetData"), ("c", "row"), ("col", "cols"), ("mergeCell", "mergeCells")):
        for elemento in raiz.iter(f"{{{_NS}}}{tag}"):
            anterior = elemento.getparent()
            if anterior is None or anterior.tag != f"{{{_NS}}}{pai}":
                raise _erro("Elemento fora da estrutura de células da planilha.")
            esperado = anterior.getparent()
            if tag == "c":
                if esperado is None or esperado.tag != f"{{{_NS}}}sheetData":
                    raise _erro("Célula fora da área de dados da planilha.")
                esperado = esperado.getparent()
            if esperado is not raiz:
                raise _erro("Elemento fora da área de dados da planilha.")
    ultima_linha = 0
    for elemento in raiz.iter(f"{{{_NS}}}row"):
        numero = _inteiro(elemento.get("r", str(ultima_linha + 1)), limite=MAX_LINHAS_FISICAS)
        if numero <= ultima_linha:
            raise _erro("Linhas duplicadas ou fora de ordem no XML da planilha.")
        ultima_linha = numero
        aba.max_linha = max(aba.max_linha, numero)
        if elemento.get("hidden") in {"1", "true"}:
            aba.ocultas.add(numero)
        ultima_coluna = 0
        for celula in elemento.findall(f"{{{_NS}}}c"):
            if "r" in celula.attrib:
                linha, coluna = _coordenada(celula.get("r"))
                if linha != numero:
                    raise _erro("Célula fora da sua linha no XML da planilha.")
            else:
                coluna = _inteiro(str(ultima_coluna + 1), limite=MAX_COLUNAS)
            if coluna <= ultima_coluna:
                raise _erro("Colunas duplicadas ou fora de ordem no XML da planilha.")
            ultima_coluna = coluna
            aba.max_coluna = max(aba.max_coluna, coluna)
    for col in raiz.findall(f"{{{_NS}}}cols/{{{_NS}}}col"):
        inicio = _inteiro(col.get("min"), limite=MAX_COLUNAS)
        fim = _inteiro(col.get("max"), limite=MAX_COLUNAS)
        if inicio > fim:
            raise _erro("Faixa de colunas inválida na planilha.")
    for merge in raiz.findall(f"{{{_NS}}}mergeCells/{{{_NS}}}mergeCell"):
        referencia = merge.get("ref", "")
        r1, c1, r2, c2 = _intervalo(referencia)
        aba.mesclagens.append((referencia, r1, c1, r2, c2))
        aba.max_linha = max(aba.max_linha, r2)
        aba.max_coluna = max(aba.max_coluna, c2)


def _preflight(caminho: str, max_bytes: int) -> list[_AbaFisica]:
    validar_pacote_xlsx(caminho, max_bytes=max_bytes)
    with zipfile.ZipFile(caminho) as pacote:
        nomes = pacote.namelist()
        if len(nomes) != len(set(nomes)):
            raise _erro("O pacote contém entradas duplicadas.")
        livro = _xml_seguro(pacote.read("xl/workbook.xml"))
        folhas = livro.findall(f"{{{_NS}}}sheets/{{{_NS}}}sheet")
        if not folhas or len(folhas) > MAX_ABAS:
            raise _erro(f"A planilha deve ter entre 1 e {MAX_ABAS} abas.")
        relacoes = _xml_seguro(pacote.read("xl/_rels/workbook.xml.rels"))
        por_id = {}
        for rel in relacoes:
            if rel.get("Id") in por_id:
                raise _erro("O pacote contém relações ambíguas.")
            por_id[rel.get("Id")] = rel
        abas = []
        for folha in folhas:
            rel = por_id.get(folha.get(f"{{{_NS_REL}}}id"))
            if rel is None or rel.get("TargetMode") == "External" or not rel.get("Type", "").endswith("/worksheet"):
                raise _erro("A base Fiscal deve conter somente abas internas de planilha.")
            alvo = rel.get("Target", "")
            interno = posixpath.normpath(alvo.lstrip("/") if alvo.startswith("/") else posixpath.join("xl", alvo))
            if not interno.startswith("xl/") or interno not in nomes:
                raise _erro("Referência de aba inválida no pacote.")
            abas.append(_AbaFisica(folha.get("name", ""), interno, folha.get("state", "visible") != "visible"))
        if len({a.nome for a in abas}) != len(abas) or len({a.caminho for a in abas}) != len(abas):
            raise _erro("O pacote contém abas ambíguas.")
        por_caminho = {a.caminho: a for a in abas}
        # Tambem confere XML auxiliares antes que o openpyxl os interprete.
        # lxml nao expande entidades, carrega DTD nem acessa rede neste passe.
        for nome in nomes:
            if nome in por_caminho or nome.lower().endswith((".xml", ".rels")):
                raiz = _xml_seguro(pacote.read(nome))
                if nome in por_caminho:
                    _medir_aba(raiz, por_caminho[nome])
        return abas


def _preenchida(celula) -> bool:
    return celula.data_type == "f" or bool(texto_celula(celula.value).strip())


def _tipo_cabecalho(linha, nome_aba: str) -> str | None:
    titulos = [chave_texto(c.value) if c.data_type != "f" else "[FORMULA]" for c in linha]
    titulos += [""] * max(0, 6 - len(titulos))
    if chave_texto(nome_aba) == "ZERADAS":
        if titulos[:3] == ["", "EMPRESA", "CNPJ"] and not any(titulos[3:]):
            return "zeradas"
    elif (titulos[:4] == ["", "RESPONSAVEL", "EMPRESAS", "RESPONSAVEL"]
          and titulos[4] in {"ZERADAS", "CNPJ"} and titulos[5] == "REGIME DA EMPRESA"
          and not any(titulos[6:])):
        return "principal"
    return None


def _mascara_cnpj(formato: str) -> bool:
    # Aceita a pontuacao literal simples, escapada ou entre aspas. Um padrao
    # generico de 14 zeros nao demonstra que o documento original era CNPJ.
    formato = re.sub(r'"([^"\r\n]*)"', r"\1", formato)
    formato = re.sub(r"\\(.)", r"\1", formato)
    return formato == "00.000.000/0000-00"


def _documento(celula) -> tuple[str, str]:
    valor = celula.value
    if celula.data_type == "f":
        return "", "formula"
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return "", "ausente"
    status = "valido"
    if isinstance(valor, bool) or celula.data_type == "e":
        return "", "invalido"
    if isinstance(valor, (int, float)):
        if (isinstance(valor, float) and not math.isfinite(valor)) or valor <= 0 or int(valor) != valor:
            return "", "invalido"
        texto = str(int(valor))
        if len(texto) == 13 and _mascara_cnpj(celula.number_format):
            texto = "0" + texto
            status = "zero_inicial_recuperado"
        elif len(texto) != 14:
            return "", "invalido"
    elif isinstance(valor, str):
        texto = valor.strip()
        if not re.fullmatch(r"[0-9. /-]+", texto):
            return "", "invalido"
        texto = re.sub(r"[. /-]", "", texto)
        if len(texto) != 14:
            return "", "invalido"
    else:
        return "", "invalido"
    try:
        return digitos_cnpj(texto), status
    except DocumentoInvalido:
        return "", "invalido"


def _mesclagens(aba: _AbaFisica, tipo: str, cabecalho: int) -> dict[int, tuple[str, int]]:
    contatos = {}
    for intervalo, r1, c1, r2, c2 in aba.mesclagens:
        if tipo != "principal" or c1 != 4 or c2 != 4 or r1 <= cabecalho or r2 <= r1:
            raise _erro(f"Mesclagem {intervalo} na aba {aba.nome}: somente contato vertical na área de dados é permitido.")
        for numero in range(r1, r2 + 1):
            if numero in contatos:
                raise _erro(f"Mesclagens de contato sobrepostas na aba {aba.nome}.")
            contatos[numero] = (intervalo, r1)
    return contatos


def _ler_aba(ws, aba: _AbaFisica, indice: int, resultado: LeituraEmpresas) -> None:
    ws.reset_dimensions()  # Dimensao declarada nao substitui os limites reais do XML.
    rows = enumerate(ws.iter_rows(max_row=max(aba.max_linha, 1), max_col=max(aba.max_coluna, 6)), 1)
    primeiras = list(islice(rows, 10))
    candidatos = [(numero, _tipo_cabecalho(linha, aba.nome)) for numero, linha in primeiras]
    candidatos = [(numero, tipo) for numero, tipo in candidatos if tipo]
    if len(candidatos) > 1:
        raise _erro(f"Cabeçalhos ambíguos na aba {aba.nome}.")
    if not candidatos:
        if any(_preenchida(c) for _, row in chain(primeiras, rows) for c in row):
            raise _erro(f"Layout não reconhecido na aba {aba.nome}; confira os cabeçalhos nas primeiras dez linhas.")
        if aba.mesclagens:
            raise _erro(f"Mesclagens fora da área de dados na aba {aba.nome}.")
        resultado.avisos.append(f"Aba vazia ignorada: {aba.nome}.")
        return
    cabecalho, tipo = candidatos[0]
    campos = _CAMPOS[tipo]
    ultima_coluna = max(campos.values())
    contatos = _mesclagens(aba, tipo, cabecalho)
    ancoras = {}
    total = 0
    for numero, row in chain(primeiras, rows):
        if any(_preenchida(c) for c in row[ultima_coluna:]):
            raise _erro(f"Coluna adicional preenchida na aba {aba.nome}, linha {numero}.")
        if numero <= cabecalho:
            continue
        if _tipo_cabecalho(row, aba.nome):
            raise _erro(f"Cabeçalho repetido na aba {aba.nome}, linha {numero}.")
        # Utilidade e definida antes de propagar o contato da celula ancorada.
        util = any(_preenchida(row[coluna - 1]) for coluna in campos.values())
        if numero in contatos and contatos[numero][1] == numero:
            ancoras[numero] = row[3]
        if not util:
            continue
        if len(resultado.linhas) >= MAX_LINHAS:
            raise _erro(f"A planilha excede o limite de {MAX_LINHAS} linhas úteis no arquivo.")
        item = LinhaEmpresa(id=f"{indice}:{numero}", aba=aba.nome, linha=numero,
            sem_movimento=True if tipo == "zeradas" else None,
            oculta=numero in aba.ocultas, aba_oculta=aba.oculta)
        valores = {}
        for campo, coluna in campos.items():
            celula = row[coluna - 1]
            origem_linha = numero
            intervalo = None
            if campo == "contato_nome" and numero in contatos:
                intervalo, origem_linha = contatos[numero]
                if origem_linha != numero and _preenchida(celula):
                    raise _erro(f"Contato preenchido dentro da mesclagem {intervalo} na aba {aba.nome}; confira as células antes da prévia.")
                celula = ancoras[origem_linha]
            formula = celula.data_type == "f"
            informado = _MARCADOR_FORMULA if formula else texto_celula(celula.value, preservar=True)
            item.origem[campo] = dict(celula=f"{get_column_letter(coluna)}{numero}",
                celula_origem=f"{get_column_letter(coluna)}{origem_linha}", valor_informado=informado,
                formato=celula.number_format, intervalo_mesclado=intervalo,
                propagado=numero != origem_linha, formula=formula)
            valores[campo] = "" if formula or celula.data_type == "e" else informado.strip()
            if formula:
                item.pendencias.append(f"Fórmula não lida em {campo} ({item.origem[campo]['celula']}).")
            elif celula.data_type == "e":
                item.pendencias.append(f"Erro de célula em {campo} ({item.origem[campo]['celula']}).")
            if campo == "documento":
                item.documento_informado = informado
                item.documento, item.documento_status = _documento(celula)
            elif campo == "regime":
                item.regime_informado = informado
        item.nome = valores["nome"]
        item.responsavel_fiscal_informado = valores.get("responsavel_fiscal", "")
        item.contato_nome = valores.get("contato_nome", "")
        item.regime = _REGIMES.get(chave_texto(valores.get("regime", "")), "")
        if not item.nome:
            item.pendencias.append("Nome da empresa ausente.")
        if item.documento_status == "ausente":
            item.pendencias.append("CNPJ ausente.")
        elif item.documento_status == "invalido":
            item.pendencias.append("CNPJ inválido; conferir o documento informado.")
        if not item.responsavel_fiscal_informado:
            item.pendencias.append("Responsável Fiscal ausente.")
        if not item.regime:
            item.pendencias.append("Regime ausente ou não reconhecido.")
        resultado.linhas.append(item)
        total += 1
    resultado.abas.append(dict(nome=aba.nome, tipo=tipo, linha_cabecalho=cabecalho, total=total, oculta=aba.oculta))


def ler_planilha_empresas(caminho, *, max_bytes: int = MAX_PLANILHA_BYTES) -> LeituraEmpresas:
    """Le XLSX reconhecido e devolve previa completa; erro nunca devolve parcial."""
    livro = None
    try:
        caminho = os.fspath(caminho)
        abas = _preflight(caminho, max_bytes)
        with open(caminho, "rb") as arquivo:
            hash_arquivo = hashlib.file_digest(arquivo, "sha256").hexdigest()
        resultado = LeituraEmpresas(hash_arquivo, [], [], [])
        livro = load_workbook(caminho, read_only=True, data_only=False, keep_links=False)
        if [ws.title for ws in livro.worksheets] != [a.nome for a in abas]:
            raise _erro("A estrutura de abas não corresponde ao pacote conferido.")
        for indice, (ws, aba) in enumerate(zip(livro.worksheets, abas), 1):
            _ler_aba(ws, aba, indice, resultado)
        if not resultado.abas:
            raise _erro("Não foi encontrada uma aba com layout reconhecido.")
        por_documento = defaultdict(list)
        for linha in resultado.linhas:
            if linha.documento:
                por_documento[linha.documento].append(linha)
        referencias = sum(len(grupo) * (len(grupo) - 1) for grupo in por_documento.values())
        if referencias > MAX_REFERENCIAS_DUPLICADAS:
            raise _erro(f"Duplicidades excessivas: a prévia excede o limite de {MAX_REFERENCIAS_DUPLICADAS} referências entre ocorrências. Separe a base em arquivos menores para conferência; nenhuma linha foi descartada ou agrupada.")
        for grupo in por_documento.values():
            if len(grupo) < 2:
                continue
            for linha in grupo:
                linha.duplicado = True
                linha.duplicado_em = [outra.id for outra in grupo if outra.id != linha.id]
                linha.pendencias.append("CNPJ duplicado na planilha; conferir todas as ocorrências.")
        return resultado
    except ErroClientes:
        raise
    except PacoteRecusado as exc:
        raise _erro(exc.mensagem) from exc
    except Exception as exc:
        # Nao inclui texto de formulas, XML nem caminhos locais nos erros da API.
        raise _erro("Não foi possível ler a planilha XLSX; confira sua estrutura e seu conteúdo.") from exc
    finally:
        if livro is not None:
            livro.close()
